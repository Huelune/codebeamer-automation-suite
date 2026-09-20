from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment

from src.gui.service_core import GuiCodebeamerService
from src.gui.service_core import GuiExcelService
from src.gui.settings_store import GuiSettings
from tests.gui_service_fixtures import CountingBatchExcelReader
from tests.gui_service_fixtures import FakeClient
from tests.gui_service_fixtures import FakeExcelReader
from tests.gui_widget_cleanup import tearDownModule  # noqa: F401


class GuiCodebeamerServiceTest(unittest.TestCase):
    def test_connection_and_tracker_loading(self) -> None:
        service = GuiCodebeamerService(client_factory=FakeClient)
        settings = GuiSettings(
            base_url="https://example.com/cb",
            username="user",
            password="secret",
            rate_limit_retry_delay_seconds=1.0,
            rate_limit_max_retries=5,
        )

        projects = service.test_connection_and_load_projects(settings)
        trackers = service.load_trackers(settings, 10)

        self.assertEqual(projects[0]["id"], 10)
        self.assertEqual(projects[0]["name"], "Project A")
        self.assertEqual(trackers[0]["id"], 1000)
        self.assertEqual(trackers[1]["name"], "Tracker 10-2")

    def test_offline_connection_and_tracker_loading_use_local_schema_snapshot(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            schema_path = Path(tmp_dir) / "offline-schema.json"
            schema_payload = FakeClient("", "", "").get_tracker_schema(0)
            schema_payload["name"] = "Offline Tracker Snapshot"
            schema_path.write_text(json.dumps(schema_payload, ensure_ascii=False), encoding="utf-8")

            service = GuiCodebeamerService(client_factory=FakeClient)
            settings = GuiSettings(
                offline_mode=True,
                offline_schema_path=str(schema_path),
                default_project_id="501",
                default_tracker_id="601",
            )

            projects = service.test_connection_and_load_projects(settings)
            trackers = service.load_trackers(settings, 501)

            self.assertEqual(projects, [{"id": 501, "name": "Offline Project"}])
            self.assertEqual(trackers, [{"id": 601, "name": "Offline Tracker Snapshot"}])


class GuiExcelServiceTest(unittest.TestCase):
    @staticmethod
    def _write_workbook(path: Path, summary: str) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Main"
        sheet.append(["Summary", "담당자"])
        sheet.append([summary, "홍길동"])
        workbook.create_sheet("Other")
        workbook.save(path)
        workbook.close()

    def test_load_preview_reads_sheet_names_headers_rows_and_summary(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "담당자", "비고"])
            sheet.append(["REQ-001", "홍길동", "메모"])
            sheet.append(["REQ-002", "김철수", "메모2"])
            workbook.create_sheet("Other")
            workbook.save(path)
            workbook.close()

            preview = GuiExcelService().load_preview(
                str(path),
                sheet_name="Main",
                header_row=1,
                max_preview_rows=5,
            )

            self.assertEqual(preview.sheet_names, ["Main", "Other"])
            self.assertEqual(preview.headers, ["Summary", "담당자", "비고"])
            self.assertEqual(preview.rows[0], ["REQ-001", "홍길동", "메모"])
            self.assertEqual(preview.suggested_summary, "Summary")

    def test_load_preview_displays_integer_like_numbers_without_decimal_suffix(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "번호", "실수"])
            sheet.append(["REQ-001", 101, 1.5])
            sheet.append(["REQ-002", 202, 2.0])
            workbook.save(path)
            workbook.close()

            preview = GuiExcelService(reader_cls=FakeExcelReader).load_preview(
                str(path),
                sheet_name="Main",
                header_row=1,
                max_preview_rows=5,
            )

            self.assertEqual(preview.rows[0], ["REQ-001", "101", "1.5"])
            self.assertEqual(preview.rows[1], ["REQ-002", "202", "2"])

    def test_load_preview_preloads_raw_data_for_all_selected_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            first_path = Path(tmp_dir) / "first.xlsx"
            second_path = Path(tmp_dir) / "second.xlsx"
            for path, summary in (
                (first_path, "REQ-001"),
                (second_path, "REQ-002"),
            ):
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Main"
                sheet.append(["Summary", "담당자"])
                sheet.append([summary, "홍길동"])
                workbook.save(path)
                workbook.close()

            CountingBatchExcelReader.reset_counts()
            preview = GuiExcelService(reader_cls=CountingBatchExcelReader).load_preview(
                str(first_path),
                file_paths=[str(first_path), str(second_path)],
                sheet_name="Main",
                header_row=1,
                summary_column="Summary",
            )

            self.assertEqual(
                sorted(preview.raw_df_by_file.keys()),
                sorted([str(first_path), str(second_path)]),
            )
            self.assertEqual(
                CountingBatchExcelReader.read_excel_calls,
                [str(first_path), str(second_path)],
            )

    def test_staged_loading_reads_metadata_then_limited_preview_before_full_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            first_path = Path(tmp_dir) / "first.xlsx"
            second_path = Path(tmp_dir) / "second.xlsx"
            self._write_workbook(first_path, "REQ-001")
            self._write_workbook(second_path, "REQ-002")
            service = GuiExcelService(reader_cls=CountingBatchExcelReader)
            CountingBatchExcelReader.reset_counts()

            metadata = service.load_metadata(str(first_path))

            self.assertEqual(metadata.sheet_names, ["Main", "Other"])
            self.assertEqual(CountingBatchExcelReader.list_sheet_calls, [str(first_path)])
            self.assertEqual(CountingBatchExcelReader.read_excel_calls, [])

            sheet_preview = service.load_sheet_preview(
                str(first_path),
                sheet_name="Main",
                header_row=1,
                summary_column="Summary",
                max_preview_rows=1,
            )

            self.assertEqual(sheet_preview.rows, [["REQ-001", "홍길동"]])
            self.assertEqual(CountingBatchExcelReader.read_preview_calls, [str(first_path)])
            self.assertEqual(CountingBatchExcelReader.read_excel_calls, [])

            preview = service.load_full_data(
                str(first_path),
                file_paths=[str(first_path), str(second_path)],
                sheet_name="Main",
                header_row=1,
                summary_column="Summary",
                sheet_preview=sheet_preview,
            )

            self.assertFalse(preview.cache_hit)
            self.assertEqual(
                CountingBatchExcelReader.read_excel_calls,
                [str(first_path), str(second_path)],
            )

            other_summary = service.load_full_data(
                str(first_path),
                file_paths=[str(first_path), str(second_path)],
                sheet_name="Main",
                header_row=1,
                summary_column="담당자",
            )
            other_representative = service.load_full_data(
                str(second_path),
                file_paths=[str(first_path), str(second_path)],
                sheet_name="Main",
                header_row=1,
                summary_column="Summary",
            )

            self.assertFalse(other_summary.cache_hit)
            self.assertEqual(other_representative.file_path, str(second_path))
            self.assertEqual(
                CountingBatchExcelReader.read_excel_calls,
                [
                    str(first_path),
                    str(second_path),
                    str(first_path),
                    str(second_path),
                ],
            )
            self.assertEqual(
                sorted(preview.raw_df_by_file),
                sorted([str(first_path), str(second_path)]),
            )

            cached = service.load_full_data(
                str(first_path),
                file_paths=[str(first_path), str(second_path)],
                sheet_name="Main",
                header_row=1,
                summary_column="Summary",
                sheet_preview=sheet_preview,
            )

            self.assertTrue(cached.cache_hit)
            self.assertEqual(
                CountingBatchExcelReader.read_excel_calls,
                [
                    str(first_path),
                    str(second_path),
                    str(first_path),
                    str(second_path),
                ],
            )

    def test_summary_change_reloads_indent_metadata_for_selected_column(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "indent-by-summary.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "Other"])
            sheet.append(["REQ-001", "Group A"])
            sheet.append(["REQ-002", "Group B"])
            sheet["A2"].alignment = Alignment(indent=0)
            sheet["A3"].alignment = Alignment(indent=2)
            sheet["B2"].alignment = Alignment(indent=3)
            sheet["B3"].alignment = Alignment(indent=0)
            workbook.save(path)
            workbook.close()

            service = GuiExcelService()
            summary_data = service.load_full_data(
                str(path),
                sheet_name="Main",
                header_row=1,
                summary_column="Summary",
            )
            other_data = service.load_full_data(
                str(path),
                sheet_name="Main",
                header_row=1,
                summary_column="Other",
            )

            self.assertEqual(
                summary_data.raw_df_by_file[str(path)]["_summary_indent"].tolist(),
                [0, 2],
            )
            self.assertEqual(
                other_data.raw_df_by_file[str(path)]["_summary_indent"].tolist(),
                [3, 0],
            )

    def test_full_data_cache_is_invalidated_when_source_file_changes(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            self._write_workbook(path, "REQ-001")
            service = GuiExcelService(reader_cls=CountingBatchExcelReader)
            CountingBatchExcelReader.reset_counts()

            first = service.load_full_data(
                str(path),
                sheet_name="Main",
                header_row=1,
                summary_column="Summary",
            )
            self._write_workbook(path, "REQ-001-수정")

            self.assertFalse(first.files_are_current())

            second = service.load_full_data(
                str(path),
                sheet_name="Main",
                header_row=1,
                summary_column="Summary",
            )

            self.assertFalse(second.cache_hit)
            self.assertEqual(second.rows[0][0], "REQ-001-수정")
            self.assertEqual(
                CountingBatchExcelReader.read_excel_calls,
                [str(path), str(path)],
            )

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from src.gui.service_core import GuiExcelService
from src.gui.settings_store import GuiSettings
from src.gui.upload_service import GuiUploadPipelineService
from src.models import PayloadStatus
from tests.gui_service_fixtures import FakeExcelReader
from tests.gui_service_fixtures import MemberReferenceDefaultFakeClient
from tests.gui_service_fixtures import TrackerItemDefaultValueFakeClient
from tests.gui_service_fixtures import TrackerItemNonTrackerConfigFakeClient
from tests.gui_service_fixtures import TrackerItemQueryFakeClient
from tests.gui_service_fixtures import TrackerItemReferenceIdConfigFakeClient
from tests.gui_service_fixtures import UserReferenceDefaultFakeClient


class GuiLookupPipelineServiceTest(unittest.TestCase):
    def test_prepare_mapping_context_uses_id_extraction_for_tracker_item_defaults(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            TrackerItemQueryFakeClient.all_search_calls = []
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "연관 요구사항"])
            sheet.append(["REQ-001", "[REQ:100]"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=TrackerItemQueryFakeClient,
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            self.assertEqual(
                mapping_context.selected_tracker_item_settings["연관 요구사항"]["mode"],
                "regex",
            )

    def test_prepare_mapping_context_uses_tracker_configuration_reference_id_when_name_differs(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "SUDS 링크"])
            sheet.append(["REQ-001", "[REQ:100]"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=TrackerItemReferenceIdConfigFakeClient,
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            self.assertEqual(
                mapping_context.selected_tracker_item_settings["SUDS 링크"]["mode"],
                "regex",
            )
            tracker_field_row = mapping_context.schema_df[
                mapping_context.schema_df["field_name"] == "SUDS 링크"
            ].iloc[0]
            self.assertEqual(int(tracker_field_row["field_id"]), 17)
            self.assertEqual(tracker_field_row["tracker_item_source_tracker_ids"], [13526611])

    def test_prepare_mapping_context_disables_query_for_non_tracker_configuration(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "연관 요구사항"])
            sheet.append(["REQ-001", "[REQ:100]"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=TrackerItemNonTrackerConfigFakeClient,
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            self.assertEqual(
                mapping_context.selected_tracker_item_settings["연관 요구사항"]["mode"],
                "regex",
            )
            tracker_item_candidate = next(
                candidate
                for candidate in mapping_context.tracker_item_field_candidates
                if candidate.schema_field == "연관 요구사항"
            )
            self.assertEqual(tracker_item_candidate.query_status, "unsupported")

    def test_validate_mapping_uses_id_extraction_without_tracker_item_query(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            TrackerItemQueryFakeClient.all_search_calls = []
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "연관 요구사항"])
            sheet.append(["REQ-001", "[REQ:100]"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=TrackerItemQueryFakeClient,
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            validation_context = service.validate_mapping(
                mapping_context,
                {"Summary": "Summary", "연관 요구사항": "연관 요구사항"},
                selected_tracker_item_settings=mapping_context.selected_tracker_item_settings,
            )

            self.assertFalse(validation_context.has_blocking_issues)
            self.assertTrue(validation_context.issue_df.empty)
            self.assertEqual(
                list(mapping_context.wizard.state.payload_df["payload_status"]),
                [PayloadStatus.READY.value],
            )
            self.assertEqual(
                TrackerItemQueryFakeClient.all_search_calls,
                [],
            )

    def test_prepare_mapping_context_includes_user_reference_field_in_default_value_candidates(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary"])
            sheet.append(["REQ-001"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=UserReferenceDefaultFakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            default_candidates = {
                candidate.schema_field: candidate
                for candidate in mapping_context.default_value_candidates
            }
            self.assertIn("담당 사용자", default_candidates)
            self.assertTrue(default_candidates["담당 사용자"].allows_custom_value)

    def test_prepare_mapping_context_includes_member_field_in_default_value_candidates(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary"])
            sheet.append(["REQ-001"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=MemberReferenceDefaultFakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            default_candidates = {
                candidate.schema_field: candidate
                for candidate in mapping_context.default_value_candidates
            }
            self.assertIn("검토 담당", default_candidates)
            self.assertTrue(default_candidates["검토 담당"].allows_custom_value)

    def test_prepare_mapping_context_includes_tracker_item_field_in_default_value_candidates(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary"])
            sheet.append(["REQ-001"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=TrackerItemDefaultValueFakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            default_candidates = {
                candidate.schema_field: candidate
                for candidate in mapping_context.default_value_candidates
            }
            self.assertIn("상위 요구사항", default_candidates)
            self.assertTrue(default_candidates["상위 요구사항"].allows_custom_value)

    def test_prepare_mapping_context_includes_multi_tracker_item_field_in_default_value_candidates(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "연관 요구사항"])
            sheet.append(["REQ-001", "REQ-100"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=TrackerItemQueryFakeClient,
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            default_candidates = {
                candidate.schema_field: candidate
                for candidate in mapping_context.default_value_candidates
            }
            self.assertIn("연관 요구사항", default_candidates)
            self.assertTrue(default_candidates["연관 요구사항"].allows_custom_value)

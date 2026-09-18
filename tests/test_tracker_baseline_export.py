from __future__ import annotations

import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from src.gui.tracker_baseline_compare import BaselineComparisonSource
from src.gui.tracker_baseline_compare import compare_tracker_items
from src.gui.tracker_baseline_export import LONG_VALUE_CELL_LINE_FEEDS
from src.gui.tracker_baseline_export import LONG_VALUE_CELL_TEXT
from src.gui.tracker_baseline_export import LONG_VALUE_SHEET_TITLE
from src.gui.tracker_baseline_export import BaselineExportError
from src.gui.tracker_baseline_export import baseline_export_fields
from src.gui.tracker_baseline_export import create_baseline_comparison_workbook
from src.gui.tracker_baseline_export import export_baseline_comparison_xlsx
from src.gui.tracker_query_models import TrackerItemSummary


def _item(
    item_id: int,
    *,
    name: str,
    status: str,
    description: str = "",
    assignees: list[dict] | None = None,
    table_rows: list[list[dict]] | None = None,
) -> TrackerItemSummary:
    return TrackerItemSummary.from_raw(
        {
            "id": item_id,
            "name": name,
            "status": {
                "id": 1 if status == "Open" else 2,
                "name": status,
                "type": "ChoiceOptionReference",
            },
            "description": description,
            "assignedTo": assignees or [],
            "customFields": [
                {
                    "fieldId": 10,
                    "name": "검증 표",
                    "type": "TableFieldValue",
                    "values": table_rows or [],
                }
            ],
        },
        tracker_id=20,
    )


def _table_row(condition: str, result: bool) -> list[dict]:
    return [
        {
            "fieldId": 501,
            "name": "조건",
            "type": "TextFieldValue",
            "value": condition,
        },
        {
            "fieldId": 502,
            "name": "결과",
            "type": "BoolFieldValue",
            "value": result,
        },
    ]


class TrackerBaselineExportTest(unittest.TestCase):
    def setUp(self) -> None:
        comparison = _item(
            1,
            name="이전 로그인 기능",
            status="Draft",
            assignees=[{"id": 7, "name": "홍길동", "type": "UserReference"}],
            table_rows=[_table_row("로그인", True), _table_row("권한 확인", False)],
        )
        reference = _item(
            1,
            name="로그인 기능",
            status="Open",
            assignees=[
                {"id": 7, "name": "홍길동", "type": "UserReference"},
                {"id": 9, "name": "김영희", "type": "UserReference"},
            ],
            table_rows=[_table_row("로그인", True), _table_row("권한 확인", True)],
        )
        added = _item(2, name="비밀번호 재설정", status="Open")
        removed = _item(3, name="기존 로그인", status="Draft")
        self.result = compare_tracker_items(
            (comparison, removed),
            (reference, added),
            before_source=BaselineComparisonSource(11),
            after_source=BaselineComparisonSource(None),
        )
        self.selected_keys = ("name", "status", "assignedTo", "custom:10")

    def test_fields_include_table_metadata_and_column_union(self) -> None:
        fields = {field.field_key: field for field in baseline_export_fields(self.result)}

        self.assertTrue(fields["custom:10"].is_table)
        self.assertEqual(
            [column.label for column in fields["custom:10"].table_columns],
            ["조건", "결과"],
        )

    def test_workbook_uses_two_header_rows_and_item_row_blocks(self) -> None:
        workbook, summary = create_baseline_comparison_workbook(
            self.result,
            tracker_name="요구사항 (ID 20)",
            reference_label="현재 상태",
            comparison_label="R1 (2026-01-01)",
            selected_field_keys=self.selected_keys,
            generated_at=datetime(2026, 8, 6, 12, 0, 0),
        )
        sheet = workbook["비교 결과"]

        self.assertEqual(workbook.sheetnames, ["요약", "비교 결과"])
        self.assertEqual(sheet["A1"].value, "결과")
        self.assertIn("A1:A2", {str(item) for item in sheet.merged_cells.ranges})
        self.assertEqual(sheet["D1"].value, "기준 · 현재 상태")
        self.assertEqual(sheet["I1"].value, "비교 · R1 (2026-01-01)")
        self.assertEqual(sheet["G2"].value, "검증 표.조건")
        self.assertEqual(sheet["H2"].value, "검증 표.결과")
        self.assertIn("A3:A4", {str(item) for item in sheet.merged_cells.ranges})
        self.assertIn("F3:F4", {str(item) for item in sheet.merged_cells.ranges})
        self.assertEqual(sheet["G3"].value, "로그인")
        self.assertEqual(sheet["H3"].value, "예")
        self.assertEqual(sheet["G4"].value, "권한 확인")
        self.assertEqual(sheet["H4"].value, "예")
        self.assertEqual(sheet["M4"].value, "아니요")
        self.assertIn("• 홍길동 (ID 7)", sheet["F3"].value)
        self.assertIn("• 김영희 (ID 9)", sheet["F3"].value)
        self.assertEqual(sheet["A5"].value, "신규")
        self.assertEqual(workbook["요약"]["A9"].value, "신규")
        self.assertEqual(workbook["요약"]["B9"].value, 1)
        self.assertEqual(sheet.freeze_panes, "D3")
        self.assertIsNone(sheet.auto_filter.ref)
        self.assertEqual(summary.item_count, 3)
        self.assertEqual(summary.data_row_count, 4)

    def test_exported_file_reopens_with_merged_ranges_and_without_formula_injection(self) -> None:
        injected = _item(4, name="=HYPERLINK('unsafe')", status="Open")
        result = compare_tracker_items(
            (),
            (injected,),
            before_source=BaselineComparisonSource(11),
            after_source=BaselineComparisonSource(None),
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "baseline.xlsx"
            export_baseline_comparison_xlsx(
                result,
                path,
                tracker_name="요구사항",
                reference_label="현재 상태",
                comparison_label="R1",
                selected_field_keys=("name",),
            )
            workbook = load_workbook(path, data_only=False)

        self.assertEqual(workbook.sheetnames, ["요약", "비교 결과"])
        self.assertEqual(workbook["비교 결과"]["A3"].value, "신규")
        self.assertEqual(workbook["비교 결과"]["D3"].value, "'=HYPERLINK('unsafe')")

    def test_export_rejects_empty_selection(self) -> None:
        with self.assertRaises(BaselineExportError):
            create_baseline_comparison_workbook(
                self.result,
                tracker_name="요구사항",
                reference_label="현재 상태",
                comparison_label="R1",
                selected_field_keys=(),
            )

    def test_long_scalar_value_moves_to_linked_sheet_without_data_loss(self) -> None:
        long_description = "현재 설명\n" + ("가" * (LONG_VALUE_CELL_TEXT + 100))
        comparison = _item(10, name="긴 설명", status="Draft", description="과거 설명")
        reference = _item(
            10,
            name="긴 설명",
            status="Open",
            description=long_description,
        )
        result = compare_tracker_items(
            (comparison,),
            (reference,),
            before_source=BaselineComparisonSource(11),
            after_source=BaselineComparisonSource(None),
        )

        workbook, summary = create_baseline_comparison_workbook(
            result,
            tracker_name="요구사항",
            reference_label="현재 상태",
            comparison_label="R1",
            selected_field_keys=("description",),
        )

        self.assertEqual(workbook.sheetnames, ["요약", "비교 결과", LONG_VALUE_SHEET_TITLE])
        comparison_sheet = workbook["비교 결과"]
        self.assertIn("전체 내용", comparison_sheet["D3"].value)
        self.assertEqual(
            comparison_sheet["D3"].hyperlink.target,
            f"#'{LONG_VALUE_SHEET_TITLE}'!H3",
        )
        long_sheet = workbook[LONG_VALUE_SHEET_TITLE]
        reconstructed = "".join(
            str(long_sheet.cell(3 + offset, 8).value or "")
            for offset in range(summary.long_value_part_count)
        )
        self.assertEqual(reconstructed, long_description)
        self.assertEqual(long_sheet["D3"].value, "설명")
        self.assertEqual(long_sheet["J3"].hyperlink.target, "#'비교 결과'!D3")
        self.assertEqual(
            workbook["요약"]["B13"].hyperlink.target,
            f"#'{LONG_VALUE_SHEET_TITLE}'!A1",
        )
        self.assertEqual(summary.long_value_count, 1)
        self.assertGreater(summary.long_value_part_count, 1)

    def test_many_line_feeds_are_split_even_below_character_limit(self) -> None:
        long_description = "\n".join(
            f"줄 {index}" for index in range(LONG_VALUE_CELL_LINE_FEEDS + 2)
        )
        result = compare_tracker_items(
            (),
            (
                _item(
                    11,
                    name="줄바꿈 설명",
                    status="Open",
                    description=long_description,
                ),
            ),
            before_source=BaselineComparisonSource(11),
            after_source=BaselineComparisonSource(None),
        )

        workbook, summary = create_baseline_comparison_workbook(
            result,
            tracker_name="요구사항",
            reference_label="현재 상태",
            comparison_label="R1",
            selected_field_keys=("description",),
        )

        long_sheet = workbook[LONG_VALUE_SHEET_TITLE]
        parts = [
            str(long_sheet.cell(3 + offset, 8).value or "")
            for offset in range(summary.long_value_part_count)
        ]
        self.assertEqual("".join(parts), long_description)
        self.assertTrue(all(part.count("\n") <= LONG_VALUE_CELL_LINE_FEEDS for part in parts))

    def test_comparison_only_long_value_links_back_to_comparison_cell(self) -> None:
        comparison_text = "과거 설명 " + ("B" * (LONG_VALUE_CELL_TEXT + 10))
        result = compare_tracker_items(
            (_item(15, name="과거 긴 값", status="Draft", description=comparison_text),),
            (_item(15, name="현재 짧은 값", status="Open", description="현재 설명"),),
            before_source=BaselineComparisonSource(11),
            after_source=BaselineComparisonSource(None),
        )

        workbook, _summary = create_baseline_comparison_workbook(
            result,
            tracker_name="요구사항",
            reference_label="현재 상태",
            comparison_label="R1",
            selected_field_keys=("description",),
        )

        self.assertEqual(
            workbook["비교 결과"]["E3"].hyperlink.target,
            f"#'{LONG_VALUE_SHEET_TITLE}'!I3",
        )
        self.assertEqual(
            workbook[LONG_VALUE_SHEET_TITLE]["J3"].hyperlink.target,
            "#'비교 결과'!E3",
        )

    def test_long_table_cell_keeps_table_row_and_column_context(self) -> None:
        long_condition = "조건 " + ("X" * (LONG_VALUE_CELL_TEXT + 200))
        reference = _item(
            12,
            name="긴 표",
            status="Open",
            table_rows=[_table_row(long_condition, True)],
        )
        comparison = _item(
            12,
            name="긴 표",
            status="Draft",
            table_rows=[_table_row("기존 조건", True)],
        )
        result = compare_tracker_items(
            (comparison,),
            (reference,),
            before_source=BaselineComparisonSource(11),
            after_source=BaselineComparisonSource(None),
        )

        workbook, summary = create_baseline_comparison_workbook(
            result,
            tracker_name="요구사항",
            reference_label="현재 상태",
            comparison_label="R1",
            selected_field_keys=("custom:10",),
        )

        comparison_sheet = workbook["비교 결과"]
        self.assertIn("전체 내용", comparison_sheet["D3"].value)
        long_sheet = workbook[LONG_VALUE_SHEET_TITLE]
        self.assertEqual(long_sheet["D3"].value, "검증 표")
        self.assertEqual(long_sheet["E3"].value, 1)
        self.assertEqual(long_sheet["F3"].value, "조건")
        reconstructed = "".join(
            str(long_sheet.cell(3 + offset, 8).value or "")
            for offset in range(summary.long_value_part_count)
        )
        self.assertEqual(reconstructed, long_condition)

    def test_long_list_value_preserves_rendered_list_content(self) -> None:
        assignees = [
            {
                "id": index,
                "name": f"담당자 {index} " + ("가" * 50),
                "type": "UserReference",
            }
            for index in range(1, 701)
        ]
        result = compare_tracker_items(
            (),
            (_item(14, name="긴 담당자 목록", status="Open", assignees=assignees),),
            before_source=BaselineComparisonSource(11),
            after_source=BaselineComparisonSource(None),
        )
        difference = next(
            field for field in result.items[0].fields if field.field_key == "assignedTo"
        )

        workbook, summary = create_baseline_comparison_workbook(
            result,
            tracker_name="요구사항",
            reference_label="현재 상태",
            comparison_label="R1",
            selected_field_keys=("assignedTo",),
        )

        long_sheet = workbook[LONG_VALUE_SHEET_TITLE]
        reconstructed = "".join(
            str(long_sheet.cell(3 + offset, 8).value or "")
            for offset in range(summary.long_value_part_count)
        )
        self.assertEqual(reconstructed, difference.reference_text())
        self.assertTrue(reconstructed.startswith("• 담당자 1"))
        self.assertIn("• 담당자 700", reconstructed)

    def test_saved_long_value_workbook_reopens_with_safe_cell_sizes(self) -> None:
        long_description = "=위험\n" + ("🙂" * (LONG_VALUE_CELL_TEXT + 100))
        result = compare_tracker_items(
            (),
            (
                _item(
                    13,
                    name="긴 이모지",
                    status="Open",
                    description=long_description,
                ),
            ),
            before_source=BaselineComparisonSource(11),
            after_source=BaselineComparisonSource(None),
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "long-baseline.xlsx"
            summary = export_baseline_comparison_xlsx(
                result,
                path,
                tracker_name="요구사항",
                reference_label="현재 상태",
                comparison_label="R1",
                selected_field_keys=("description",),
            )
            workbook = load_workbook(path, data_only=False)

        self.assertEqual(summary.long_value_count, 1)
        self.assertIn(LONG_VALUE_SHEET_TITLE, workbook.sheetnames)
        self.assertEqual(
            workbook["비교 결과"]["D3"].hyperlink.target,
            f"#'{LONG_VALUE_SHEET_TITLE}'!H3",
        )
        self.assertEqual(
            workbook[LONG_VALUE_SHEET_TITLE]["J3"].hyperlink.target,
            "#'비교 결과'!D3",
        )
        for row in workbook[LONG_VALUE_SHEET_TITLE].iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    self.assertLessEqual(len(cell.value.encode("utf-16-le")) // 2, 32767)
                    self.assertLessEqual(cell.value.count("\n"), 253)
        self.assertTrue(str(workbook[LONG_VALUE_SHEET_TITLE]["H3"].value).startswith("'="))

    def test_failed_save_keeps_existing_file_and_removes_temporary_file(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "baseline.xlsx"
            path.write_bytes(b"existing workbook")

            def fail_after_partial_write(_workbook, output_path) -> None:
                Path(output_path).write_bytes(b"partial workbook")
                raise OSError("disk write failed")

            with patch(
                "src.gui.tracker_baseline_export.Workbook.save",
                new=fail_after_partial_write,
            ), self.assertRaisesRegex(
                BaselineExportError,
                "Excel 파일을 저장하지 못했습니다",
            ):
                export_baseline_comparison_xlsx(
                    self.result,
                    path,
                    tracker_name="요구사항",
                    reference_label="현재 상태",
                    comparison_label="R1",
                    selected_field_keys=("name",),
                )

            self.assertEqual(path.read_bytes(), b"existing workbook")
            self.assertEqual([item.name for item in Path(temp_dir).iterdir()], ["baseline.xlsx"])


if __name__ == "__main__":
    unittest.main()

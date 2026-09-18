from __future__ import annotations

import json
import re
import unittest
from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parent.parent
SAMPLE_ROOTS = (
    REPO_ROOT / "data" / "gui-offline-sample",
    REPO_ROOT / "tests" / "fixtures",
    REPO_ROOT / "templates" / "codebeamer-upload-starter",
)
SENSITIVE_KEYS = {
    "access_token",
    "api_key",
    "authorization",
    "client_secret",
    "password",
    "refresh_token",
    "secret",
    "token",
}
EMAIL_PATTERN = re.compile(r"\b[A-Z0-9._%+-]+@([A-Z0-9.-]+\.[A-Z]{2,})\b", re.IGNORECASE)
ALLOWED_EMAIL_DOMAINS = {"example.com", "example.test", "invalid"}


def _walk_json_values(value, *, path: str = "$"):
    if isinstance(value, dict):
        for key, child in value.items():
            child_path = f"{path}.{key}"
            yield str(key), child, child_path
            yield from _walk_json_values(child, path=child_path)
    elif isinstance(value, list):
        for index, child in enumerate(value):
            yield from _walk_json_values(child, path=f"{path}[{index}]")


class SampleDataSecurityTest(unittest.TestCase):
    def test_sample_json_does_not_contain_credential_values(self) -> None:
        for root in SAMPLE_ROOTS:
            for path in root.rglob("*.json"):
                payload = json.loads(path.read_text(encoding="utf-8"))
                for key, value, json_path in _walk_json_values(payload):
                    normalized_key = key.strip().lower().replace("-", "_")
                    if normalized_key not in SENSITIVE_KEYS:
                        continue
                    self.assertIn(
                        value,
                        (None, ""),
                        f"{path.relative_to(REPO_ROOT)}:{json_path}",
                    )

    def test_sample_assets_do_not_contain_personal_email_domains(self) -> None:
        text_extensions = {".csv", ".json", ".jsonl", ".md", ".txt"}
        for root in SAMPLE_ROOTS:
            for path in root.rglob("*"):
                if not path.is_file() or path.suffix.lower() not in text_extensions:
                    continue
                text = path.read_text(encoding="utf-8")
                for domain in EMAIL_PATTERN.findall(text):
                    self.assertIn(
                        domain.lower(),
                        ALLOWED_EMAIL_DOMAINS,
                        str(path.relative_to(REPO_ROOT)),
                    )

            for path in root.rglob("*.xlsx"):
                workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    for sheet in workbook.worksheets:
                        for row in sheet.iter_rows(values_only=True):
                            for value in row:
                                if not isinstance(value, str):
                                    continue
                                for domain in EMAIL_PATTERN.findall(value):
                                    self.assertIn(
                                        domain.lower(),
                                        ALLOWED_EMAIL_DOMAINS,
                                        f"{path.relative_to(REPO_ROOT)}:{sheet.title}",
                                    )
                finally:
                    workbook.close()

    def test_env_example_uses_explicit_placeholders(self) -> None:
        env_values = {}
        for line in (REPO_ROOT / ".env.example").read_text(encoding="utf-8").splitlines():
            if "=" not in line:
                continue
            key, value = line.split("=", 1)
            env_values[key] = value

        self.assertEqual(env_values["CODEBEAMER_USERNAME"], "your_username")
        self.assertEqual(env_values["CODEBEAMER_PASSWORD"], "your_password")
        self.assertIn("your-codebeamer-host", env_values["CODEBEAMER_BASE_URL"])

    def test_gitignore_excludes_environment_and_private_key_files(self) -> None:
        ignore_rules = set(
            (REPO_ROOT / ".gitignore").read_text(encoding="utf-8").splitlines()
        )

        self.assertTrue({".env*", "*.key", "*.pem", "*.p12", "*.pfx"} <= ignore_rules)
        self.assertIn("!.env.example", ignore_rules)


if __name__ == "__main__":
    unittest.main()

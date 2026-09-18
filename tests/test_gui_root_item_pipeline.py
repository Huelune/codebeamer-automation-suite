from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from src.gui.root_item_service import ROOT_ITEM_MODE_GROUP_BY_COLUMN
from src.gui.root_item_service import ROOT_SOURCE_GROUP_VALUE
from src.gui.service_core import GuiExcelService
from src.gui.settings_store import GuiSettings
from src.gui.upload_service import GuiUploadPipelineService
from tests.gui_service_fixtures import FakeClient
from tests.gui_service_fixtures import FakeExcelReader
from tests.gui_service_fixtures import TrackerItemQueryFakeClient
from tests.gui_widget_cleanup import tearDownModule  # noqa: F401


class GuiRootItemPipelineServiceTest(unittest.TestCase):
    def test_build_root_item_preview_context_parses_named_groups(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary"])
            sheet.append(["REQ-001"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "file_paths": [str(path)],
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            preview_context = service.build_root_item_preview_context(
                mapping_context,
                {
                    "regex_pattern": r"^(?P<project>[A-Z]+)_(?P<title>.+)$",
                    "regex_target": "file_stem",
                    "field_sources": {"Summary": "title"},
                },
            )

            self.assertFalse(preview_context.has_blocking_issues)
            self.assertIn("project", preview_context.preview_columns)
            self.assertIn("title", preview_context.preview_columns)
            self.assertEqual(preview_context.preview_rows[0]["project"], "ABC")
            self.assertEqual(preview_context.preview_rows[0]["title"], "REQ-001")
            self.assertEqual(
                preview_context.field_assignments["Summary"],
                {
                    "enabled": True,
                    "mode": "file_source",
                    "value": "title",
                },
            )
            status_candidate = next(
                candidate
                for candidate in preview_context.field_candidates
                if candidate.schema_field == "Status"
            )
            summary_candidate = next(
                candidate
                for candidate in preview_context.field_candidates
                if candidate.schema_field == "Summary"
            )
            self.assertTrue(status_candidate.allows_fixed_value)
            self.assertEqual(status_candidate.fixed_options, ["Open", "Review"])
            self.assertTrue(summary_candidate.allows_fixed_value)
            self.assertTrue(summary_candidate.allows_custom_value)

    def test_build_root_item_preview_context_allows_fixed_value_for_multi_tracker_item_field(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary"])
            sheet.append(["REQ-001"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=TrackerItemQueryFakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "file_paths": [str(path)],
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            preview_context = service.build_root_item_preview_context(
                mapping_context,
                {"regex_pattern": "", "regex_target": "file_stem", "field_assignments": {}},
            )

            related_candidate = next(
                candidate
                for candidate in preview_context.field_candidates
                if candidate.schema_field == "연관 요구사항"
            )

            self.assertTrue(related_candidate.allows_fixed_value)
            self.assertTrue(related_candidate.allows_custom_value)

    def test_build_root_item_preview_context_groups_top_level_rows_by_column(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "Folder"])
            sheet.append(["REQ-001", "EMS"])
            sheet.append(["REQ-002", "EMS"])
            sheet.append(["REQ-003", "VCU"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "file_paths": [str(path)],
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            preview_context = service.build_root_item_preview_context(
                mapping_context,
                {
                    "root_mode": ROOT_ITEM_MODE_GROUP_BY_COLUMN,
                    "group_by_column": "Folder",
                    "regex_pattern": "",
                    "regex_target": "file_stem",
                },
            )

            self.assertFalse(preview_context.has_blocking_issues)
            self.assertEqual(preview_context.root_mode, ROOT_ITEM_MODE_GROUP_BY_COLUMN)
            self.assertEqual(preview_context.group_by_column, "Folder")
            self.assertEqual(len(preview_context.preview_rows), 2)
            self.assertEqual(
                [row[ROOT_SOURCE_GROUP_VALUE] for row in preview_context.preview_rows],
                ["EMS", "VCU"],
            )
            self.assertEqual(
                preview_context.field_assignments["Summary"],
                {
                    "enabled": True,
                    "mode": "file_source",
                    "value": ROOT_SOURCE_GROUP_VALUE,
                },
            )

    def test_build_root_item_preview_context_allows_group_mode_without_group_column(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "Folder"])
            sheet.append(["REQ-001", "EMS"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "file_paths": [str(path)],
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            preview_context = service.build_root_item_preview_context(
                mapping_context,
                {
                    "root_mode": ROOT_ITEM_MODE_GROUP_BY_COLUMN,
                    "group_by_column": "",
                    "regex_pattern": "",
                    "regex_target": "file_stem",
                },
            )

            self.assertFalse(preview_context.has_blocking_issues)
            self.assertEqual(preview_context.group_by_column, "")
            self.assertTrue(preview_context.enabled)
            self.assertFalse(preview_context.group_enabled)
            self.assertIn("파일명 파싱 결과", preview_context.status_message)

    def test_build_root_item_payload_spec_uses_regex_mapped_name(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary"])
            sheet.append(["REQ-001"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "file_paths": [str(path)],
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )
            mapping_context.root_item_config = {
                "regex_pattern": r"^(?P<project>[A-Z]+)_(?P<title>.+)$",
                "regex_target": "file_stem",
                "field_assignments": {
                    "Summary": {
                        "enabled": True,
                        "mode": "file_source",
                        "value": "title",
                    },
                    "Status": {
                        "enabled": True,
                        "mode": "fixed_value",
                        "value": "Open",
                    },
                },
            }

            root_item_name, root_field_values = service.build_root_item_payload_spec(mapping_context, str(path))

            self.assertEqual(root_item_name, "REQ-001")
            self.assertEqual(root_field_values["Summary"], "REQ-001")
            self.assertEqual(root_field_values["Status"], "Open")

    def test_build_root_item_payload_spec_accepts_custom_scalar_fixed_value(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary"])
            sheet.append(["REQ-001"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "file_paths": [str(path)],
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )
            mapping_context.root_item_config = {
                "regex_pattern": "",
                "regex_target": "file_stem",
                "field_assignments": {
                    "담당자": {
                        "enabled": True,
                        "mode": "fixed_value",
                        "value": "홍길동",
                    },
                },
            }

            root_item_name, root_field_values = service.build_root_item_payload_spec(mapping_context, str(path))

            self.assertEqual(root_item_name, "ABC_REQ-001")
            self.assertEqual(root_field_values["담당자"], "홍길동")

    def test_build_root_item_payload_specs_groups_rows_by_column_within_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "Folder"])
            sheet.append(["REQ-001", "EMS"])
            sheet.append(["REQ-002", "EMS"])
            sheet.append(["REQ-003", "VCU"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            file_state = {
                "file_path": str(path),
                "file_paths": [str(path)],
                "preview_file_path": str(path),
                "sheet_name": "Main",
                "header_row": 1,
                "summary_column": "Summary",
            }
            mapping_context = service.prepare_mapping_context(settings, file_state)
            mapping_context.root_item_config = {
                "root_mode": ROOT_ITEM_MODE_GROUP_BY_COLUMN,
                "group_by_column": "Folder",
                "regex_pattern": "",
                "regex_target": "file_stem",
                "field_assignments": {
                    "Summary": {
                        "enabled": True,
                        "mode": "file_source",
                        "value": ROOT_SOURCE_GROUP_VALUE,
                    },
                    "Status": {
                        "enabled": True,
                        "mode": "fixed_value",
                        "value": "Open",
                    },
                },
            }

            wizard = service._prepare_wizard_for_file(
                settings,
                mapping_context,
                file_path=str(path),
                sheet_name="Main",
                header_row=1,
                summary_col="Summary",
            )
            root_item_specs = service.build_root_item_payload_specs(mapping_context, wizard, str(path))

            self.assertEqual(
                [(spec.name, spec.row_ids) for spec in root_item_specs],
                [("EMS", [0, 1]), ("VCU", [2])],
            )
            self.assertEqual(root_item_specs[0].field_values["Summary"], "EMS")
            self.assertEqual(root_item_specs[0].field_values["Status"], "Open")
            self.assertEqual(root_item_specs[1].field_values["Summary"], "VCU")

    def test_build_root_item_payload_specs_can_create_file_root_and_group_folders_together(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "Folder"])
            sheet.append(["REQ-001", "EMS"])
            sheet.append(["REQ-002", "EMS"])
            sheet.append(["REQ-003", "VCU"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            file_state = {
                "file_path": str(path),
                "file_paths": [str(path)],
                "preview_file_path": str(path),
                "sheet_name": "Main",
                "header_row": 1,
                "summary_column": "Summary",
            }
            mapping_context = service.prepare_mapping_context(settings, file_state)
            mapping_context.root_item_config = {
                "enabled": True,
                "group_enabled": True,
                "group_by_column": "Folder",
                "regex_pattern": "",
                "regex_target": "file_stem",
                "field_assignments": {
                    "Summary": {
                        "enabled": True,
                        "mode": "file_source",
                        "value": ROOT_SOURCE_GROUP_VALUE,
                    },
                    "Status": {
                        "enabled": True,
                        "mode": "fixed_value",
                        "value": "Open",
                    },
                },
            }

            wizard = service._prepare_wizard_for_file(
                settings,
                mapping_context,
                file_path=str(path),
                sheet_name="Main",
                header_row=1,
                summary_col="Summary",
            )
            root_item_specs = service.build_root_item_payload_specs(mapping_context, wizard, str(path))

            self.assertEqual(
                [
                    (spec.name, spec.kind, spec.parent_key, spec.row_ids)
                    for spec in root_item_specs
                ],
                [
                    ("ABC_REQ-001", "file_root", None, []),
                    ("EMS", "group_root", str(path), [0, 1]),
                    ("VCU", "group_root", str(path), [2]),
                ],
            )
            self.assertEqual(root_item_specs[0].field_values["Status"], "Open")
            self.assertEqual(root_item_specs[1].field_values["Status"], "Open")
            self.assertNotIn("Summary", root_item_specs[1].field_values)
            self.assertNotIn("Summary", root_item_specs[2].field_values)

    def test_build_root_item_payload_specs_can_customize_file_root_name_when_group_folders_are_enabled(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "Folder"])
            sheet.append(["REQ-001", "EMS"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            file_state = {
                "file_path": str(path),
                "file_paths": [str(path)],
                "preview_file_path": str(path),
                "sheet_name": "Main",
                "header_row": 1,
                "summary_column": "Summary",
            }
            mapping_context = service.prepare_mapping_context(settings, file_state)
            mapping_context.root_item_config = {
                "enabled": True,
                "group_enabled": True,
                "group_by_column": "Folder",
                "regex_pattern": r"^(?P<project>[A-Z]+)_(?P<title>.+)$",
                "regex_target": "file_stem",
                "field_assignments": {
                    "Summary": {
                        "enabled": True,
                        "mode": "file_source",
                        "value": "title",
                    },
                },
            }

            wizard = service._prepare_wizard_for_file(
                settings,
                mapping_context,
                file_path=str(path),
                sheet_name="Main",
                header_row=1,
                summary_col="Summary",
            )
            root_item_specs = service.build_root_item_payload_specs(mapping_context, wizard, str(path))

            self.assertEqual(root_item_specs[0].kind, "file_root")
            self.assertEqual(root_item_specs[0].name, "REQ-001")
            self.assertEqual(root_item_specs[0].field_values["Summary"], "REQ-001")
            self.assertEqual(root_item_specs[1].name, "EMS")

    def test_build_root_item_payload_specs_falls_back_to_file_mode_when_group_column_is_blank(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "Folder"])
            sheet.append(["REQ-001", "EMS"])
            sheet.append(["REQ-002", "VCU"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            file_state = {
                "file_path": str(path),
                "file_paths": [str(path)],
                "preview_file_path": str(path),
                "sheet_name": "Main",
                "header_row": 1,
                "summary_column": "Summary",
            }
            mapping_context = service.prepare_mapping_context(settings, file_state)
            mapping_context.root_item_config = {
                "root_mode": ROOT_ITEM_MODE_GROUP_BY_COLUMN,
                "group_by_column": "",
                "regex_pattern": "",
                "regex_target": "file_stem",
            }

            wizard = service._prepare_wizard_for_file(
                settings,
                mapping_context,
                file_path=str(path),
                sheet_name="Main",
                header_row=1,
                summary_col="Summary",
            )
            root_item_specs = service.build_root_item_payload_specs(mapping_context, wizard, str(path))

            self.assertEqual(len(root_item_specs), 1)
            self.assertEqual(root_item_specs[0].name, "ABC_REQ-001")
            self.assertEqual(root_item_specs[0].row_ids, [0, 1])

    def test_build_root_item_preview_context_does_not_block_when_root_item_is_disabled(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary"])
            sheet.append(["REQ-001"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "file_paths": [str(path)],
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            preview_context = service.build_root_item_preview_context(
                mapping_context,
                {
                    "enabled": False,
                    "regex_pattern": "(",
                    "field_assignments": {
                        "Summary": {
                            "enabled": True,
                            "mode": "file_source",
                            "value": "missing",
                        }
                    },
                },
            )

            self.assertFalse(preview_context.enabled)
            self.assertFalse(preview_context.has_blocking_issues)
            self.assertEqual(preview_context.preview_columns, ["file_name"])
            self.assertIn("무시", preview_context.status_message)

    def test_build_root_item_payload_spec_returns_none_when_root_item_is_disabled(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary"])
            sheet.append(["REQ-001"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "file_paths": [str(path)],
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )
            mapping_context.root_item_config = {
                "enabled": False,
                "regex_pattern": r"^(?P<title>.+)$",
                "field_assignments": {
                    "Summary": {
                        "enabled": True,
                        "mode": "file_source",
                        "value": "title",
                    },
                },
            }

            root_item_name, root_field_values = service.build_root_item_payload_spec(mapping_context, str(path))

            self.assertIsNone(root_item_name)
            self.assertEqual(root_field_values, {})


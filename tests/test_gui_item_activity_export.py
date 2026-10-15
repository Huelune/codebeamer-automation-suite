from __future__ import annotations

import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from src.gui.excel_long_value import LONG_VALUE_SHEET_TITLE
from src.gui.item_activity import BaselineMarker
from src.gui.item_activity import ItemActivityEntry
from src.gui.item_activity import ItemActivityKind
from src.gui.item_activity import summarize
from src.gui.item_activity_export import HISTORY_SHEET_TITLE
from src.gui.item_activity_export import ItemActivityExportError
from src.gui.item_activity_export import ItemActivityExportRequest
from src.gui.item_activity_export import create_item_activity_workbook
from src.gui.item_activity_export import export_item_activity_xlsx
from src.gui.tracker_item_context_models import ItemFieldChange


FIELD = ItemActivityKind.FIELD_CHANGE
GENERATED = datetime(2026, 10, 13, 9, 0, 0)


def _request(*entries: ItemActivityEntry, baseline: BaselineMarker | None = None) -> ItemActivityExportRequest:
    return ItemActivityExportRequest(
        item_id=9001001,
        item_name="Sample requirement",
        tracker_name="Offline Requirements (ID 24680001)",
        scope_label="Offline Requirements Baseline 이후",
        kind_labels=("필드 변경", "댓글", "첨부"),
        entries=entries,
        summary=summarize(entries),
        baseline=baseline,
    )


ENTRIES = (
    ItemActivityEntry(
        FIELD,
        at="2026-07-20T09:00:00Z",
        by="Sample User A",
        version=3,
        changes=(ItemFieldChange("Risk Level", None, "Medium", 101), ItemFieldChange("Summary", "Old", "New", 3)),
    ),
    ItemActivityEntry(
        ItemActivityKind.COMMENT,
        at="2026-07-19T09:10:00Z",
        by="Sample User B",
        title="댓글 #51001",
        body="Review note",
        attachment_names=("anonymous-evidence.txt",),
    ),
    ItemActivityEntry(
        ItemActivityKind.ATTACHMENT, at="2026-07-18T11:00:00Z", title="anonymous-spec.txt", body="2,048 bytes"
    ),
    ItemActivityEntry(FIELD, at="2026-07-17T00:00:00Z", by="Sample User A", version=2, change_summary="Moved"),
)


def _rows(sheet) -> list[tuple]:
    return [tuple(cell.value for cell in row) for row in sheet.iter_rows()]


class ItemActivityExportTest(unittest.TestCase):
    def test_summary_field_and_history_sheets(self) -> None:
        workbook, summary = create_item_activity_workbook(_request(*ENTRIES), generated_at=GENERATED)

        self.assertEqual(workbook.sheetnames, ["요약", "필드별 요약", HISTORY_SHEET_TITLE])
        self.assertEqual((summary.entry_count, summary.row_count, summary.long_value_count), (4, 5, 0))
        overview = dict(_rows(workbook["요약"]))
        self.assertEqual(overview["아이템"], "#9001001 Sample requirement")
        self.assertEqual(overview["구간"], "Offline Requirements Baseline 이후")
        self.assertEqual(overview["유형"], "필드 변경, 댓글, 첨부")
        self.assertEqual(overview["생성 시각"], "2026-10-13T09:00:00")
        self.assertEqual((overview["버전"], overview["필드 변경"], overview["바뀐 필드"]), (2, 2, 2))
        self.assertEqual((overview["필드 변경 정보 없는 버전"], overview["댓글"], overview["첨부"]), (1, 1, 1))
        self.assertEqual(overview["긴 값 분할"], "없음")

        fields = _rows(workbook["필드별 요약"])
        self.assertEqual(
            fields[0], ("필드", "구간 처음 값", "구간 마지막 값", "변경 횟수", "마지막 수정 시각", "마지막 수정자")
        )
        self.assertIn(("Risk Level", "-", "Medium", 1, "2026-07-20T09:00:00Z", "Sample User A"), fields[1:])

        history = workbook[HISTORY_SHEET_TITLE]
        rows = _rows(history)
        self.assertEqual(rows[0], ("순번", "유형", "버전", "시각", "사용자", "대상", "이전 값", "새 값", "구간"))
        self.assertEqual(
            rows[1:],
            [
                (1, "필드 변경", 3, "2026-07-20T09:00:00Z", "Sample User A", "Risk Level", "-", "Medium", None),
                (1, "필드 변경", 3, "2026-07-20T09:00:00Z", "Sample User A", "Summary", "Old", "New", None),
                (2, "댓글", None, "2026-07-19T09:10:00Z", "Sample User B", "댓글 #51001", None,
                 "Review note\n\n첨부: anonymous-evidence.txt", None),
                (3, "첨부", None, "2026-07-18T11:00:00Z", None, "anonymous-spec.txt", None, "2,048 bytes", None),
                (4, "필드 변경", 2, "2026-07-17T00:00:00Z", "Sample User A", None, None, "Moved", None),
            ],
        )
        self.assertEqual(history.auto_filter.ref, "A1:I6")
        self.assertEqual(history.freeze_panes, "A2")

    def test_baseline_marker_fills_the_scope_column(self) -> None:
        marker = BaselineMarker("Baseline #7", version=2)
        workbook, _summary = create_item_activity_workbook(_request(*ENTRIES, baseline=marker), generated_at=GENERATED)

        scope_column = [row[8] for row in _rows(workbook[HISTORY_SHEET_TITLE])[1:]]
        # 필드 변경은 버전으로, 생성 시각을 모르는 Baseline이면 댓글·첨부는 표시하지 않는다.
        self.assertEqual(scope_column, ["Baseline 이후", "Baseline 이후", None, None, None])

    def test_long_values_move_to_the_long_value_sheet_with_links_both_ways(self) -> None:
        long_text = "가" * 40000
        entry = ItemActivityEntry(
            FIELD, at="2026-07-20T09:00:00Z", by="Sample User A", version=3,
            changes=(ItemFieldChange("Description", "짧은 설명", long_text, 80),),
        )

        workbook, summary = create_item_activity_workbook(_request(entry), generated_at=GENERATED)

        self.assertEqual(summary.long_value_count, 1)
        self.assertEqual(workbook.sheetnames[-1], LONG_VALUE_SHEET_TITLE)
        history_cell = workbook[HISTORY_SHEET_TITLE]["H2"]
        self.assertIn(LONG_VALUE_SHEET_TITLE, history_cell.value)
        self.assertEqual(history_cell.hyperlink.target, f"#'{LONG_VALUE_SHEET_TITLE}'!F2")
        long_rows = _rows(workbook[LONG_VALUE_SHEET_TITLE])
        self.assertEqual(long_rows[0], ("순번", "유형", "대상", "칸", "부분", "전체 내용", "이동"))
        self.assertEqual([row[4] for row in long_rows[1:]], ["1/2", "2/2"])
        self.assertEqual("".join(row[5] for row in long_rows[1:]), long_text)
        self.assertEqual(workbook[LONG_VALUE_SHEET_TITLE]["G2"].hyperlink.target, f"#'{HISTORY_SHEET_TITLE}'!H2")
        overview = dict(_rows(workbook["요약"]))
        self.assertEqual(overview["긴 값 분할"], "1개 값")

    def test_formula_like_values_are_written_as_text(self) -> None:
        entry = ItemActivityEntry(
            FIELD, at="2026-07-20T09:00:00Z", by="Sample User A", version=3,
            changes=(ItemFieldChange("=Field", "+1", "=HYPERLINK(\"x\")", 5), ItemFieldChange("Note", "-", "@x", 6)),
        )

        workbook, _summary = create_item_activity_workbook(_request(entry), generated_at=GENERATED)

        rows = _rows(workbook[HISTORY_SHEET_TITLE])
        self.assertEqual(rows[1][5:8], ("'=Field", "'+1", "'=HYPERLINK(\"x\")"))
        # 빈 값 표시 "-"는 그대로 둔다.
        self.assertEqual(rows[2][5:8], ("Note", "-", "'@x"))
        for row in workbook[HISTORY_SHEET_TITLE].iter_rows(min_row=2):
            for cell in row:
                self.assertNotEqual(cell.data_type, "f")

    def test_empty_request_is_rejected(self) -> None:
        with self.assertRaises(ItemActivityExportError):
            create_item_activity_workbook(_request())

    def test_export_saves_xlsx_without_leaving_temporary_files(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "item_history"

            summary = export_item_activity_xlsx(_request(*ENTRIES), target)

            saved = Path(summary.output_path)
            self.assertEqual(saved.suffix, ".xlsx")
            self.assertTrue(saved.is_file())
            self.assertEqual(sorted(path.name for path in Path(directory).iterdir()), ["item_history.xlsx"])
            self.assertEqual(load_workbook(saved).sheetnames, ["요약", "필드별 요약", HISTORY_SHEET_TITLE])


if __name__ == "__main__":
    unittest.main()

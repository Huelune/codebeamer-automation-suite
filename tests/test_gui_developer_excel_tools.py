from __future__ import annotations

import csv
import sys
import tempfile
import unittest
from datetime import date
from datetime import datetime
from datetime import time
from datetime import timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl import load_workbook

from src.gui.developer_excel_tools import DeveloperExcelToolError
from src.gui.developer_excel_tools import DeveloperExcelToolService
from tests.gui_widget_cleanup import tearDownModule  # noqa: F401


class DeveloperExcelToolServiceTest(unittest.TestCase):
    def setUp(self) -> None:
        self.service = DeveloperExcelToolService()

    def _workbook(self, root: str, name: str = "source.xlsx") -> Path:
        path = Path(root) / name
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Main"
        sheet.append(["Summary", " Summary ", "Table.Step"])
        sheet.append(["REQ-001", "first", "step 1"])
        sheet.merge_cells("D2:E2")
        sheet.row_dimensions[2].hidden = True
        sheet.column_dimensions["C"].hidden = True
        hidden = workbook.create_sheet("Hidden")
        hidden.sheet_state = "hidden"
        workbook.save(path)
        workbook.close()
        return path

    def test_inspect_reports_headers_table_shape_and_hidden_content(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = self._workbook(temp_dir)

            report = self.service.inspect(path, sheet_name="Main")

            codes = {issue.code for issue in report.issues}
            self.assertEqual(report.selected_sheet, "Main")
            self.assertEqual(report.hidden_sheet_count, 1)
            self.assertEqual(report.hidden_row_count, 1)
            self.assertEqual(report.hidden_column_count, 1)
            self.assertEqual(report.merged_range_count, 1)
            self.assertIn("TABLE_HEADERS", codes)
            self.assertIn("MERGED_CELLS", codes)
            self.assertIn("HIDDEN_CONTENT", codes)

    def test_inspect_reports_duplicate_empty_and_missing_summary_headers(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "headers.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Name", "name", None])
            sheet.append(["A", "B", "C"])
            workbook.save(path)
            workbook.close()

            report = self.service.inspect(path, summary_column="Summary")

            codes = {issue.code for issue in report.issues}
            self.assertIn("DUPLICATE_HEADERS", codes)
            self.assertIn("EMPTY_HEADERS", codes)
            self.assertIn("SUMMARY_MISSING", codes)

    def test_csv_inspection_supports_cp949_and_custom_delimiter(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "source.csv"
            path.write_text("Summary;설명\nREQ-001;한글\n", encoding="cp949")

            report = self.service.inspect(
                path,
                csv_encoding="cp949",
                csv_delimiter=";",
            )

            self.assertEqual(report.headers, ("Summary", "설명"))
            self.assertEqual(report.row_count, 1)

    def test_value_xlsx_conversion_does_not_change_source_and_reopens(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.csv"
            source.write_text(
                " Summary ,Description\nREQ-001,=unsafe\n",
                encoding="utf-8-sig",
            )
            before = source.read_bytes()
            output = Path(temp_dir) / "converted.xlsx"

            result = self.service.convert_to_value_xlsx(
                source,
                output,
                normalize_headers=True,
                summary_column="Summary",
                renamed_summary="요약",
            )

            self.assertEqual(source.read_bytes(), before)
            self.assertEqual(result.row_count, 1)
            workbook = load_workbook(output, data_only=False)
            try:
                sheet = workbook.active
                self.assertEqual(sheet["A1"].value, "요약")
                self.assertEqual(sheet["B2"].value, "'=unsafe")
                self.assertEqual(sheet.freeze_panes, "A2")
            finally:
                workbook.close()

    def test_value_xlsx_conversion_rejects_original_output_path(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            workbook = Workbook()
            workbook.active.append(["Summary"])
            workbook.active.append(["REQ-001"])
            workbook.save(source)
            workbook.close()

            with self.assertRaisesRegex(DeveloperExcelToolError, "원본 파일"):
                self.service.convert_to_value_xlsx(source, source)

    def test_value_xlsx_conversion_warns_that_xlsm_macros_are_not_preserved(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsm"
            workbook = Workbook()
            workbook.active.append(["Summary"])
            workbook.active.append(["REQ-001"])
            workbook.save(source)
            workbook.close()

            result = self.service.convert_to_value_xlsx(
                source,
                Path(temp_dir) / "converted.xlsx",
            )

            self.assertIn("매크로", " ".join(result.warnings))

    def test_inspection_rejects_template_extensions(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xltx"
            source.write_bytes(b"not-a-template")

            with self.assertRaisesRegex(DeveloperExcelToolError, "지원하지 않는"):
                self.service.inspect(source)

    def test_value_xlsx_conversion_rejects_long_csv_cell_without_truncation(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.csv"
            with source.open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(["Summary", "Description"])
                writer.writerow(["REQ-001", "A" * 32_768])

            with self.assertRaisesRegex(DeveloperExcelToolError, "길이 제한"):
                self.service.convert_to_value_xlsx(
                    source,
                    Path(temp_dir) / "converted.xlsx",
                )

    def test_value_xlsx_conversion_rejects_uncalculated_formula(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "formula.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Summary", "Value"])
            sheet.append(["REQ-001", "=1+1"])
            workbook.save(source)
            workbook.close()

            with self.assertRaisesRegex(DeveloperExcelToolError, "계산된 값"):
                self.service.convert_to_value_xlsx(
                    source,
                    Path(temp_dir) / "converted.xlsx",
                )

    def test_csv_conversion_protects_formula_like_values(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Summary", "Tab", "CR", "LF", "Space", "Safe"])
            sheet.append(
                [
                    "REQ-001",
                    "\t=unsafe",
                    "\r+unsafe",
                    "\n-unsafe",
                    " @unsafe",
                    "  safe",
                ]
            )
            workbook.save(source)
            workbook.close()
            output = Path(temp_dir) / "converted.csv"

            self.service.convert_to_csv(source, output)

            with output.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle))
            self.assertEqual(rows[1][0], "REQ-001")
            # csv.reader may normalize an embedded standalone CR to LF on
            # Windows. The security contract is that the apostrophe precedes
            # the original leading whitespace and formula prefix.
            for value, formula_prefix in zip(rows[1][1:5], "=+-@", strict=True):
                with self.subTest(value=value):
                    self.assertTrue(value.startswith("'"))
                    self.assertIn(value[1], " \t\r\n")
                    self.assertTrue(
                        value[1:].lstrip(" \t\r\n").startswith(formula_prefix)
                    )
            self.assertEqual(rows[1][5], "  safe")

    def test_value_xlsx_conversion_preserves_excel_scalar_types(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Summary", "Date", "DateTime", "Time", "Duration"])
            sheet.append(
                [
                    "REQ-001",
                    date(2026, 8, 9),
                    datetime(2026, 8, 9, 10, 11, 12),
                    time(13, 14, 15),
                    timedelta(hours=27, minutes=30),
                ]
            )
            workbook.save(source)
            workbook.close()
            output = Path(temp_dir) / "converted.xlsx"

            self.service.convert_to_value_xlsx(source, output)

            converted = load_workbook(output, data_only=True)
            try:
                values = [converted.active.cell(2, column).value for column in range(2, 6)]
                self.assertIsInstance(values[0], (date, datetime))
                self.assertIsInstance(values[1], datetime)
                self.assertIsInstance(values[2], time)
                self.assertIsInstance(values[3], timedelta)
            finally:
                converted.close()

    def test_xls_inspection_and_conversion_preserve_raw_matrix_and_header_row(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xls"
            source.write_bytes(b"fake-xls")
            source_before = source.read_bytes()
            output = Path(temp_dir) / "converted.xlsx"
            matrix = [
                ["설명", None, None],
                [None, "Name", "When"],
                [None, None, None],
                [None, "Item", datetime(2026, 8, 9, 10, 0, 0)],
            ]
            fake_xlwings, apps = self._fake_xlwings(matrix)

            with patch.dict(sys.modules, {"xlwings": fake_xlwings}):
                report = self.service.inspect(
                    source,
                    sheet_name="Raw",
                    header_row=2,
                    summary_column="Name",
                )
                result = self.service.convert_to_value_xlsx(
                    source,
                    output,
                    sheet_name="Raw",
                    header_row=2,
                )

            self.assertEqual(report.headers, ("Unnamed_0", "Name", "When"))
            self.assertEqual(report.row_count, 2)
            self.assertEqual(report.column_count, 3)
            self.assertEqual(result.row_count, 2)
            self.assertEqual(source.read_bytes(), source_before)
            self.assertTrue(all(app.quit_called for app in apps))
            converted = load_workbook(output, data_only=True)
            try:
                sheet = converted["Raw"]
                self.assertEqual(sheet.max_row, 4)
                self.assertEqual(sheet.max_column, 3)
                self.assertEqual(sheet["A1"].value, "설명")
                self.assertEqual(sheet["B2"].value, "Name")
                self.assertIsNone(sheet["B3"].value)
                self.assertEqual(sheet["B4"].value, "Item")
                self.assertIsInstance(sheet["C4"].value, datetime)
                self.assertEqual(sheet.freeze_panes, "A3")
            finally:
                converted.close()

    def test_xls_reports_missing_xlwings_explicitly(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xls"
            source.write_bytes(b"fake-xls")

            with patch.dict(sys.modules, {"xlwings": None}), self.assertRaisesRegex(
                DeveloperExcelToolError,
                "Microsoft Excel과 xlwings",
            ):
                self.service.inspect(source)

    def test_xls_reports_unavailable_excel_explicitly(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xls"
            source.write_bytes(b"fake-xls")

            def fail_to_start(**_kwargs):
                raise RuntimeError("Excel is unavailable")

            fake_xlwings = SimpleNamespace(App=fail_to_start)
            with patch.dict(sys.modules, {"xlwings": fake_xlwings}), self.assertRaisesRegex(
                DeveloperExcelToolError,
                "Microsoft Excel을 시작할 수 없습니다",
            ):
                self.service.inspect(source)

    @staticmethod
    def _fake_xlwings(matrix):
        class FakeRange:
            def __init__(self, values):
                self.value = values

            def options(self, **_kwargs):
                return self

        class FakeSheet:
            name = "Raw"

            def __init__(self, values):
                self._values = values
                self.used_range = SimpleNamespace(
                    last_cell=SimpleNamespace(
                        row=len(values),
                        column=max((len(row) for row in values), default=0),
                    )
                )

            def range(self, _start, _end):
                return FakeRange(self._values)

        class FakeSheets(list):
            def __getitem__(self, key):
                if isinstance(key, str):
                    for sheet in self:
                        if sheet.name == key:
                            return sheet
                    raise KeyError(key)
                return super().__getitem__(key)

        class FakeWorkbook:
            def __init__(self, values):
                self.sheets = FakeSheets([FakeSheet(values)])

            def close(self):
                return None

        class FakeBooks:
            def __init__(self, values):
                self._values = values

            def open(self, _path, **_kwargs):
                return FakeWorkbook(self._values)

        class FakeApp:
            def __init__(self, values):
                self.books = FakeBooks(values)
                self.display_alerts = True
                self.screen_updating = True
                self.quit_called = False

            def quit(self):
                self.quit_called = True

        apps = []

        def create_app(**_kwargs):
            app = FakeApp(matrix)
            apps.append(app)
            return app

        return SimpleNamespace(App=create_app), apps

    def test_csv_conversion_respects_input_encoding_and_delimiter(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.csv"
            source.write_text("Summary;설명\nREQ-001;한글\n", encoding="cp949")
            output = Path(temp_dir) / "converted.csv"

            self.service.convert_to_csv(
                source,
                output,
                input_encoding="cp949",
                input_delimiter=";",
                output_encoding="utf-8-sig",
                delimiter=",",
            )

            with output.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle))
            self.assertEqual(rows, [["Summary", "설명"], ["REQ-001", "한글"]])


if __name__ == "__main__":
    unittest.main()

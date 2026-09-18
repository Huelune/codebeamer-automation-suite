from __future__ import annotations

import pandas as pd

from src.gui.service_core import GuiExcelService


class FakeExcelReader:
    def __init__(self, header_row: int = 1, summary_col: str = "Summary", logger=None) -> None:
        del logger
        self.header_row = header_row
        self.summary_col = summary_col

    def list_sheet_names(self, file_path: str) -> list[str]:
        from openpyxl import load_workbook
        wb = load_workbook(file_path, read_only=True, data_only=True)
        try:
            return wb.sheetnames
        finally:
            wb.close()

    def read_headers(self, file_path: str, sheet_name: str | int) -> list[str]:
        from openpyxl import load_workbook
        wb = load_workbook(file_path, read_only=True, data_only=True)
        try:
            ws = wb[sheet_name]
            row = next(ws.iter_rows(min_row=self.header_row, max_row=self.header_row, values_only=True), ())
            return [str(value).strip() if value is not None else f"Unnamed_{index}" for index, value in enumerate(row)]
        finally:
            wb.close()

    def read_preview_rows(
        self,
        file_path: str,
        sheet_name: str | int,
        *,
        max_rows: int = 10,
    ) -> tuple[list[str], list[list[object]]]:
        from openpyxl import load_workbook
        wb = load_workbook(file_path, read_only=True, data_only=True)
        try:
            ws = wb[sheet_name]
            headers = self.read_headers(file_path, sheet_name)
            rows: list[list[object]] = []
            if max(int(max_rows), 0) == 0:
                return headers, rows
            for values in ws.iter_rows(
                min_row=self.header_row + 1,
                values_only=True,
            ):
                normalized = list(values[: len(headers)])
                normalized += [None] * (len(headers) - len(normalized))
                if all(value is None or str(value).strip() == "" for value in normalized):
                    continue
                rows.append(normalized)
                if len(rows) >= max(int(max_rows), 0):
                    break
            return headers, rows
        finally:
            wb.close()

    def read_excel(self, file_path: str, sheet_name: str | int = 0, visible: bool = False) -> pd.DataFrame:
        del visible
        from openpyxl import load_workbook
        wb = load_workbook(file_path, read_only=True, data_only=True)
        try:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            headers = self.read_headers(file_path, sheet_name)
            records = []
            for excel_row, values in enumerate(rows[self.header_row:], start=self.header_row + 1):
                normalized = list(values)
                if all(value is None or str(value).strip() == "" for value in normalized[:len(headers)]):
                    continue
                record = {
                    header: normalized[index] if index < len(normalized) else None
                    for index, header in enumerate(headers)
                }
                record["_excel_row"] = excel_row
                record["_summary_indent"] = 0
                records.append(record)
            return pd.DataFrame(records, dtype=object)
        finally:
            wb.close()


class FakeClient:
    def __init__(self, base_url, username, password, logger=None, **kwargs) -> None:
        del logger, kwargs
        self.base_url = base_url
        self.username = username
        self.password = password

    def get_projects(self):
        return [
            {"id": 10, "name": "Project A"},
            {"id": 20, "name": "Project B"},
        ]

    def get_trackers(self, project_id: int):
        return [
            {"id": project_id * 100, "name": f"Tracker {project_id}-1"},
            {"id": project_id * 100 + 1, "name": f"Tracker {project_id}-2"},
        ]

    def get_tracker_schema(self, tracker_id: int):
        del tracker_id
        return {
            "fields": [
                {
                    "id": 1,
                    "name": "Summary",
                    "type": "TextField",
                    "trackerItemField": "name",
                    "valueModel": "TextFieldValue",
                    "multipleValues": False,
                },
                {
                    "id": 2,
                    "name": "Status",
                    "type": "OptionChoiceField",
                    "trackerItemField": "status",
                    "valueModel": "ChoiceFieldValue<ChoiceOptionReference>",
                    "multipleValues": False,
                    "options": [
                        {"id": 201, "name": "Open", "type": "ChoiceOptionReference"},
                        {"id": 202, "name": "Review", "type": "ChoiceOptionReference"},
                    ],
                },
                {
                    "id": 3,
                    "name": "담당자",
                    "type": "TextField",
                    "trackerItemField": None,
                    "valueModel": "TextFieldValue",
                    "multipleValues": False,
                },
                {
                    "id": 4,
                    "name": "id",
                    "type": "TextField",
                    "trackerItemField": "id",
                    "valueModel": "TextFieldValue",
                    "multipleValues": False,
                },
                {
                    "id": 5,
                    "name": "parent",
                    "type": "TextField",
                    "trackerItemField": "parent",
                    "valueModel": "TextFieldValue",
                    "multipleValues": False,
                },
                {
                    "id": 6,
                    "name": "테이블필드",
                    "type": "TableField",
                    "trackerItemField": None,
                    "valueModel": "TableFieldValue",
                    "multipleValues": False,
                    "columns": [
                        {
                            "id": 31,
                            "name": "컬럼A",
                            "type": "TextField",
                            "valueModel": "TextFieldValue",
                        }
                    ],
                },
            ]
        }

    def create_item(self, tracker_id: int, payload: dict, parent_item_id: int | None = None):
        del tracker_id, payload, parent_item_id
        return {"id": 1}


class UpdateModeFakeClient(FakeClient):
    all_update_calls: list[tuple[int, dict]] = []
    all_get_item_calls: list[int] = []

    @classmethod
    def reset_calls(cls) -> None:
        cls.all_update_calls = []
        cls.all_get_item_calls = []

    def get_item(self, item_id: int):
        self.__class__.all_get_item_calls.append(int(item_id))
        return {
            "id": int(item_id),
            "name": f"기존-{item_id}",
            "description": "기존 설명",
            "status": {"id": 201, "name": "Open", "type": "ChoiceOptionReference"},
            "customFields": [
                {
                    "fieldId": 3,
                    "name": "담당자",
                    "type": "TextFieldValue",
                    "value": "기존 담당자",
                },
                {
                    "fieldId": 999,
                    "name": "기타",
                    "type": "TextFieldValue",
                    "value": "보존",
                },
            ],
        }

    def update_item(self, item_id: int, payload: dict):
        self.__class__.all_update_calls.append((int(item_id), dict(payload)))
        return {"id": int(item_id)}


class UpsertModeFakeClient(UpdateModeFakeClient):
    all_create_item_calls: list[dict[str, object]] = []

    @classmethod
    def reset_calls(cls) -> None:
        super().reset_calls()
        cls.all_create_item_calls = []

    def create_item(self, tracker_id: int, payload: dict, parent_item_id: int | None = None):
        self.__class__.all_create_item_calls.append({
            "tracker_id": int(tracker_id),
            "payload": dict(payload),
            "parent_item_id": parent_item_id,
        })
        return {"id": 1000 + len(self.__class__.all_create_item_calls)}


class TrackerItemQueryFakeClient(FakeClient):
    all_search_calls: list[tuple[int, str]] = []

    def __init__(self, base_url, username, password, logger=None, **kwargs) -> None:
        super().__init__(base_url, username, password, logger=logger, **kwargs)
        self.search_calls: list[tuple[int, str]] = []

    def get_tracker_schema(self, tracker_id: int):
        schema = super().get_tracker_schema(tracker_id)
        schema["fields"].append({
            "id": 7,
            "name": "연관 요구사항",
            "type": "TrackerItemChoiceField",
            "valueModel": "ChoiceFieldValue<TrackerItemReference>",
            "multipleValues": True,
        })
        return schema

    def get_tracker_configuration(self, tracker_id: int):
        del tracker_id
        return {
            "basicInformation": {
                "trackerId": 1000,
                "name": "Test Tracker",
            },
            "fields": [
                {
                    "label": "연관 요구사항",
                    "choiceOptionSetting": {
                        "referenceFilters": [
                            {
                                "domainType": "TRACKER",
                                "domainId": 13526611,
                            }
                        ]
                    },
                }
            ],
        }

    def search_tracker_items_by_name(self, *, tracker_id: int, name: str, **kwargs):
        del kwargs
        self.search_calls.append((tracker_id, name))
        self.__class__.all_search_calls.append((tracker_id, name))
        lookup = {
            "REQ-100": [{"id": 9001, "name": "REQ-100", "type": "TrackerItemReference"}],
            "REQ-200": [{"id": 9002, "name": "REQ-200", "type": "TrackerItemReference"}],
        }
        return lookup.get(name, [])


class TrackerItemNonTrackerConfigFakeClient(TrackerItemQueryFakeClient):
    def get_tracker_configuration(self, tracker_id: int):
        del tracker_id
        return {
            "basicInformation": {
                "trackerId": 1000,
                "name": "Test Tracker",
            },
            "fields": [
                {
                    "label": "연관 요구사항",
                    "choiceOptionSetting": {
                        "referenceFilters": [
                            {
                                "domainType": "PROJECT",
                                "domainId": 13526611,
                            }
                        ]
                    },
                }
            ],
        }


class TrackerItemReferenceIdConfigFakeClient(TrackerItemQueryFakeClient):
    def get_tracker_schema(self, tracker_id: int):
        schema = super().get_tracker_schema(tracker_id)
        for field in schema["fields"]:
            if field.get("id") == 7:
                field["id"] = 17
                field["name"] = "SUDS 링크"
        return schema

    def get_tracker_configuration(self, tracker_id: int):
        del tracker_id
        return {
            "basicInformation": {
                "trackerId": 1000,
                "name": "Test Tracker",
            },
            "fields": [
                {
                    "referenceId": 17,
                    "label": "Software Unit Design Specification",
                    "choiceOptionSetting": {
                        "referenceFilters": [
                            {
                                "domainType": "TRACKER",
                                "domainId": 13526611,
                            }
                        ]
                    },
                }
            ],
        }


class UserReferenceDefaultFakeClient(FakeClient):
    def get_tracker_schema(self, tracker_id: int):
        schema = super().get_tracker_schema(tracker_id)
        schema["fields"].append({
            "id": 8,
            "name": "담당 사용자",
            "type": "ReferenceField",
            "referenceType": "UserReference",
            "valueModel": "ChoiceFieldValue<UserReference>",
            "multipleValues": False,
        })
        return schema


class MemberReferenceDefaultFakeClient(FakeClient):
    def get_tracker_schema(self, tracker_id: int):
        schema = super().get_tracker_schema(tracker_id)
        schema["fields"].append({
            "id": 9,
            "name": "검토 담당",
            "type": "MemberField",
            "valueModel": "ChoiceFieldValue",
            "multipleValues": False,
            "memberTypes": ["USER", "ROLE", "GROUP"],
        })
        return schema


class TrackerItemDefaultValueFakeClient(FakeClient):
    def get_tracker_schema(self, tracker_id: int):
        schema = super().get_tracker_schema(tracker_id)
        schema["fields"].append({
            "id": 1007,
            "name": "상위 요구사항",
            "type": "TrackerItemChoiceField",
            "referenceType": "TrackerItemReference",
            "valueModel": "ChoiceFieldValue<TrackerItemReference>",
            "multipleValues": False,
        })
        return schema

    def get_tracker_configuration(self, tracker_id: int):
        del tracker_id
        return {
            "basicInformation": {
                "trackerId": 1000,
                "name": "Test Tracker",
            },
            "fields": [
                {
                    "referenceId": 1007,
                    "label": "상위 요구사항",
                    "choiceOptionSetting": {
                        "referenceFilters": [
                            {
                                "domainType": "TRACKER",
                                "domainId": 13526611,
                            }
                        ]
                    },
                }
            ],
        }


class CountingGuiExcelService(GuiExcelService):
    def __init__(self, **kwargs) -> None:
        super().__init__(**kwargs)
        self.load_preview_calls = 0

    def load_preview(self, *args, **kwargs):
        self.load_preview_calls += 1
        return super().load_preview(*args, **kwargs)


class CountingBatchExcelReader(FakeExcelReader):
    read_excel_calls: list[str] = []
    read_headers_calls: list[str] = []
    list_sheet_calls: list[str] = []
    read_preview_calls: list[str] = []

    @classmethod
    def reset_counts(cls) -> None:
        cls.read_excel_calls = []
        cls.read_headers_calls = []
        cls.list_sheet_calls = []
        cls.read_preview_calls = []

    def list_sheet_names(self, file_path: str) -> list[str]:
        self.__class__.list_sheet_calls.append(str(file_path))
        return super().list_sheet_names(file_path)

    def read_headers(self, file_path: str, sheet_name: str | int) -> list[str]:
        self.__class__.read_headers_calls.append(str(file_path))
        return super().read_headers(file_path, sheet_name)

    def read_excel(self, file_path: str, sheet_name: str | int = 0, visible: bool = False) -> pd.DataFrame:
        self.__class__.read_excel_calls.append(str(file_path))
        return super().read_excel(file_path, sheet_name=sheet_name, visible=visible)

    def read_preview_rows(
        self,
        file_path: str,
        sheet_name: str | int,
        *,
        max_rows: int = 10,
    ) -> tuple[list[str], list[list[object]]]:
        self.__class__.read_preview_calls.append(str(file_path))
        return super().read_preview_rows(
            file_path,
            sheet_name,
            max_rows=max_rows,
        )

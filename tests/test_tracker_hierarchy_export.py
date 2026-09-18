from __future__ import annotations

import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from src.gui.tracker_hierarchy import build_tracker_hierarchy
from src.gui.tracker_hierarchy_export import TrackerHierarchyExportError
from src.gui.tracker_hierarchy_export import build_tracker_hierarchy_export_snapshot
from src.gui.tracker_hierarchy_export import build_tracker_hierarchy_snapshot
from src.gui.tracker_hierarchy_export import create_tracker_hierarchy_workbook
from src.gui.tracker_hierarchy_export import export_tracker_hierarchy_xlsx
from src.gui.tracker_hierarchy_export import hierarchy_export_fields_from_schema
from src.gui.tracker_query_models import TrackerItemSummary
from tests.gui_widget_cleanup import tearDownModule  # noqa: F401


SCHEMA = {
    "id": 20,
    "fields": [
        {
            "id": 1,
            "name": "Summary",
            "type": "TextField",
            "trackerItemField": "name",
        },
        {
            "id": 2,
            "name": "Description",
            "type": "TextField",
        },
        {
            "id": 7,
            "name": "Status",
            "type": "OptionChoiceField",
            "trackerItemField": "status",
        },
        {
            "id": 8,
            "name": "연관 요구사항",
            "type": "TrackerItemChoiceField",
            "valueModel": "ChoiceFieldValue<TrackerItemReference>",
            "multipleValues": True,
        },
        {
            "id": 10,
            "name": "검증 표",
            "type": "TableField",
            "valueModel": "TableFieldValue",
            "columns": [
                {"id": 101, "name": "조건", "type": "TextField"},
                {"id": 102, "name": "결과", "type": "BoolField"},
            ],
        },
        {"id": 99, "name": "Hidden", "type": "TextField", "hidden": True},
    ],
}


def _table_row(condition: str, result: bool) -> list[dict]:
    return [
        {"fieldId": 101, "name": "조건", "value": condition},
        {"fieldId": 102, "name": "결과", "value": result},
    ]


def _item(
    item_id: int,
    name: str,
    *,
    parent_id: int | None = None,
    child_ids: tuple[int, ...] = (),
    ordinal: int = 0,
    table_rows: list[list[dict]] | None = None,
    related_items: list[dict] | None = None,
) -> TrackerItemSummary:
    payload = {
        "id": item_id,
        "name": name,
        "tracker": {"id": 20, "name": "Requirements"},
        "description": f"Description {item_id}",
        "status": {"id": 1, "name": "Open", "type": "ChoiceOptionReference"},
        "ordinal": ordinal,
        "children": [
            {"id": child_id, "name": f"Item {child_id}", "type": "TrackerItemReference"}
            for child_id in child_ids
        ],
        "customFields": [
            {
                "fieldId": 8,
                "name": "연관 요구사항",
                "type": "ChoiceFieldValue<TrackerItemReference>",
                "values": related_items or [],
            },
            {
                "fieldId": 10,
                "name": "검증 표",
                "type": "TableFieldValue",
                "values": table_rows or [],
            }
        ],
    }
    if parent_id is not None:
        payload["parent"] = {"id": parent_id, "name": f"Item {parent_id}"}
    return TrackerItemSummary.from_raw(payload, tracker_id=20)


class TrackerHierarchyExportTest(unittest.TestCase):
    def setUp(self) -> None:
        root = _item(1, "Root", child_ids=(2, 3))
        table_child = _item(
            2,
            "Child table",
            parent_id=1,
            ordinal=2,
            table_rows=[_table_row("A", True), _table_row("B", False)],
        )
        child = _item(3, "Child", parent_id=1, child_ids=(4,), ordinal=3)
        nested = _item(4, "Nested", parent_id=3, ordinal=1)
        self.snapshot = build_tracker_hierarchy_snapshot(
            (nested, child, root, table_child),
            (root,),
            SCHEMA,
            tracker_id=20,
        )

    def test_schema_fields_fix_structural_columns_and_default_table_off(self) -> None:
        fields = hierarchy_export_fields_from_schema(SCHEMA)
        by_key = {field.field_key: field for field in fields}

        self.assertNotIn("name", by_key)
        self.assertNotIn("custom:99", by_key)
        self.assertIn("description", by_key)
        self.assertTrue(by_key["status"].default_selected)
        self.assertTrue(by_key["custom:8"].is_tracker_item_choice)
        self.assertFalse(by_key["custom:10"].default_selected)
        self.assertEqual(
            [column.label for column in by_key["custom:10"].table_columns],
            ["조건", "결과"],
        )

    def test_snapshot_uses_parent_before_children_and_preserves_child_order(self) -> None:
        self.assertEqual(
            [(node.item.item_id, node.parent_id, node.depth) for node in self.snapshot.nodes],
            [(1, None, 0), (2, 1, 1), (3, 1, 1), (4, 3, 2)],
        )

    def test_existing_baseline_hierarchy_can_be_exported_without_rebuilding_it(self) -> None:
        items = tuple(node.item for node in self.snapshot.nodes)
        hierarchy = build_tracker_hierarchy(items, tracker_id=20)

        snapshot = build_tracker_hierarchy_export_snapshot(hierarchy, SCHEMA)
        workbook, summary = create_tracker_hierarchy_workbook(
            snapshot,
            tracker_name="Requirements (ID 20)",
            project_name="Vehicle",
            selected_field_keys=("status",),
            baseline_id=701,
            baseline_name="Release 1",
            generated_at=datetime(2026, 8, 20, 12, 0, 0),
        )

        info_values = {
            workbook["내보내기 정보"].cell(row, 1).value:
            workbook["내보내기 정보"].cell(row, 2).value
            for row in range(1, workbook["내보내기 정보"].max_row + 1)
        }
        self.assertEqual(info_values["조회 기준"], "Baseline")
        self.assertEqual(info_values["Baseline"], "Release 1")
        self.assertEqual(info_values["Baseline ID"], 701)
        self.assertEqual(summary.item_count, 4)

    def test_snapshot_rejects_missing_hierarchy_metadata_without_per_item_fallback(self) -> None:
        root = _item(1, "Root")
        unresolved = _item(2, "Unknown parent")

        with self.assertRaisesRegex(TrackerHierarchyExportError, "상위 관계"):
            build_tracker_hierarchy_snapshot(
                (root, unresolved),
                (root,),
                SCHEMA,
                tracker_id=20,
            )

    def test_snapshot_rejects_conflicting_parent_references(self) -> None:
        root = _item(1, "Root", child_ids=(3,))
        other_root = _item(2, "Other root")
        child = _item(3, "Child", parent_id=2)

        with self.assertRaisesRegex(TrackerHierarchyExportError, "두 개의 상위"):
            build_tracker_hierarchy_snapshot(
                (root, other_root, child),
                (root, other_root),
                SCHEMA,
                tracker_id=20,
            )

    def test_snapshot_rejects_disconnected_cycle(self) -> None:
        root = _item(1, "Root")
        first = _item(2, "First", parent_id=3)
        second = _item(3, "Second", parent_id=2)

        with self.assertRaisesRegex(TrackerHierarchyExportError, "순환 계층"):
            build_tracker_hierarchy_snapshot(
                (root, first, second),
                (root,),
                SCHEMA,
                tracker_id=20,
            )

    def test_workbook_preserves_indent_outline_and_table_rows(self) -> None:
        workbook, summary = create_tracker_hierarchy_workbook(
            self.snapshot,
            tracker_name="Requirements (ID 20)",
            project_name="Vehicle",
            selected_field_keys=("status", "custom:10"),
            generated_at=datetime(2026, 8, 13, 12, 0, 0),
        )
        sheet = workbook["트래커 항목"]

        self.assertEqual(workbook.sheetnames, ["내보내기 정보", "트래커 항목"])
        self.assertEqual(sheet["A1"].value, "ID")
        self.assertEqual(sheet["B1"].value, "Summary")
        self.assertEqual(sheet["E1"].value, "Status")
        self.assertEqual(sheet["F1"].value, "검증 표")
        self.assertEqual(sheet["F2"].value, "조건")
        self.assertEqual(sheet["G2"].value, "결과")
        self.assertIn("A4:A5", {str(value) for value in sheet.merged_cells.ranges})
        self.assertEqual(sheet["F4"].value, "A")
        self.assertEqual(sheet["G4"].value, "예")
        self.assertEqual(sheet["F5"].value, "B")
        self.assertEqual(sheet["G5"].value, "아니요")
        self.assertEqual(sheet["B4"].alignment.indent, 1.0)
        self.assertEqual(sheet.row_dimensions[4].outlineLevel, 1)
        self.assertEqual(sheet.row_dimensions[7].outlineLevel, 2)
        self.assertEqual(summary.item_count, 4)
        self.assertEqual(summary.data_row_count, 5)

    def test_tracker_item_choice_field_exports_names_only(self) -> None:
        root = _item(
            11,
            "Root",
            related_items=[
                {"id": 501, "name": "요구사항 A", "type": "TrackerItemReference"},
                {
                    "id": 502,
                    "name": "요구사항 B",
                    "type": "TrackerItemReference",
                    "tracker": {"id": 20, "name": "Requirements"},
                },
                {"id": 503, "type": "TrackerItemReference"},
            ],
        )
        snapshot = build_tracker_hierarchy_snapshot(
            (root,),
            (root,),
            SCHEMA,
            tracker_id=20,
        )

        workbook, _summary = create_tracker_hierarchy_workbook(
            snapshot,
            tracker_name="Requirements",
            selected_field_keys=("custom:8",),
        )

        value = workbook["트래커 항목"]["E3"].value
        self.assertEqual(value, "요구사항 A\n요구사항 B")
        self.assertNotIn("501", value)

    def test_fixed_columns_only_and_formula_injection_are_safe(self) -> None:
        injected = _item(9, "=HYPERLINK('unsafe')")
        snapshot = build_tracker_hierarchy_snapshot(
            (injected,),
            (injected,),
            SCHEMA,
            tracker_id=20,
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "hierarchy.xlsx"
            summary = export_tracker_hierarchy_xlsx(
                snapshot,
                path,
                tracker_name="Requirements",
                selected_field_keys=(),
            )
            workbook = load_workbook(path, data_only=False)

        self.assertEqual(workbook["트래커 항목"]["B3"].value, "'=HYPERLINK('unsafe')")
        self.assertEqual(summary.selected_field_count, 0)
        self.assertEqual(summary.output_path, str(path))

    def test_long_value_is_preserved_in_linked_sheet(self) -> None:
        long_item = _item(10, "Long value")
        payload = dict(long_item.raw_reference)
        payload["description"] = "가" * 30100
        long_item = TrackerItemSummary.from_raw(payload, tracker_id=20)
        snapshot = build_tracker_hierarchy_snapshot(
            (long_item,),
            (long_item,),
            SCHEMA,
            tracker_id=20,
        )

        workbook, summary = create_tracker_hierarchy_workbook(
            snapshot,
            tracker_name="Requirements",
            selected_field_keys=("description",),
        )

        self.assertIn("긴 값 전체보기", workbook.sheetnames)
        self.assertEqual(summary.long_value_count, 1)
        self.assertIsNotNone(workbook["트래커 항목"]["E3"].hyperlink)
        restored = "".join(
            workbook["긴 값 전체보기"].cell(row, 5).value or ""
            for row in range(2, workbook["긴 값 전체보기"].max_row + 1)
        )
        self.assertEqual(restored, payload["description"])

    def test_column_limit_is_checked_before_workbook_population(self) -> None:
        with (
            patch("src.gui.tracker_hierarchy_export.EXCEL_MAX_COLUMNS", 4),
            self.assertRaisesRegex(TrackerHierarchyExportError, "열 한도"),
        ):
            create_tracker_hierarchy_workbook(
                self.snapshot,
                tracker_name="Requirements",
                selected_field_keys=("status",),
            )

    def test_failed_save_keeps_existing_target_and_removes_temporary_file(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "hierarchy.xlsx"
            path.write_bytes(b"existing")

            with (
                patch("src.gui.tracker_hierarchy_export.Workbook.save", side_effect=OSError("disk")),
                self.assertRaisesRegex(TrackerHierarchyExportError, "저장하지 못했습니다"),
            ):
                export_tracker_hierarchy_xlsx(
                    self.snapshot,
                    path,
                    tracker_name="Requirements",
                    selected_field_keys=(),
                )

            self.assertEqual(path.read_bytes(), b"existing")
            self.assertEqual([value.name for value in Path(temp_dir).iterdir()], ["hierarchy.xlsx"])


if __name__ == "__main__":
    unittest.main()

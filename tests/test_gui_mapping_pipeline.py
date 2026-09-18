from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from src.gui.service_core import GuiExcelService
from src.gui.settings_store import GuiSettings
from src.gui.upload_service import GuiUploadPipelineService
from src.models import PayloadStatus
from src.upload_policy import UPLOAD_MODE_UPDATE as GUI_UPLOAD_MODE_UPDATE
from src.upload_policy import UPLOAD_MODE_UPSERT as GUI_UPLOAD_MODE_UPSERT
from tests.gui_service_fixtures import CountingGuiExcelService
from tests.gui_service_fixtures import FakeClient
from tests.gui_service_fixtures import FakeExcelReader
from tests.gui_service_fixtures import UpdateModeFakeClient
from tests.gui_service_fixtures import UpsertModeFakeClient


class GuiMappingPipelineServiceTest(unittest.TestCase):
    def test_prepare_mapping_context_builds_upload_columns_and_default_mapping(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "담당자", "테이블필드.컬럼A", "id", "parent"])
            sheet.append(["REQ-001", "홍길동", "값1", "1", ""])
            workbook.save(path)
            workbook.close()

            excel_service = GuiExcelService(reader_cls=FakeExcelReader)
            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=excel_service,
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            self.assertIn("Summary", mapping_context.upload_columns)
            self.assertIn("담당자", mapping_context.upload_columns)
            self.assertIn("테이블필드.컬럼A", mapping_context.upload_columns)
            self.assertNotIn("id", mapping_context.upload_columns)
            self.assertNotIn("parent", mapping_context.upload_columns)
            self.assertNotIn("id", mapping_context.schema_df["field_name"].tolist())
            self.assertNotIn("parent", mapping_context.schema_df["field_name"].tolist())
            self.assertEqual(mapping_context.selected_mapping["Summary"], "Summary")
            self.assertEqual(mapping_context.selected_mapping["담당자"], "담당자")
            self.assertEqual(mapping_context.selected_mapping["테이블필드.컬럼A"], "테이블필드")
            self.assertEqual(
                [candidate.schema_field for candidate in mapping_context.default_value_candidates],
                ["Summary", "Status", "담당자"],
            )
            self.assertTrue(mapping_context.default_value_candidates[0].allows_custom_value)
            self.assertEqual(mapping_context.default_value_candidates[1].options, ["Open", "Review"])
            self.assertEqual(mapping_context.file_paths, [str(path)])

    def test_prepare_mapping_context_keeps_update_mode(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["id", "Summary", "담당자"])
            sheet.append([101, "REQ-101", "홍길동"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=UpdateModeFakeClient,
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                upload_mode=GUI_UPLOAD_MODE_UPDATE,
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            self.assertEqual(mapping_context.upload_mode, GUI_UPLOAD_MODE_UPDATE)
            self.assertEqual(mapping_context.wizard.state.upload_mode, GUI_UPLOAD_MODE_UPDATE)

    def test_validate_mapping_blocks_update_mode_when_id_column_missing(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "담당자"])
            sheet.append(["REQ-101", "홍길동"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                upload_mode=GUI_UPLOAD_MODE_UPDATE,
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )
            validation_context = service.validate_mapping(
                mapping_context,
                {"Summary": "Summary"},
            )

            self.assertTrue(validation_context.has_blocking_issues)
            self.assertTrue(
                validation_context.issue_df["message"].str.contains("id 열", regex=False).any()
            )
            self.assertEqual(
                list(mapping_context.wizard.state.payload_df["payload_status"]),
                [PayloadStatus.FAILED.value],
            )

    def test_run_batch_upload_uses_update_mode_put_payload(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            UpdateModeFakeClient.reset_calls()
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["id", "Summary", "담당자"])
            sheet.append([101, "REQ-101", "홍길동"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=UpdateModeFakeClient,
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                upload_mode=GUI_UPLOAD_MODE_UPDATE,
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            file_state = {
                "file_path": str(path),
                "file_paths": [str(path)],
                "preview_file_path": str(path),
                "sheet_name": "Main",
                "header_row": 1,
                "summary_column": "Summary",
            }

            mapping_context = service.prepare_mapping_context(settings, file_state)
            validation_context = service.validate_mapping(
                mapping_context,
                {"Summary": "Summary", "담당자": "담당자"},
                {"Status": "Review"},
            )

            self.assertFalse(validation_context.has_blocking_issues)
            payload = mapping_context.wizard.state.payload_df.iloc[0]["payload_json"]
            self.assertEqual(payload["id"], 101)
            self.assertEqual(payload["name"], "REQ-101")
            self.assertEqual(payload["status"]["name"], "Review")
            self.assertEqual(UpdateModeFakeClient.all_get_item_calls, [])

            result = service.run_batch_upload(
                settings,
                file_state,
                mapping_context,
                dry_run=False,
                continue_on_error=True,
                output_dir=str(Path(tmp_dir) / "output"),
            )

            self.assertEqual(len(result["success_df"]), 1)
            self.assertTrue(result["failed_df"].empty)
            self.assertTrue(result["unresolved_df"].empty)
            self.assertEqual(UpdateModeFakeClient.all_get_item_calls, [101])
            self.assertEqual(UpdateModeFakeClient.all_update_calls[0][0], 101)
            self.assertEqual(UpdateModeFakeClient.all_update_calls[0][1]["name"], "REQ-101")
            self.assertEqual(UpdateModeFakeClient.all_update_calls[0][1]["description"], "기존 설명")
            self.assertEqual(UpdateModeFakeClient.all_update_calls[0][1]["status"]["name"], "Review")
            self.assertEqual(
                next(
                    field["value"]
                    for field in UpdateModeFakeClient.all_update_calls[0][1]["customFields"]
                    if field["fieldId"] == 3
                ),
                "홍길동",
            )
            self.assertEqual(
                next(
                    field["value"]
                    for field in UpdateModeFakeClient.all_update_calls[0][1]["customFields"]
                    if field["fieldId"] == 999
                ),
                "보존",
            )
            self.assertEqual(mapping_context.representative_file_path, str(path))

    def test_run_batch_upload_uses_upsert_mode_for_create_and_update_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            UpsertModeFakeClient.reset_calls()
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["id", "Summary", "담당자"])
            sheet.append([101, "REQ-101", "홍길동"])
            sheet.append([None, "REQ-NEW", "신규담당자"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=UpsertModeFakeClient,
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                upload_mode=GUI_UPLOAD_MODE_UPSERT,
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            file_state = {
                "file_path": str(path),
                "file_paths": [str(path)],
                "preview_file_path": str(path),
                "sheet_name": "Main",
                "header_row": 1,
                "summary_column": "Summary",
            }

            mapping_context = service.prepare_mapping_context(settings, file_state)
            validation_context = service.validate_mapping(
                mapping_context,
                {"Summary": "Summary", "담당자": "담당자"},
                {"Status": "Review"},
            )

            self.assertFalse(validation_context.has_blocking_issues)
            emitted_events: list[dict[str, object]] = []

            result = service.run_batch_upload(
                settings,
                file_state,
                mapping_context,
                dry_run=False,
                continue_on_error=True,
                output_dir=str(Path(tmp_dir) / "output"),
                event_callback=lambda event: emitted_events.append(dict(event)),
            )

            self.assertEqual(len(result["success_df"]), 2)
            self.assertEqual(result["success_df"]["phase"].tolist(), ["insert", "update"])
            self.assertEqual(
                result["phase_results"],
                {
                    "insert": {"total": 1, "success": 1, "failed": 0, "unresolved": 0},
                    "update": {"total": 1, "success": 1, "failed": 0, "unresolved": 0},
                },
            )
            self.assertEqual(len(UpsertModeFakeClient.all_create_item_calls), 1)
            self.assertEqual(UpsertModeFakeClient.all_create_item_calls[0]["payload"]["name"], "REQ-NEW")
            self.assertIsNone(UpsertModeFakeClient.all_create_item_calls[0]["parent_item_id"])
            self.assertEqual(UpsertModeFakeClient.all_update_calls[0][0], 101)
            self.assertEqual(UpsertModeFakeClient.all_update_calls[0][1]["name"], "REQ-101")
            phase_started_events = [event for event in emitted_events if event.get("type") == "phase_started"]
            phase_finished_events = [event for event in emitted_events if event.get("type") == "phase_finished"]
            self.assertEqual(
                [(event.get("phase"), event.get("total")) for event in phase_started_events],
                [("insert", 1), ("update", 1)],
            )
            self.assertEqual(
                [(event.get("phase"), event.get("success"), event.get("failed")) for event in phase_finished_events],
                [("insert", 1, 0), ("update", 1, 0)],
            )

    def test_validate_mapping_blocks_duplicate_update_ids_across_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            first_path = Path(tmp_dir) / "first.xlsx"
            second_path = Path(tmp_dir) / "second.xlsx"

            for path, summary in (
                (first_path, "REQ-101-A"),
                (second_path, "REQ-101-B"),
            ):
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Main"
                sheet.append(["id", "Summary"])
                sheet.append([101, summary])
                workbook.save(path)
                workbook.close()

            service = GuiUploadPipelineService(
                client_factory=UpdateModeFakeClient,
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                upload_mode=GUI_UPLOAD_MODE_UPDATE,
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            file_state = {
                "file_path": str(first_path),
                "file_paths": [str(first_path), str(second_path)],
                "preview_file_path": str(first_path),
                "sheet_name": "Main",
                "header_row": 1,
                "summary_column": "Summary",
            }

            mapping_context = service.prepare_mapping_context(settings, file_state)
            validation_context = service.validate_mapping(
                mapping_context,
                {"Summary": "Summary"},
            )

            self.assertTrue(validation_context.has_blocking_issues)
            self.assertEqual(mapping_context.batch_duplicate_update_item_ids, {101})
            self.assertTrue(
                validation_context.issue_df["message"]
                .str.contains("여러 파일에서 같은 item id가 중복", regex=False)
                .any()
            )

    def test_prepare_mapping_context_uses_file_selection_header_and_summary(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["메타", "메타", "메타"])
            sheet.append(["요약", "담당자", "비고"])
            sheet.append(["REQ-001", "홍길동", "메모"])
            workbook.save(path)
            workbook.close()

            excel_service = GuiExcelService(reader_cls=FakeExcelReader)
            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=excel_service,
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 2,
                    "summary_column": "요약",
                },
            )

            self.assertEqual(mapping_context.wizard.reader.header_row, 2)
            self.assertEqual(mapping_context.wizard.reader.summary_col, "요약")
            self.assertEqual(mapping_context.wizard.processor.summary_col, "요약")
            self.assertIn("요약", mapping_context.upload_columns)
            self.assertEqual(
                mapping_context.wizard.state.upload_df.iloc[0]["upload_name"],
                "REQ-001",
            )

    def test_prepare_mapping_context_keeps_original_upload_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "담당자", "비고"])
            sheet.append(["REQ-001", "홍길동", "메모"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            upload_df = mapping_context.wizard.state.upload_df

            self.assertEqual(list(upload_df["upload_name"]), ["REQ-001"])
            self.assertNotIn("_synthetic_root", upload_df.columns)

    def test_apply_saved_workflow_values_keeps_auto_mapping_when_saved_columns_do_not_match(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "담당자", "테이블필드.컬럼A", "id", "parent"])
            sheet.append(["REQ-001", "홍길동", "값1", "1", ""])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            service.apply_saved_workflow_values(
                mapping_context,
                root_item_config={"enabled": False},
                selected_mapping={
                    "Summary": "Status",
                    "예전담당자": "담당자",
                },
                selected_mapping_modes={
                    "Summary": {"create": False, "update": True},
                    "예전담당자": {"create": True, "update": False},
                },
                selected_default_values={
                    "Status": "Open",
                    "담당자": "홍길동",
                    "없는필드": "무시",
                },
                selected_default_value_modes={
                    "Status": {"create": True, "update": False},
                    "담당자": {"create": False, "update": True},
                    "없는필드": {"create": False, "update": True},
                },
                selected_tracker_item_settings={
                    "없는필드": {
                        "mode": "query",
                        "source_tracker_ids": [13526611],
                    }
                },
            )

            self.assertFalse(mapping_context.root_item_config["enabled"])
            self.assertEqual(mapping_context.selected_mapping["Summary"], "Status")
            self.assertEqual(mapping_context.selected_mapping["담당자"], "담당자")
            self.assertEqual(mapping_context.selected_mapping["테이블필드.컬럼A"], "테이블필드")
            self.assertNotIn("예전담당자", mapping_context.selected_mapping)
            self.assertEqual(mapping_context.selected_mapping_modes["Summary"], {"create": False, "update": True})
            self.assertEqual(mapping_context.selected_mapping_modes["담당자"], {"create": True, "update": False})
            self.assertEqual(mapping_context.selected_default_values, {"Status": "Open", "담당자": "홍길동"})
            self.assertEqual(mapping_context.selected_default_value_modes["Status"], {"create": True, "update": False})
            self.assertEqual(mapping_context.selected_default_value_modes["담당자"], {"create": False, "update": False})
            self.assertEqual(mapping_context.selected_tracker_item_settings, {})

    def test_prepare_mapping_context_reuses_cached_preview_data(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "담당자"])
            sheet.append(["REQ-001", "홍길동"])
            workbook.save(path)
            workbook.close()

            excel_service = CountingGuiExcelService(reader_cls=FakeExcelReader)
            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=excel_service,
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            preview = excel_service.load_preview(
                str(path),
                sheet_name="Main",
                header_row=1,
                summary_column="Summary",
            )

            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                    "preview_data": preview,
                },
            )

            self.assertEqual(excel_service.load_preview_calls, 1)
            self.assertIsNotNone(mapping_context.preview_data)
            self.assertEqual(mapping_context.preview_data.file_path, str(path))

    def test_prepare_mapping_context_rejects_full_data_when_source_file_changed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "담당자"])
            sheet.append(["REQ-001", "홍길동"])
            workbook.save(path)
            workbook.close()

            excel_service = GuiExcelService(reader_cls=FakeExcelReader)
            preview = excel_service.load_full_data(
                str(path),
                sheet_name="Main",
                header_row=1,
                summary_column="Summary",
            )

            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "담당자"])
            sheet.append(["REQ-001-수정", "홍길동"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=excel_service,
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            with self.assertRaisesRegex(ValueError, "전체 데이터를 다시 불러오세요"):
                service.prepare_mapping_context(
                    settings,
                    {
                        "file_path": str(path),
                        "preview_file_path": str(path),
                        "sheet_name": "Main",
                        "header_row": 1,
                        "summary_column": "Summary",
                        "preview_data": preview,
                    },
                )

    def test_prepare_mapping_context_supports_offline_schema_snapshot(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            schema_path = Path(tmp_dir) / "offline-schema.json"
            schema_payload = FakeClient("", "", "").get_tracker_schema(0)
            schema_payload["name"] = "Offline Tracker Snapshot"
            schema_path.write_text(json.dumps(schema_payload, ensure_ascii=False), encoding="utf-8")

            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary", "담당자"])
            sheet.append(["REQ-001", "홍길동"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                offline_mode=True,
                offline_schema_path=str(schema_path),
                default_project_id="501",
                default_tracker_id="601",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            mapping_context = service.prepare_mapping_context(
                settings,
                {
                    "file_path": str(path),
                    "preview_file_path": str(path),
                    "sheet_name": "Main",
                    "header_row": 1,
                    "summary_column": "Summary",
                },
            )

            self.assertEqual(mapping_context.selected_mapping["Summary"], "Summary")
            self.assertEqual(mapping_context.selected_mapping["담당자"], "담당자")
            self.assertEqual(mapping_context.wizard.state.tracker_id, 601)

    def test_run_batch_upload_supports_offline_mode_with_dry_run(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            schema_path = Path(tmp_dir) / "offline-schema.json"
            schema_payload = FakeClient("", "", "").get_tracker_schema(0)
            schema_payload["name"] = "Offline Tracker Snapshot"
            schema_path.write_text(json.dumps(schema_payload, ensure_ascii=False), encoding="utf-8")

            path = Path(tmp_dir) / "sample.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Main"
            sheet.append(["Summary"])
            sheet.append(["REQ-001"])
            workbook.save(path)
            workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                offline_mode=True,
                offline_schema_path=str(schema_path),
                default_project_id="501",
                default_tracker_id="601",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            file_state = {
                "file_path": str(path),
                "file_paths": [str(path)],
                "preview_file_path": str(path),
                "sheet_name": "Main",
                "header_row": 1,
                "summary_column": "Summary",
            }

            mapping_context = service.prepare_mapping_context(settings, file_state)
            mapping_context.root_item_config = {"enabled": False}
            validation_context = service.validate_mapping(
                mapping_context,
                {"Summary": "Summary"},
            )

            self.assertFalse(validation_context.has_blocking_issues)

            result = service.run_batch_upload(
                settings,
                file_state,
                mapping_context,
                dry_run=True,
                continue_on_error=True,
                output_dir=str(Path(tmp_dir) / "output"),
            )

            self.assertEqual(len(result["success_df"]), 1)
            self.assertEqual(result["success_df"].iloc[0]["created_item_id"], "DRYRUN-0")

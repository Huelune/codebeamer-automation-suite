from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from src.gui.root_item_service import ROOT_ITEM_MODE_GROUP_BY_COLUMN
from src.gui.root_item_service import ROOT_SOURCE_GROUP_VALUE
from src.gui.service_core import GuiExcelService
from src.gui.settings_store import GuiSettings
from src.gui.upload_service import GuiUploadPipelineService
from src.models import UserInfo
from tests.gui_service_fixtures import CountingBatchExcelReader
from tests.gui_service_fixtures import FakeClient
from tests.gui_service_fixtures import FakeExcelReader


class CachedMemberLookupFakeClient(FakeClient):
    all_user_lookup_calls: list[str] = []
    all_group_lookup_calls = 0
    all_role_lookup_calls: list[tuple[int, int]] = []

    @classmethod
    def reset_calls(cls) -> None:
        cls.all_user_lookup_calls = []
        cls.all_group_lookup_calls = 0
        cls.all_role_lookup_calls = []

    def get_tracker_schema(self, tracker_id: int):
        schema = super().get_tracker_schema(tracker_id)
        schema["fields"].append({
            "id": 9,
            "name": "검토 담당",
            "type": "MemberField",
            "valueModel": "ChoiceFieldValue",
            "multipleValues": False,
            "memberTypes": ["USER", "ROLE", "GROUP"],
        })
        return schema

    def get_user_by_name(self, name: str) -> UserInfo:
        self.__class__.all_user_lookup_calls.append(name)
        return UserInfo(id=101, name=name)

    def get_user_groups(self):
        self.__class__.all_group_lookup_calls += 1
        return [{"id": 301, "name": "QA Group", "type": "GroupReference"}]

    def get_tracker_field_permissions(self, tracker_id: int, field_id: int):
        self.__class__.all_role_lookup_calls.append((int(tracker_id), int(field_id)))
        return [{
            "status": {"id": 1, "name": "Open", "type": "ChoiceOptionReference"},
            "permissions": [{
                "role": {"id": 201, "name": "Developer", "type": "RoleReference"},
            }],
        }]

class GuiBatchPipelineServiceTest(unittest.TestCase):
    def test_prepare_mapping_context_rejects_mismatched_batch_headers(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            first_path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            second_path = Path(tmp_dir) / "DEF_REQ-002.xlsx"

            first_book = Workbook()
            first_sheet = first_book.active
            first_sheet.title = "Main"
            first_sheet.append(["Summary", "담당자", "비고"])
            first_sheet.append(["REQ-001", "홍길동", "메모"])
            first_book.save(first_path)
            first_book.close()

            second_book = Workbook()
            second_sheet = second_book.active
            second_sheet.title = "Main"
            second_sheet.append(["Summary", "상태", "비고"])
            second_sheet.append(["REQ-002", "Open", "메모2"])
            second_book.save(second_path)
            second_book.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )

            with self.assertRaisesRegex(ValueError, "헤더가 기준 파일과 다릅니다"):
                service.prepare_mapping_context(
                    settings,
                    {
                        "file_path": str(first_path),
                        "file_paths": [str(first_path), str(second_path)],
                        "preview_file_path": str(first_path),
                        "sheet_name": "Main",
                        "header_row": 1,
                        "summary_column": "Summary",
                    },
                )

    def test_run_batch_upload_aggregates_multiple_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            first_path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            second_path = Path(tmp_dir) / "DEF_REQ-002.xlsx"

            for path, summary in (
                (first_path, "REQ-001"),
                (second_path, "REQ-002"),
            ):
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Main"
                sheet.append(["Summary"])
                sheet.append([summary])
                workbook.save(path)
                workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            file_state = {
                "file_path": str(first_path),
                "file_paths": [str(first_path), str(second_path)],
                "preview_file_path": str(first_path),
                "sheet_name": "Main",
                "header_row": 1,
                "summary_column": "Summary",
            }

            mapping_context = service.prepare_mapping_context(settings, file_state)
            mapping_context.root_item_config = {
                "regex_pattern": r"^(?P<project>[A-Z]+)_(?P<title>.+)$",
                "regex_target": "file_stem",
                "field_assignments": {
                    "Summary": {
                        "enabled": True,
                        "mode": "file_source",
                        "value": "title",
                    },
                    "Status": {
                        "enabled": True,
                        "mode": "fixed_value",
                        "value": "Open",
                    },
                },
            }
            validation_context = service.validate_mapping(
                mapping_context,
                mapping_context.selected_mapping,
            )

            self.assertFalse(validation_context.has_blocking_issues)
            self.assertEqual(validation_context.summary_stats["file_count"], 2)
            self.assertEqual(validation_context.summary_stats["batch_total_rows"], 2)

            result = service.run_batch_upload(
                settings,
                file_state,
                mapping_context,
                dry_run=True,
                continue_on_error=True,
                output_dir=str(Path(tmp_dir) / "output"),
            )

            success_df = result["success_df"]
            self.assertEqual(len(success_df), 4)
            self.assertEqual(set(success_df["source_file"].tolist()), {"ABC_REQ-001.xlsx", "DEF_REQ-002.xlsx"})
            root_rows = success_df[success_df["_row_id"].isna()].reset_index(drop=True)
            self.assertEqual(root_rows["upload_name"].tolist(), ["REQ-001", "REQ-002"])
            self.assertTrue(result["failed_df"].empty)
            self.assertTrue(result["unresolved_df"].empty)

    def test_run_batch_upload_reuses_validation_member_caches_across_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            first_path = Path(tmp_dir) / "first.xlsx"
            second_path = Path(tmp_dir) / "second.xlsx"

            for path, summary in (
                (first_path, "REQ-001"),
                (second_path, "REQ-002"),
            ):
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Main"
                sheet.append(["Summary", "검토 담당"])
                sheet.append([summary, "Kim QA"])
                workbook.save(path)
                workbook.close()

            CachedMemberLookupFakeClient.reset_calls()
            service = GuiUploadPipelineService(
                client_factory=CachedMemberLookupFakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            file_state = {
                "file_path": str(first_path),
                "file_paths": [str(first_path), str(second_path)],
                "preview_file_path": str(first_path),
                "sheet_name": "Main",
                "header_row": 1,
                "summary_column": "Summary",
            }

            mapping_context = service.prepare_mapping_context(settings, file_state)
            validation_context = service.validate_mapping(
                mapping_context,
                mapping_context.selected_mapping,
            )

            self.assertFalse(validation_context.has_blocking_issues)
            self.assertEqual(CachedMemberLookupFakeClient.all_user_lookup_calls, ["Kim QA"])
            self.assertEqual(CachedMemberLookupFakeClient.all_group_lookup_calls, 1)
            self.assertEqual(CachedMemberLookupFakeClient.all_role_lookup_calls, [(1000, 9)])
            self.assertTrue(mapping_context.user_lookup_cache)
            self.assertTrue(mapping_context.member_lookup_cache)
            self.assertTrue(mapping_context.group_lookup_cache)
            self.assertTrue(mapping_context.tracker_role_cache)

            service.run_batch_upload(
                settings,
                file_state,
                mapping_context,
                dry_run=True,
                continue_on_error=True,
                output_dir=str(Path(tmp_dir) / "output"),
            )

            self.assertEqual(CachedMemberLookupFakeClient.all_user_lookup_calls, ["Kim QA"])
            self.assertEqual(CachedMemberLookupFakeClient.all_group_lookup_calls, 1)
            self.assertEqual(CachedMemberLookupFakeClient.all_role_lookup_calls, [(1000, 9)])

    def test_run_batch_upload_creates_grouped_root_items_per_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            first_path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            second_path = Path(tmp_dir) / "DEF_REQ-002.xlsx"

            for path, rows in (
                (first_path, [("REQ-001", "EMS"), ("REQ-002", "VCU")]),
                (second_path, [("REQ-003", "EMS")]),
            ):
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Main"
                sheet.append(["Summary", "Folder"])
                for summary, folder in rows:
                    sheet.append([summary, folder])
                workbook.save(path)
                workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            file_state = {
                "file_path": str(first_path),
                "file_paths": [str(first_path), str(second_path)],
                "preview_file_path": str(first_path),
                "sheet_name": "Main",
                "header_row": 1,
                "summary_column": "Summary",
            }

            mapping_context = service.prepare_mapping_context(settings, file_state)
            mapping_context.root_item_config = {
                "root_mode": ROOT_ITEM_MODE_GROUP_BY_COLUMN,
                "group_by_column": "Folder",
                "regex_pattern": "",
                "regex_target": "file_stem",
                "field_assignments": {
                    "Summary": {
                        "enabled": True,
                        "mode": "file_source",
                        "value": ROOT_SOURCE_GROUP_VALUE,
                    },
                },
            }
            validation_context = service.validate_mapping(
                mapping_context,
                mapping_context.selected_mapping,
            )

            self.assertFalse(validation_context.has_blocking_issues)

            result = service.run_batch_upload(
                settings,
                file_state,
                mapping_context,
                dry_run=True,
                continue_on_error=True,
                output_dir=str(Path(tmp_dir) / "output"),
            )

            success_df = result["success_df"]
            root_rows = success_df[success_df["_row_id"].isna()].reset_index(drop=True)
            self.assertEqual(len(root_rows), 3)
            self.assertEqual(
                root_rows[["source_file", "upload_name"]].values.tolist(),
                [
                    ["ABC_REQ-001.xlsx", "EMS"],
                    ["ABC_REQ-001.xlsx", "VCU"],
                    ["DEF_REQ-002.xlsx", "EMS"],
                ],
            )
            self.assertEqual(len(success_df), 6)

    def test_validate_mapping_aggregates_batch_file_issues_beyond_representative_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            first_path = Path(tmp_dir) / "first.xlsx"
            second_path = Path(tmp_dir) / "second.xlsx"

            for path, status in (
                (first_path, "Open"),
                (second_path, "Unknown"),
            ):
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Main"
                sheet.append(["Summary", "Status"])
                sheet.append([f"REQ-{path.stem}", status])
                workbook.save(path)
                workbook.close()

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=GuiExcelService(reader_cls=FakeExcelReader),
                reader_cls=FakeExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            file_state = {
                "file_path": str(first_path),
                "file_paths": [str(first_path), str(second_path)],
                "preview_file_path": str(first_path),
                "sheet_name": "Main",
                "header_row": 1,
                "summary_column": "Summary",
            }

            mapping_context = service.prepare_mapping_context(settings, file_state)
            validation_context = service.validate_mapping(
                mapping_context,
                mapping_context.selected_mapping,
            )

            self.assertTrue(validation_context.has_blocking_issues)
            self.assertEqual(validation_context.summary_stats["file_count"], 2)
            self.assertEqual(validation_context.summary_stats["total_rows"], 2)
            self.assertEqual(validation_context.summary_stats["error_rows"], 1)
            self.assertIn("second.xlsx", validation_context.issue_df["source_file"].tolist())
            self.assertTrue(
                validation_context.issue_df["message"].str.contains("옵션 목록에 없습니다").any()
            )

    def test_run_batch_upload_reuses_preloaded_raw_data_for_all_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            first_path = Path(tmp_dir) / "ABC_REQ-001.xlsx"
            second_path = Path(tmp_dir) / "DEF_REQ-002.xlsx"

            for path, summary in (
                (first_path, "REQ-001"),
                (second_path, "REQ-002"),
            ):
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Main"
                sheet.append(["Summary"])
                sheet.append([summary])
                workbook.save(path)
                workbook.close()

            CountingBatchExcelReader.reset_counts()
            excel_service = GuiExcelService(reader_cls=CountingBatchExcelReader)
            preview = excel_service.load_preview(
                str(first_path),
                file_paths=[str(first_path), str(second_path)],
                sheet_name="Main",
                header_row=1,
                summary_column="Summary",
            )

            service = GuiUploadPipelineService(
                client_factory=FakeClient,
                excel_service=excel_service,
                reader_cls=CountingBatchExcelReader,
            )
            settings = GuiSettings(
                base_url="https://example.com/cb",
                username="user",
                password="secret",
                default_project_id="10",
                default_tracker_id="1000",
                excel_header_row=1,
                summary_column="Summary",
                excel_sheet_name="Main",
            )
            file_state = {
                "file_path": str(first_path),
                "file_paths": [str(first_path), str(second_path)],
                "preview_file_path": str(first_path),
                "sheet_name": "Main",
                "header_row": 1,
                "summary_column": "Summary",
                "preview_data": preview,
            }

            mapping_context = service.prepare_mapping_context(settings, file_state)
            self.assertEqual(
                CountingBatchExcelReader.read_excel_calls,
                [str(first_path), str(second_path)],
            )

            service.validate_mapping(
                mapping_context,
                mapping_context.selected_mapping,
            )
            result = service.run_batch_upload(
                settings,
                file_state,
                mapping_context,
                dry_run=True,
                continue_on_error=True,
                output_dir=str(Path(tmp_dir) / "output"),
            )

            self.assertEqual(
                CountingBatchExcelReader.read_excel_calls,
                [str(first_path), str(second_path)],
            )
            self.assertEqual(len(result["success_df"]), 4)

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.gui.developer_excel_tools import DeveloperExcelToolError
from src.gui.upload_workbook_tools import UploadWorkbookService


class UploadWorkbookServiceTest(unittest.TestCase):
    def setUp(self) -> None:
        self.service = UploadWorkbookService()

    def test_validation_report_contains_summary_and_user_visible_issues(self) -> None:
        issues = pd.DataFrame(
            [
                {
                    "severity": "오류",
                    "source_file": "sample.xlsx",
                    "row_label": "Excel 2행",
                    "item_name": "REQ-001",
                    "column": "담당자",
                    "field": "Owner",
                    "raw_value": "=unsafe",
                    "message": "사용자를 찾지 못했습니다.",
                    "action": "입력값을 확인하세요.",
                    "payload_json": "secret payload",
                }
            ]
        )
        summary = {"file_count": 1, "total_rows": 1, "ready_rows": 0, "error_rows": 1}

        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "validation.xlsx"
            result = self.service.export_validation_report(issues, summary, output)

            self.assertEqual(result.row_count, 1)
            workbook = load_workbook(output, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames, ["요약", "문제 목록"])
                sheet = workbook["문제 목록"]
                headers = [cell.value for cell in sheet[1]]
                self.assertNotIn("payload_json", headers)
                self.assertEqual(sheet.cell(row=2, column=headers.index("입력값") + 1).value, "'=unsafe")
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertTrue(sheet.auto_filter.ref)
            finally:
                workbook.close()

    def test_validation_report_rejects_overlong_cell_without_replacing_existing_file(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "validation.xlsx"
            output.write_bytes(b"existing")
            issues = pd.DataFrame([{"severity": "오류", "message": "A" * 32_768}])

            with self.assertRaisesRegex(DeveloperExcelToolError, "길이 제한"):
                self.service.export_validation_report(issues, {}, output)

            self.assertEqual(output.read_bytes(), b"existing")

    def test_tracker_template_expands_table_fields_and_static_options(self) -> None:
        schema_df = pd.DataFrame(
            [
                {
                    "field_name": "Summary",
                    "field_type": "TextField",
                    "tracker_item_field": "name",
                    "mandatory": True,
                    "is_supported": True,
                },
                {
                    "field_name": "Status",
                    "field_label": "상태",
                    "field_type": "OptionChoiceField",
                    "tracker_item_field": "status",
                    "mandatory": False,
                    "is_supported": True,
                    "options": [{"id": 1, "name": "Open"}, {"id": 2, "name": "Closed"}],
                },
                {
                    "field_name": "Description",
                    "field_label": "설명",
                    "field_type": "TextField",
                    "mandatory": False,
                    "mandatory_mode": "conditional",
                    "mandatory_status_names": ["In Progress"],
                    "is_supported": True,
                },
                {
                    "field_name": "Tags",
                    "field_label": "태그",
                    "field_type": "OptionChoiceField",
                    "multiple_values": True,
                    "is_supported": True,
                    "options": [{"id": 10, "name": "One"}, {"id": 11, "name": "Two"}],
                },
                {
                    "field_name": "Steps",
                    "field_label": "절차",
                    "field_type": "TableField",
                    "is_table_field": True,
                    "mandatory": True,
                    "mandatory_mode": "always",
                    "is_supported": True,
                    "table_columns": [
                        {
                            "name": "Action",
                            "mandatory": True,
                            "mandatory_mode": "always",
                        },
                        {
                            "name": "Expected",
                            "mandatory": False,
                            "mandatory_mode": "conditional",
                            "mandatory_status_names": ["Closed", "Verified"],
                        },
                    ],
                },
                {
                    "field_name": "Unsupported",
                    "field_type": "ReferenceField",
                    "is_supported": False,
                },
            ]
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "template.xlsx"
            result = self.service.export_tracker_template(schema_df, output, upload_mode="upsert")

            self.assertEqual(result.column_count, 7)
            workbook = load_workbook(output, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames, ["업로드", "사용 안내", "선택값"])
                self.assertEqual(workbook["선택값"].sheet_state, "hidden")
                headers = [cell.value for cell in workbook["업로드"][1]]
                self.assertEqual(
                    headers,
                    [
                        "id",
                        "Summary",
                        "Status",
                        "Description",
                        "Tags",
                        "Steps.Action",
                        "Steps.Expected",
                    ],
                )
                self.assertNotIn("Unsupported", headers)
                guide = workbook["사용 안내"]
                self.assertEqual(
                    [cell.value for cell in guide[1]],
                    ["컬럼", "필수 기준", "필수 상태", "입력 안내"],
                )
                guide_rows = {
                    str(row[0].value): tuple(cell.value for cell in row[1:])
                    for row in guide.iter_rows(min_row=2)
                    if row[0].value in {"Description", "Steps.Action", "Steps.Expected"}
                }
                self.assertEqual(
                    guide_rows["Description"][:2],
                    ("상태별 필수", "In Progress"),
                )
                self.assertEqual(
                    guide_rows["Steps.Action"][:2],
                    ("항상 필수", "모든 상태"),
                )
                self.assertEqual(
                    guide_rows["Steps.Expected"][:2],
                    ("상태별 필수", "Closed, Verified"),
                )
                validations = workbook["업로드"].data_validations.dataValidation
                self.assertEqual(len(validations), 1)
                self.assertEqual(validations[0].formula1, "CB_OPTIONS_1")
                self.assertIn("CB_OPTIONS_1", workbook.defined_names)
            finally:
                workbook.close()

    def test_tracker_template_update_keeps_id_and_summary_once(self) -> None:
        schema_df = pd.DataFrame(
            [
                {
                    "field_name": "Summary",
                    "field_type": "TextField",
                    "tracker_item_field": "name",
                    "is_supported": True,
                },
                {
                    "field_name": "id",
                    "field_type": "IntegerField",
                    "tracker_item_field": "id",
                    "is_supported": True,
                },
            ]
        )
        before = schema_df.copy(deep=True)

        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "template.xlsx"
            self.service.export_tracker_template(schema_df, output, upload_mode="update")
            workbook = load_workbook(output)
            try:
                self.assertEqual([cell.value for cell in workbook["업로드"][1]], ["id", "Summary"])
            finally:
                workbook.close()

        pd.testing.assert_frame_equal(schema_df, before)

    def test_tracker_template_rejects_unsupported_output_extension(self) -> None:
        schema_df = pd.DataFrame([{"field_name": "Summary", "is_supported": True}])
        with (
            tempfile.TemporaryDirectory() as temp_dir,
            self.assertRaisesRegex(DeveloperExcelToolError, ".xlsx"),
        ):
            self.service.export_tracker_template(schema_df, Path(temp_dir) / "template.xls")

    def test_failed_upload_report_uses_whitelist_and_redacts_sensitive_values(self) -> None:
        failed_df = pd.DataFrame(
            [
                {
                    "source_file": "sample.xlsx",
                    "source_file_path": "/private/work/sample.xlsx",
                    "_row_id": 3,
                    "upload_name": "REQ-003",
                    "phase": "insert",
                    "status": "failed",
                    "error_status_code": 400,
                    "error": '{"message":"Denied","token":"top-secret","details":{"password":"pw-123"}}',
                    "payload_json": {"name": "private payload", "token": "payload-token"},
                    "error_response_json": {"authorization": "Bearer hidden"},
                    "raw_value": "private source value",
                }
            ]
        )
        unresolved_df = pd.DataFrame(
            [
                {
                    "source_file": "sample.xlsx",
                    "_row_id": 4,
                    "upload_name": "REQ-004",
                    "phase": "insert",
                    "status": "unresolved_parent",
                    "error": "A" * 40_000,
                    "payload_json": {"name": "hidden"},
                }
            ]
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "failed.xlsx"
            result = self.service.export_failed_upload_report(
                failed_df,
                unresolved_df,
                output,
            )

            self.assertEqual(result.row_count, 2)
            workbook = load_workbook(output, data_only=False)
            try:
                self.assertEqual(workbook.sheetnames, ["요약", "실패", "미해결"])
                failed_sheet = workbook["실패"]
                headers = [cell.value for cell in failed_sheet[1]]
                for forbidden in (
                    "payload_json",
                    "error_response_json",
                    "source_file_path",
                    "raw_value",
                ):
                    self.assertNotIn(forbidden, headers)
                exported_text = " ".join(
                    str(cell.value or "")
                    for sheet in workbook.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                )
                self.assertNotIn("top-secret", exported_text)
                self.assertNotIn("pw-123", exported_text)
                self.assertNotIn("payload-token", exported_text)
                self.assertNotIn("private source value", exported_text)
                unresolved_headers = [cell.value for cell in workbook["미해결"][1]]
                unresolved_error = workbook["미해결"].cell(
                    row=2,
                    column=unresolved_headers.index("오류 요약") + 1,
                ).value
                self.assertLessEqual(len(str(unresolved_error)), 4_000)
            finally:
                workbook.close()

    def test_failed_upload_report_masks_wrapped_credentials_and_private_locations(self) -> None:
        failed_df = pd.DataFrame(
            [
                {
                    "source_file": "/Users/report-owner/work/sample.xlsx",
                    "_row_id": 3,
                    "upload_name": "report-owner@example.test",
                    "phase": "insert",
                    "status": "failed",
                    "error_status_code": 401,
                    "error": (
                        'HTTP 401 Denied: {"token":"json-secret"}; '
                        "Authorization=Bearer bearer-secret; "
                        "X-Api-Key: header-secret\n"
                        "https://private.example.test/api person@example.test "
                        "/Users/private/work/file.xlsx"
                    ),
                }
            ]
        )

        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "failed.xlsx"
            self.service.export_failed_upload_report(
                failed_df,
                pd.DataFrame(),
                output,
            )
            workbook = load_workbook(output, data_only=False)
            try:
                sheet = workbook["실패"]
                headers = [cell.value for cell in sheet[1]]
                error_text = str(
                    sheet.cell(
                        row=2,
                        column=headers.index("오류 요약") + 1,
                    ).value
                    or ""
                )
                exported_text = " ".join(
                    str(cell.value or "")
                    for row in sheet.iter_rows()
                    for cell in row
                )
            finally:
                workbook.close()

        self.assertIn("HTTP 401 Denied", error_text)
        for secret in (
            "json-secret",
            "bearer-secret",
            "header-secret",
            "private.example",
            "person@example",
            "/Users/private",
            "report-owner",
        ):
            self.assertNotIn(secret, exported_text)
        self.assertIn("<URL>", error_text)
        self.assertIn("<EMAIL>", error_text)
        self.assertIn("<PATH>", error_text)


if __name__ == "__main__":
    unittest.main()

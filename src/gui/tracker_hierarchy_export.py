from __future__ import annotations

import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from .excel_export_style import BODY_FONT
from .excel_export_style import CENTER_WRAP
from .excel_export_style import EXCEL_MAX_CELL_LINE_FEEDS
from .excel_export_style import EXCEL_MAX_CELL_TEXT
from .excel_export_style import GROUP_FILL
from .excel_export_style import HEADER_FILL
from .excel_export_style import HEADER_FONT
from .excel_export_style import LINK_FONT
from .excel_export_style import THIN_BORDER
from .excel_export_style import TOP_WRAP
from .excel_export_style import WHITE_FONT
from .excel_export_style import excel_text_units
from .excel_export_style import normalize_excel_text
from .tracker_baseline_compare import TrackerItemFieldValue
from .tracker_baseline_compare import TrackerTableColumn
from .tracker_baseline_compare import display_tracker_value
from .tracker_baseline_compare import table_field_rows
from .tracker_baseline_compare import tracker_item_fields
from .tracker_hierarchy import TrackerHierarchyError
from .tracker_hierarchy import TrackerHierarchySnapshot
from .tracker_hierarchy import build_tracker_hierarchy
from .tracker_query_models import TrackerItemSummary


EXCEL_MAX_COLUMNS = 16384
EXCEL_MAX_ROWS = 1048576
LONG_VALUE_CELL_TEXT = 30000
LONG_VALUE_CELL_LINE_FEEDS = 200
LONG_VALUE_PREVIEW_TEXT = 180
LONG_VALUE_PREVIEW_LINE_FEEDS = 3
LONG_VALUE_SHEET_TITLE = "긴 값 전체보기"
_FIXED_COLUMN_COUNT = 4
_STRUCTURAL_FIELD_KEYS = {"id", "name", "summary", "parent", "children"}
_INTERNAL_FIELD_KEYS = {"tracker", "project", "childCount", "hasChildren"}


class TrackerHierarchyExportError(RuntimeError):
    pass


@dataclass(frozen=True)
class TrackerHierarchyExportField:
    field_key: str
    label: str
    is_table: bool = False
    table_columns: tuple[TrackerTableColumn, ...] = ()
    default_selected: bool = True
    is_tracker_item_choice: bool = False

    @property
    def column_count(self) -> int:
        return max(len(self.table_columns), 1)


@dataclass(frozen=True)
class TrackerHierarchyExportNode:
    item: TrackerItemSummary
    parent_id: int | None
    depth: int
    fields: tuple[TrackerItemFieldValue, ...]


@dataclass(frozen=True)
class TrackerHierarchyExportSnapshot:
    tracker_id: int
    nodes: tuple[TrackerHierarchyExportNode, ...]
    fields: tuple[TrackerHierarchyExportField, ...]


@dataclass(frozen=True)
class TrackerHierarchyExportSummary:
    item_count: int
    data_row_count: int
    selected_field_count: int
    long_value_count: int = 0
    long_value_part_count: int = 0
    output_path: str = ""


@dataclass(frozen=True)
class _LongValueRecord:
    item_id: int
    item_name: str
    field_label: str
    parts: tuple[str, ...]
    source_coordinate: str
    start_row: int


class _LongValueCollector:
    def __init__(self) -> None:
        self.records: list[_LongValueRecord] = []
        self.part_count = 0

    def add_if_needed(
        self,
        *,
        item_id: int,
        item_name: str,
        field_label: str,
        value: str,
        source_coordinate: str,
    ) -> _LongValueRecord | None:
        normalized = normalize_excel_text(value)
        if not _requires_long_value_sheet(normalized):
            return None
        parts = _split_excel_text(
            normalized,
            max_text_units=LONG_VALUE_CELL_TEXT,
            max_line_feeds=LONG_VALUE_CELL_LINE_FEEDS,
        )
        start_row = 3 + self.part_count
        if start_row + len(parts) - 1 > EXCEL_MAX_ROWS:
            raise TrackerHierarchyExportError(
                f"아이템 #{item_id}의 긴 값을 분할하는 중 Excel 행 한도"
                f"({EXCEL_MAX_ROWS:,}행)를 초과합니다."
            )
        record = _LongValueRecord(
            item_id=item_id,
            item_name=item_name,
            field_label=field_label,
            parts=parts,
            source_coordinate=source_coordinate,
            start_row=start_row,
        )
        self.records.append(record)
        self.part_count += len(parts)
        return record




def hierarchy_export_fields_from_schema(
    tracker_schema: dict[str, Any],
) -> tuple[TrackerHierarchyExportField, ...]:
    """스키마에서 사용자에게 노출할 선택 필드를 안정적인 내부 key로 만든다."""
    raw_fields = tracker_schema.get("fields") if isinstance(tracker_schema, dict) else None
    if not isinstance(raw_fields, list):
        return ()
    fields: list[TrackerHierarchyExportField] = []
    seen: set[str] = set()
    for raw_field in raw_fields:
        if not isinstance(raw_field, dict) or bool(raw_field.get("hidden", False)):
            continue
        builtin_key = str(raw_field.get("trackerItemField") or "").strip()
        if not builtin_key:
            normalized_name = str(raw_field.get("name") or "").strip().casefold().replace(" ", "")
            builtin_key = {
                "id": "id",
                "summary": "name",
                "name": "name",
                "description": "description",
                "status": "status",
                "assignedto": "assignedTo",
                "assignee": "assignedTo",
            }.get(normalized_name, "")
        field_id = _positive_field_id(raw_field)
        field_key = builtin_key or (f"custom:{field_id}" if field_id is not None else "")
        if (
            not field_key
            or field_key in seen
            or field_key in _STRUCTURAL_FIELD_KEYS
            or field_key in _INTERNAL_FIELD_KEYS
        ):
            continue
        seen.add(field_key)
        label = str(raw_field.get("name") or raw_field.get("label") or "").strip()
        if not label:
            label = (
                f"사용자 정의 필드 #{field_id}"
                if field_id is not None
                else "필드"
            )
        type_name = " ".join(
            str(raw_field.get(key) or "") for key in ("type", "valueModel")
        ).casefold()
        is_table = "tablefield" in type_name
        is_tracker_item_choice = "trackeritemchoicefield" in type_name
        fields.append(
            TrackerHierarchyExportField(
                field_key=field_key,
                label=label,
                is_table=is_table,
                table_columns=_schema_table_columns(raw_field) if is_table else (),
                default_selected=not is_table,
                is_tracker_item_choice=is_tracker_item_choice,
            )
        )
    return tuple(fields)


def build_tracker_hierarchy_snapshot(
    items: tuple[TrackerItemSummary, ...] | list[TrackerItemSummary],
    roots: tuple[TrackerItemSummary, ...] | list[TrackerItemSummary],
    tracker_schema: dict[str, Any],
    *,
    tracker_id: int,
) -> TrackerHierarchyExportSnapshot:
    normalized_tracker_id = int(tracker_id)
    try:
        hierarchy = build_tracker_hierarchy(
            items,
            tracker_id=normalized_tracker_id,
            root_ids=(root.item_id for root in roots),
        )
    except TrackerHierarchyError as exc:
        raise TrackerHierarchyExportError(str(exc)) from exc

    return build_tracker_hierarchy_export_snapshot(hierarchy, tracker_schema)


def build_tracker_hierarchy_export_snapshot(
    hierarchy: TrackerHierarchySnapshot,
    tracker_schema: dict[str, Any],
) -> TrackerHierarchyExportSnapshot:
    """이미 조회한 계층에 schema 필드를 결합해 Excel snapshot을 만든다."""
    fields = list(hierarchy_export_fields_from_schema(tracker_schema))
    field_indexes = {field.field_key: index for index, field in enumerate(fields)}
    nodes: list[TrackerHierarchyExportNode] = []
    for hierarchy_node in hierarchy.nodes:
        item = hierarchy_node.item
        normalized_fields = tracker_item_fields(item, tracker_schema=tracker_schema)
        for value in normalized_fields:
            index = field_indexes.get(value.field_key)
            if index is None:
                continue
            current = fields[index]
            merged_columns = _merge_columns(current.table_columns, value.table_columns)
            fields[index] = TrackerHierarchyExportField(
                field_key=current.field_key,
                label=current.label or value.label,
                is_table=current.is_table or value.is_table,
                table_columns=merged_columns,
                default_selected=current.default_selected,
                is_tracker_item_choice=current.is_tracker_item_choice,
            )
        nodes.append(
            TrackerHierarchyExportNode(
                item=item,
                parent_id=hierarchy_node.parent_id,
                depth=hierarchy_node.depth,
                fields=normalized_fields,
            )
        )
    return TrackerHierarchyExportSnapshot(
        tracker_id=hierarchy.tracker_id,
        nodes=tuple(nodes),
        fields=tuple(fields),
    )


def create_tracker_hierarchy_workbook(
    snapshot: TrackerHierarchyExportSnapshot,
    *,
    tracker_name: str,
    project_name: str = "",
    selected_field_keys: tuple[str, ...] | list[str] = (),
    generated_at: datetime | None = None,
    baseline_id: int | None = None,
    baseline_name: str = "",
) -> tuple[Workbook, TrackerHierarchyExportSummary]:
    field_by_key = {field.field_key: field for field in snapshot.fields}
    selected_fields = tuple(
        field_by_key[key]
        for key in dict.fromkeys(str(key) for key in selected_field_keys)
        if key in field_by_key
    )
    total_columns = _FIXED_COLUMN_COUNT + sum(field.column_count for field in selected_fields)
    if total_columns > EXCEL_MAX_COLUMNS:
        raise TrackerHierarchyExportError(
            f"선택한 필드가 Excel 열 한도({EXCEL_MAX_COLUMNS:,}열)를 초과합니다."
        )

    workbook = Workbook()
    info_sheet = workbook.active
    info_sheet.title = "내보내기 정보"
    item_sheet = workbook.create_sheet("트래커 항목")
    generated = generated_at or datetime.now().astimezone()
    long_values = _LongValueCollector()
    data_row_count = _populate_item_sheet(
        item_sheet,
        snapshot=snapshot,
        selected_fields=selected_fields,
        long_values=long_values,
    )
    if long_values.records:
        _populate_long_value_sheet(workbook.create_sheet(LONG_VALUE_SHEET_TITLE), long_values)
    _populate_info_sheet(
        info_sheet,
        snapshot=snapshot,
        tracker_name=tracker_name,
        project_name=project_name,
        selected_fields=selected_fields,
        generated_at=generated,
        data_row_count=data_row_count,
        long_values=long_values,
        baseline_id=baseline_id,
        baseline_name=baseline_name,
    )
    workbook.properties.title = (
        "Codebeamer Baseline 트래커 계층 내보내기"
        if baseline_id is not None
        else "Codebeamer 트래커 계층 내보내기"
    )
    workbook.properties.subject = str(tracker_name or snapshot.tracker_id)
    workbook.properties.creator = "Codebeamer Automation Suite"
    return workbook, TrackerHierarchyExportSummary(
        item_count=len(snapshot.nodes),
        data_row_count=data_row_count,
        selected_field_count=len(selected_fields),
        long_value_count=len(long_values.records),
        long_value_part_count=long_values.part_count,
    )


def export_tracker_hierarchy_xlsx(
    snapshot: TrackerHierarchyExportSnapshot,
    output_path: str | Path,
    *,
    tracker_name: str,
    project_name: str = "",
    selected_field_keys: tuple[str, ...] | list[str] = (),
    baseline_id: int | None = None,
    baseline_name: str = "",
) -> TrackerHierarchyExportSummary:
    path = Path(output_path).expanduser()
    if path.suffix.casefold() != ".xlsx":
        path = path.with_suffix(".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook, summary = create_tracker_hierarchy_workbook(
        snapshot,
        tracker_name=tracker_name,
        project_name=project_name,
        selected_field_keys=selected_field_keys,
        baseline_id=baseline_id,
        baseline_name=baseline_name,
    )
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{path.stem}-",
            suffix=".xlsx",
            dir=path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        temporary_path.replace(path)
    except Exception as exc:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
        raise TrackerHierarchyExportError("Excel 파일을 저장하지 못했습니다.") from exc
    return TrackerHierarchyExportSummary(
        item_count=summary.item_count,
        data_row_count=summary.data_row_count,
        selected_field_count=summary.selected_field_count,
        long_value_count=summary.long_value_count,
        long_value_part_count=summary.long_value_part_count,
        output_path=str(path),
    )


def _tracker_item_choice_names(value: Any) -> str:
    """TrackerItemChoiceField 참조에서 이름만 간결하게 표시한다."""
    candidates = value if isinstance(value, (list, tuple)) else (value,)
    names: list[str] = []
    for candidate in candidates:
        if not isinstance(candidate, dict):
            continue
        name = str(candidate.get("name") or "").strip()
        if name:
            names.append(name)
    return "\n".join(names) if names else "-"


def _populate_item_sheet(
    sheet,
    *,
    snapshot: TrackerHierarchyExportSnapshot,
    selected_fields: tuple[TrackerHierarchyExportField, ...],
    long_values: _LongValueCollector,
) -> int:
    fixed_headers = ("ID", "Summary", "계층 단계", "상위 아이템 ID")
    for column, label in enumerate(fixed_headers, start=1):
        sheet.merge_cells(start_row=1, start_column=column, end_row=2, end_column=column)
        _style_header(sheet.cell(1, column, label), group=True)

    flattened = _flattened_fields(selected_fields)
    column = _FIXED_COLUMN_COUNT + 1
    for field in selected_fields:
        width = field.column_count
        if field.is_table and field.table_columns:
            sheet.merge_cells(
                start_row=1,
                start_column=column,
                end_row=1,
                end_column=column + width - 1,
            )
            _style_header(sheet.cell(1, column, field.label), group=True)
            for offset, table_column in enumerate(field.table_columns):
                _style_header(sheet.cell(2, column + offset, table_column.label))
        else:
            sheet.merge_cells(start_row=1, start_column=column, end_row=2, end_column=column)
            _style_header(sheet.cell(1, column, field.label), group=True)
        column += width

    row_index = 3
    for node in snapshot.nodes:
        values_by_key = {field.field_key: field for field in node.fields}
        row_count = _node_row_count(values_by_key, selected_fields)
        if row_index + row_count - 1 > EXCEL_MAX_ROWS:
            raise TrackerHierarchyExportError(
                f"아이템 #{node.item.item_id}에서 Excel 행 한도({EXCEL_MAX_ROWS:,}행)를 초과합니다."
            )
        start_row = row_index
        end_row = row_index + row_count - 1
        _merge_value(sheet, start_row, end_row, 1, node.item.item_id, context="아이템 ID")
        _merge_value(sheet, start_row, end_row, 2, node.item.name, context=f"#{node.item.item_id} Summary")
        summary_cell = sheet.cell(start_row, 2)
        summary_cell.alignment = Alignment(
            vertical="top",
            wrap_text=True,
            indent=min(max(node.depth, 0), 250),
        )
        _merge_value(sheet, start_row, end_row, 3, node.depth, context="계층 단계")
        _merge_value(
            sheet,
            start_row,
            end_row,
            4,
            node.parent_id,
            context="상위 아이템 ID",
        )

        for offset, (field, table_column) in enumerate(flattened, start=_FIXED_COLUMN_COUNT + 1):
            value = values_by_key.get(field.field_key)
            if value is None:
                _merge_value(sheet, start_row, end_row, offset, "-", context=field.label)
                continue
            if not field.is_table or table_column is None:
                text = (
                    _tracker_item_choice_names(value.value)
                    if field.is_tracker_item_choice
                    else value.display_text()
                )
                _write_scalar_value(
                    sheet,
                    start_row=start_row,
                    end_row=end_row,
                    column=offset,
                    item=node.item,
                    field_label=field.label,
                    value=text,
                    long_values=long_values,
                )
                continue
            rows = table_field_rows(value.value)
            for row_offset in range(row_count):
                cell_value = (
                    rows[row_offset].get(table_column.column_key)
                    if row_offset < len(rows)
                    else None
                )
                text = display_tracker_value(cell_value)
                _write_scalar_value(
                    sheet,
                    start_row=start_row + row_offset,
                    end_row=start_row + row_offset,
                    column=offset,
                    item=node.item,
                    field_label=f"{field.label}.{table_column.label}",
                    value=text,
                    long_values=long_values,
                )

        for row in range(start_row, end_row + 1):
            sheet.row_dimensions[row].outlineLevel = min(max(node.depth, 0), 7)
            sheet.row_dimensions[row].collapsed = False
            for target_column in range(1, len(flattened) + _FIXED_COLUMN_COUNT + 1):
                cell = sheet.cell(row, target_column)
                if isinstance(cell, Cell):
                    cell.border = THIN_BORDER
                    if cell.alignment is None or cell.alignment.indent == 0:
                        cell.alignment = TOP_WRAP
        row_index = end_row + 1

    widths = {1: 13, 2: 42, 3: 12, 4: 18}
    for target_column, width in widths.items():
        sheet.column_dimensions[get_column_letter(target_column)].width = width
    for target_column in range(_FIXED_COLUMN_COUNT + 1, len(flattened) + _FIXED_COLUMN_COUNT + 1):
        sheet.column_dimensions[get_column_letter(target_column)].width = 24
    sheet.freeze_panes = "E3"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.outlinePr.summaryBelow = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:2"
    return max(row_index - 3, 0)


def _populate_info_sheet(
    sheet,
    *,
    snapshot: TrackerHierarchyExportSnapshot,
    tracker_name: str,
    project_name: str,
    selected_fields: tuple[TrackerHierarchyExportField, ...],
    generated_at: datetime,
    data_row_count: int,
    long_values: _LongValueCollector,
    baseline_id: int | None,
    baseline_name: str,
) -> None:
    rows = [
        ("항목", "값"),
        ("프로젝트", project_name or "-"),
        ("트래커", tracker_name or f"ID {snapshot.tracker_id}"),
        ("트래커 ID", snapshot.tracker_id),
    ]
    if baseline_id is not None:
        rows.extend(
            (
                ("조회 기준", "Baseline"),
                ("Baseline", baseline_name or f"Baseline #{baseline_id}"),
                ("Baseline ID", baseline_id),
            )
        )
    else:
        rows.append(("조회 기준", "현재 상태"))
    rows.extend((
        ("생성 시각", generated_at.isoformat(timespec="seconds")),
        ("아이템 수", len(snapshot.nodes)),
        ("데이터 행 수", data_row_count),
        ("선택 필드 수", len(selected_fields)),
        ("선택 필드", "\n".join(field.label for field in selected_fields) or "고정 열만"),
        ("긴 값 수", len(long_values.records)),
        ("긴 값 분할 행 수", long_values.part_count),
    ))
    for row_index, (label, value) in enumerate(rows, start=1):
        _set_safe_value(sheet.cell(row_index, 1), label, context="내보내기 정보")
        _set_safe_value(sheet.cell(row_index, 2), value, context=f"내보내기 정보 {label}")
        for column in (1, 2):
            sheet.cell(row_index, column).border = THIN_BORDER
            sheet.cell(row_index, column).alignment = TOP_WRAP
        if row_index == 1:
            _style_header(sheet.cell(row_index, 1), group=True)
            _style_header(sheet.cell(row_index, 2), group=True)
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 60
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False


def _populate_long_value_sheet(sheet, collector: _LongValueCollector) -> None:
    headers = ("아이템 ID", "아이템명", "필드", "분할", "전체 값")
    for column, label in enumerate(headers, start=1):
        _style_header(sheet.cell(1, column, label), group=True)
    row_index = 2
    for record in collector.records:
        for part_index, part in enumerate(record.parts, start=1):
            values = (
                record.item_id,
                record.item_name,
                record.field_label,
                f"{part_index}/{len(record.parts)}",
                part,
            )
            for column, value in enumerate(values, start=1):
                _set_safe_value(
                    sheet.cell(row_index, column),
                    value,
                    context=f"#{record.item_id} 긴 값 {part_index}/{len(record.parts)}",
                )
                sheet.cell(row_index, column).border = THIN_BORDER
                sheet.cell(row_index, column).alignment = TOP_WRAP
            if part_index == 1:
                sheet.cell(row_index, 1).hyperlink = f"#'트래커 항목'!{record.source_coordinate}"
                sheet.cell(row_index, 1).font = LINK_FONT
            row_index += 1
    widths = {1: 14, 2: 30, 3: 32, 4: 10, 5: 80}
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.auto_filter.ref = f"A1:E{max(row_index - 1, 1)}"
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False


def _write_scalar_value(
    sheet,
    *,
    start_row: int,
    end_row: int,
    column: int,
    item: TrackerItemSummary,
    field_label: str,
    value: str,
    long_values: _LongValueCollector,
) -> None:
    source = sheet.cell(start_row, column).coordinate
    record = long_values.add_if_needed(
        item_id=item.item_id,
        item_name=item.name,
        field_label=field_label,
        value=value,
        source_coordinate=source,
    )
    display_value = _long_value_preview(value) if record is not None else value
    _merge_value(
        sheet,
        start_row,
        end_row,
        column,
        display_value,
        context=f"#{item.item_id} {field_label}",
    )
    if record is not None:
        cell = sheet.cell(start_row, column)
        cell.hyperlink = f"#'{LONG_VALUE_SHEET_TITLE}'!E{record.start_row}"
        cell.font = LINK_FONT


def _merge_value(
    sheet,
    start_row: int,
    end_row: int,
    column: int,
    value: Any,
    *,
    context: str,
) -> None:
    if end_row > start_row:
        sheet.merge_cells(
            start_row=start_row,
            start_column=column,
            end_row=end_row,
            end_column=column,
        )
    _set_safe_value(sheet.cell(start_row, column), value, context=context)
    sheet.cell(start_row, column).alignment = TOP_WRAP


def _style_header(cell: Cell, *, group: bool = False) -> None:
    cell.fill = GROUP_FILL if group else HEADER_FILL
    cell.font = WHITE_FONT if group else HEADER_FONT
    cell.alignment = CENTER_WRAP
    cell.border = THIN_BORDER


def _flattened_fields(
    fields: tuple[TrackerHierarchyExportField, ...],
) -> tuple[tuple[TrackerHierarchyExportField, TrackerTableColumn | None], ...]:
    flattened: list[tuple[TrackerHierarchyExportField, TrackerTableColumn | None]] = []
    for field in fields:
        if field.is_table and field.table_columns:
            flattened.extend((field, column) for column in field.table_columns)
        else:
            flattened.append((field, None))
    return tuple(flattened)


def _node_row_count(
    values_by_key: dict[str, TrackerItemFieldValue],
    fields: tuple[TrackerHierarchyExportField, ...],
) -> int:
    row_count = 1
    for field in fields:
        if not field.is_table:
            continue
        value = values_by_key.get(field.field_key)
        if value is not None:
            row_count = max(row_count, len(table_field_rows(value.value)))
    return row_count


def _schema_table_columns(raw_field: dict[str, Any]) -> tuple[TrackerTableColumn, ...]:
    raw_columns = raw_field.get("columns")
    if not isinstance(raw_columns, list):
        return ()
    columns: list[TrackerTableColumn] = []
    for index, raw_column in enumerate(raw_columns):
        if not isinstance(raw_column, dict) or bool(raw_column.get("hidden", False)):
            continue
        field_id = _positive_field_id(raw_column)
        name = str(raw_column.get("name") or raw_column.get("label") or "").strip()
        key = f"id:{field_id}" if field_id is not None else f"index:{index}"
        columns.append(TrackerTableColumn(key, name or f"열 {index + 1}"))
    return tuple(columns)


def _merge_columns(
    *groups: tuple[TrackerTableColumn, ...],
) -> tuple[TrackerTableColumn, ...]:
    merged: list[TrackerTableColumn] = []
    seen: set[str] = set()
    for group in groups:
        for column in group:
            if column.column_key in seen:
                continue
            seen.add(column.column_key)
            merged.append(column)
    return tuple(merged)


def _positive_field_id(value: dict[str, Any]) -> int | None:
    return _optional_positive_int(value.get("fieldId") or value.get("id"))


def _optional_positive_int(value: Any) -> int | None:
    if value in (None, "") or isinstance(value, bool):
        return None
    try:
        normalized = int(value)
    except (TypeError, ValueError):
        return None
    return normalized if normalized > 0 else None




def _requires_long_value_sheet(value: str) -> bool:
    return (
        excel_text_units(value) > LONG_VALUE_CELL_TEXT
        or value.count("\n") > LONG_VALUE_CELL_LINE_FEEDS
    )


def _split_excel_text(
    value: str,
    *,
    max_text_units: int,
    max_line_feeds: int,
) -> tuple[str, ...]:
    normalized = normalize_excel_text(value)
    if not normalized:
        return ("",)
    parts: list[str] = []
    start = 0
    while start < len(normalized):
        units = 0
        line_feeds = 0
        cursor = start
        last_break = -1
        while cursor < len(normalized):
            character = normalized[cursor]
            character_units = 2 if ord(character) > 0xFFFF else 1
            character_lines = 1 if character == "\n" else 0
            if units + character_units > max_text_units or line_feeds + character_lines > max_line_feeds:
                break
            units += character_units
            line_feeds += character_lines
            if character.isspace():
                last_break = cursor
            cursor += 1
        if cursor >= len(normalized):
            end = len(normalized)
        elif last_break >= start + max((cursor - start) // 2, 1):
            end = last_break + 1
        else:
            end = max(cursor, start + 1)
        parts.append(normalized[start:end])
        start = end
    return tuple(parts)


def _long_value_preview(value: str) -> str:
    suffix = f"\n\n[전체 내용은 '{LONG_VALUE_SHEET_TITLE}' 시트에서 확인]"
    budget = max(LONG_VALUE_PREVIEW_TEXT - excel_text_units(suffix), 1)
    line_budget = max(LONG_VALUE_PREVIEW_LINE_FEEDS - suffix.count("\n"), 0)
    prefix = _split_excel_text(
        value,
        max_text_units=budget,
        max_line_feeds=line_budget,
    )[0]
    return f"{prefix}{suffix}"


def _set_safe_value(cell: Cell, value: Any, *, context: str) -> None:
    cell.font = BODY_FONT
    if isinstance(value, str):
        normalized = normalize_excel_text(value)
        text_units = excel_text_units(normalized)
        if text_units > EXCEL_MAX_CELL_TEXT:
            raise TrackerHierarchyExportError(
                f"{context} 내용이 Excel 셀 길이 한도({EXCEL_MAX_CELL_TEXT:,}자)를 초과합니다."
            )
        if normalized.count("\n") > EXCEL_MAX_CELL_LINE_FEEDS:
            raise TrackerHierarchyExportError(
                f"{context} 내용이 Excel 셀 줄바꿈 한도"
                f"({EXCEL_MAX_CELL_LINE_FEEDS:,}개)를 초과합니다."
            )
        if normalized.startswith(("=", "+", "@")) or (
            normalized.startswith("-") and normalized != "-"
        ):
            normalized = f"'{normalized}"
        cell.value = normalized
        return
    cell.value = value


__all__ = [
    "TrackerHierarchyExportError",
    "TrackerHierarchyExportField",
    "TrackerHierarchyExportNode",
    "TrackerHierarchyExportSnapshot",
    "TrackerHierarchyExportSummary",
    "build_tracker_hierarchy_export_snapshot",
    "build_tracker_hierarchy_snapshot",
    "create_tracker_hierarchy_workbook",
    "export_tracker_hierarchy_xlsx",
    "hierarchy_export_fields_from_schema",
]

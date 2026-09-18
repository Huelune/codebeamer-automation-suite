from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from src.diagnostics import sanitize_diagnostic_text
from src.diagnostics import sanitize_diagnostic_value

from .developer_excel_tools import DeveloperExcelToolError
from .developer_excel_tools import _atomic_workbook_save
from .developer_excel_tools import _safe_output_value
from .excel_export_style import fit_column_widths


_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
_ERROR_FILL = PatternFill("solid", fgColor="FCE4D6")
_WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
_SUPPORTED_UPLOAD_MODES = {"create", "update", "upsert"}
_REPORT_COLUMN_LABELS = {
    "severity": "구분",
    "source_file": "파일",
    "row_label": "Excel 행",
    "item_name": "아이템명",
    "column": "Excel 컬럼",
    "field": "트래커 필드",
    "raw_value": "입력값",
    "message": "문제 내용",
    "action": "해결 방법",
}
_SUMMARY_LABELS = {
    "file_count": "파일 수",
    "batch_total_rows": "전체 입력 행",
    "total_rows": "검증 대상 아이템",
    "ready_rows": "업로드 가능",
    "error_rows": "오류 아이템",
    "warning_rows": "경고 아이템",
    "config_errors": "설정 오류",
    "config_warnings": "설정 경고",
}
_FAILED_REPORT_COLUMN_LABELS = {
    "source_file": "파일",
    "_row_id": "행 식별자",
    "upload_name": "아이템명",
    "phase": "작업",
    "target_item_id": "대상 아이템 ID",
    "status": "상태",
    "error_status_code": "HTTP 상태",
    "error": "오류 요약",
}
_ILLEGAL_EXCEL_TEXT_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
_FAILURE_MESSAGE_MAX_CHARACTERS = 4_000
_MANDATORY_MODE_LABELS = {
    "always": "항상 필수",
    "conditional": "상태별 필수",
    "never": "선택",
}


@dataclass(frozen=True)
class WorkbookExportResult:
    output_path: str
    sheet_names: tuple[str, ...]
    row_count: int
    column_count: int


def _output_path(value: str | Path) -> Path:
    target = Path(value).expanduser().resolve()
    if target.suffix.lower() != ".xlsx":
        raise DeveloperExcelToolError("출력 파일 확장자는 .xlsx여야 합니다.")
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def _style_header(sheet, row: int = 1) -> None:
    for cell in sheet[row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _option_label(value: Any) -> str:
    if isinstance(value, dict):
        for key in ("name", "label", "title", "id"):
            text = str(value.get(key) or "").strip()
            if text:
                return text
        return ""
    return str(value or "").strip()


def _options(value: Any) -> list[Any]:
    return list(value) if isinstance(value, (list, tuple)) else []


def _truthy(value: Any, *, default: bool = False) -> bool:
    if value is None:
        return default
    if isinstance(value, float) and pd.isna(value):
        return default
    if isinstance(value, str):
        return value.strip().lower() not in {"", "0", "false", "no", "off", "nan"}
    return bool(value)


def _field_label(row: pd.Series | dict[str, Any]) -> str:
    return str(row.get("field_label") or row.get("field_name") or "").strip()


def _table_columns(row: pd.Series | dict[str, Any]) -> list[dict[str, Any]]:
    raw = row.get("table_columns")
    if not isinstance(raw, list):
        return []
    return [column for column in raw if isinstance(column, dict)]


def _mandatory_metadata(
    value: pd.Series | dict[str, Any],
    *,
    mandatory: bool,
) -> tuple[str, tuple[str, ...]]:
    raw_mode = value.get("mandatory_mode")
    mode = str(raw_mode or "").strip().casefold() if isinstance(raw_mode, str) else ""
    if mode not in _MANDATORY_MODE_LABELS:
        mode = "always" if mandatory else "never"
    raw_status_names = value.get("mandatory_status_names")
    status_names = tuple(
        dict.fromkeys(
            str(name).strip()
            for name in (
                raw_status_names
                if isinstance(raw_status_names, (list, tuple, set))
                else ()
            )
            if str(name).strip()
        )
    )
    return mode, status_names


def _mandatory_status_label(mode: str, status_names: tuple[str, ...]) -> str:
    if mode == "conditional":
        return ", ".join(status_names) if status_names else "상태 정보 없음"
    if mode == "always":
        return "모든 상태"
    return "-"


def _safe_failure_message(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    normalized: Any = value
    if isinstance(value, str):
        try:
            normalized = json.loads(value)
        except (TypeError, ValueError, json.JSONDecodeError):
            normalized = value
    normalized = sanitize_diagnostic_value(normalized)
    if isinstance(normalized, (dict, list)):
        text = json.dumps(normalized, ensure_ascii=False, sort_keys=True)
    else:
        text = sanitize_diagnostic_text(
            normalized,
            limit=_FAILURE_MESSAGE_MAX_CHARACTERS,
        )
    text = sanitize_diagnostic_text(
        text,
        limit=_FAILURE_MESSAGE_MAX_CHARACTERS,
    )
    text = _ILLEGAL_EXCEL_TEXT_RE.sub(" ", text)
    if len(text) > _FAILURE_MESSAGE_MAX_CHARACTERS:
        text = text[: _FAILURE_MESSAGE_MAX_CHARACTERS - 13] + " ... [생략됨]"
    return text


def _failure_display_value(column: str, value: Any) -> Any:
    if column == "error":
        return _safe_failure_message(value)
    if value is None or (not isinstance(value, (dict, list, tuple)) and pd.isna(value)):
        return ""
    text = str(value).strip()
    if column == "phase":
        text = {"insert": "생성", "create": "생성", "update": "수정"}.get(
            text.lower(), text
        )
    elif column == "status":
        text = {
            "failed": "실패",
            "unresolved_parent": "상위 항목 미해결",
            "payload_failed": "데이터 변환 실패",
        }.get(text.lower(), text)
    text = sanitize_diagnostic_text(
        text,
        limit=_FAILURE_MESSAGE_MAX_CHARACTERS,
    )
    text = _ILLEGAL_EXCEL_TEXT_RE.sub(" ", text)
    if len(text) > _FAILURE_MESSAGE_MAX_CHARACTERS:
        text = text[: _FAILURE_MESSAGE_MAX_CHARACTERS - 13] + " ... [생략됨]"
    return text


class UploadWorkbookService:
    """사용자가 검증 결과를 공유하고 안전한 업로드 양식을 만들 수 있게 한다."""

    def export_validation_report(
        self,
        issue_df: pd.DataFrame | None,
        summary_stats: dict[str, Any] | None,
        output_path: str | Path,
    ) -> WorkbookExportResult:
        target = _output_path(output_path)
        issues = issue_df.copy() if isinstance(issue_df, pd.DataFrame) else pd.DataFrame()
        summary = dict(summary_stats or {})

        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "요약"
        summary_sheet.append(["검증 항목", "결과"])
        for key, label in _SUMMARY_LABELS.items():
            if key in summary:
                summary_sheet.append([label, _safe_output_value(summary[key], cell=label)])
        summary_sheet.append(["문제 목록 수", len(issues.index)])
        _style_header(summary_sheet)
        summary_sheet.freeze_panes = "A2"
        fit_column_widths(summary_sheet)

        issue_sheet = workbook.create_sheet("문제 목록")
        visible_columns = [column for column in _REPORT_COLUMN_LABELS if column in issues.columns]
        if not visible_columns:
            visible_columns = list(_REPORT_COLUMN_LABELS)
        issue_sheet.append([_REPORT_COLUMN_LABELS[column] for column in visible_columns])
        for row_number, (_, row) in enumerate(issues.iterrows(), start=2):
            for column_number, column in enumerate(visible_columns, start=1):
                cell_name = f"{get_column_letter(column_number)}{row_number}"
                issue_sheet.cell(
                    row=row_number,
                    column=column_number,
                    value=_safe_output_value(row.get(column), cell=cell_name),
                )
        _style_header(issue_sheet)
        issue_sheet.freeze_panes = "A2"
        issue_sheet.auto_filter.ref = issue_sheet.dimensions
        issue_sheet.sheet_view.showGridLines = False
        for row in issue_sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if "severity" in visible_columns and issue_sheet.max_row >= 2:
            severity_column = get_column_letter(visible_columns.index("severity") + 1)
            data_range = f"A2:{get_column_letter(issue_sheet.max_column)}{issue_sheet.max_row}"
            issue_sheet.conditional_formatting.add(
                data_range,
                FormulaRule(formula=[f'${severity_column}2="오류"'], fill=_ERROR_FILL),
            )
            issue_sheet.conditional_formatting.add(
                data_range,
                FormulaRule(formula=[f'${severity_column}2="경고"'], fill=_WARNING_FILL),
            )
        fit_column_widths(issue_sheet)

        try:
            _atomic_workbook_save(workbook, target)
        finally:
            workbook.close()
        return WorkbookExportResult(
            output_path=str(target),
            sheet_names=("요약", "문제 목록"),
            row_count=len(issues.index),
            column_count=len(visible_columns),
        )

    def export_tracker_template(
        self,
        schema_df: pd.DataFrame,
        output_path: str | Path,
        *,
        upload_mode: str = "create",
        summary_column: str = "Summary",
    ) -> WorkbookExportResult:
        mode = str(upload_mode or "").strip().lower()
        if mode not in _SUPPORTED_UPLOAD_MODES:
            raise DeveloperExcelToolError("업로드 방식은 create, update, upsert 중 하나여야 합니다.")
        if not isinstance(schema_df, pd.DataFrame) or schema_df.empty:
            raise DeveloperExcelToolError("트래커 필드 정보가 없습니다.")

        target = _output_path(output_path)
        headers: list[str] = []
        header_meta: list[dict[str, Any]] = []
        option_lists: list[tuple[str, list[str]]] = []

        if mode in {"update", "upsert"}:
            headers.append("id")
            header_meta.append(
                {
                    "mandatory": mode == "update",
                    "mandatory_mode": "always" if mode == "update" else "never",
                    "mandatory_status_names": (),
                    "description": (
                        "수정할 아이템 ID"
                        if mode == "update"
                        else "기존 아이템 수정 행은 ID 입력, 새 아이템 생성 행은 비움"
                    ),
                }
            )
        headers.append(summary_column)
        header_meta.append(
            {
                "mandatory": mode != "update",
                "mandatory_mode": "always" if mode != "update" else "never",
                "mandatory_status_names": (),
                "description": "아이템명",
            }
        )

        normalized_summary = summary_column.casefold()
        for _, row in schema_df.iterrows():
            field_name = str(row.get("field_name") or "").strip()
            tracker_item_field = str(row.get("tracker_item_field") or "").strip().casefold()
            if not field_name or field_name.casefold() == normalized_summary:
                continue
            if tracker_item_field in {"id", "name"}:
                continue
            if not _truthy(row.get("is_supported"), default=True):
                continue

            field_label = _field_label(row) or field_name
            mandatory = _truthy(row.get("mandatory"), default=False)
            mandatory_mode, mandatory_status_names = _mandatory_metadata(
                row,
                mandatory=mandatory,
            )
            is_table = _truthy(row.get("is_table_field"), default=False) or str(
                row.get("field_type") or ""
            ) == "TableField"
            if is_table:
                for column in _table_columns(row):
                    column_name = str(column.get("name") or column.get("label") or "").strip()
                    if not column_name:
                        continue
                    header = f"{field_name}.{column_name}"
                    if header in headers:
                        continue
                    headers.append(header)
                    column_mandatory = _truthy(column.get("mandatory"), default=False)
                    column_mode, column_status_names = _mandatory_metadata(
                        column,
                        mandatory=column_mandatory,
                    )
                    header_meta.append(
                        {
                            "mandatory": column_mandatory,
                            "mandatory_mode": column_mode,
                            "mandatory_status_names": column_status_names,
                            "description": f"{field_label} 표의 {column_name} 열",
                        }
                    )
                    options = [_option_label(value) for value in _options(column.get("options"))]
                    options = [value for value in options if value]
                    if options and not _truthy(column.get("multipleValues"), default=False):
                        option_lists.append((header, options))
                continue

            if field_name in headers:
                continue
            headers.append(field_name)
            description_parts = [field_label]
            if _truthy(row.get("multiple_values"), default=False):
                description_parts.append("여러 값은 줄바꿈 또는 세미콜론으로 구분")
            header_meta.append(
                {
                    "mandatory": mandatory,
                    "mandatory_mode": mandatory_mode,
                    "mandatory_status_names": mandatory_status_names,
                    "description": " · ".join(description_parts),
                }
            )
            options = [_option_label(value) for value in _options(row.get("options"))]
            options = [value for value in options if value]
            if options and not _truthy(row.get("multiple_values"), default=False):
                option_lists.append((field_name, options))

        workbook = Workbook()
        upload_sheet = workbook.active
        upload_sheet.title = "업로드"
        upload_sheet.append(
            [_safe_output_value(header, cell=f"업로드!{get_column_letter(index)}1") for index, header in enumerate(headers, start=1)]
        )
        _style_header(upload_sheet)
        upload_sheet.freeze_panes = "A2"
        upload_sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        upload_sheet.sheet_view.showGridLines = False
        for index, meta in enumerate(header_meta, start=1):
            cell = upload_sheet.cell(row=1, column=index)
            cell.comment = None
            if meta["mandatory"]:
                cell.fill = PatternFill("solid", fgColor="C65911")
            upload_sheet.column_dimensions[get_column_letter(index)].width = min(
                max(len(headers[index - 1]) + 4, 14),
                42,
            )

        guide_sheet = workbook.create_sheet("사용 안내")
        guide_sheet.append(["컬럼", "필수 기준", "필수 상태", "입력 안내"])
        for header, meta in zip(headers, header_meta, strict=False):
            mandatory_mode = str(meta["mandatory_mode"])
            mandatory_status_names = tuple(meta["mandatory_status_names"])
            guide_sheet.append(
                [
                    _safe_output_value(header, cell=f"사용 안내!A{guide_sheet.max_row + 1}"),
                    _MANDATORY_MODE_LABELS[mandatory_mode],
                    _safe_output_value(
                        _mandatory_status_label(
                            mandatory_mode,
                            mandatory_status_names,
                        ),
                        cell=f"사용 안내!C{guide_sheet.max_row + 1}",
                    ),
                    _safe_output_value(meta["description"], cell=f"사용 안내!D{guide_sheet.max_row + 1}"),
                ]
            )
        guide_sheet.append([])
        guide_sheet.append(["사용 순서", "", "", "파일 단계에서 이 양식을 선택하고, 데이터 불러오기를 실행한 뒤 매핑을 확인하세요."])
        guide_sheet.append(["주의", "", "", "헤더 이름을 바꾸면 자동 매핑이 달라질 수 있습니다. TableField 열은 필드명.열이름 형식을 유지하세요."])
        _style_header(guide_sheet)
        guide_sheet.freeze_panes = "A2"
        for row in guide_sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        fit_column_widths(guide_sheet)

        if option_lists:
            option_sheet = workbook.create_sheet("선택값")
            for list_index, (header, values) in enumerate(option_lists, start=1):
                option_sheet.cell(row=1, column=list_index, value=header)
                for row_index, value in enumerate(dict.fromkeys(values), start=2):
                    option_sheet.cell(
                        row=row_index,
                        column=list_index,
                        value=_safe_output_value(value, cell=f"선택값!{get_column_letter(list_index)}{row_index}"),
                    )
                header_index = headers.index(header) + 1
                option_column = get_column_letter(list_index)
                last_row = max(len(dict.fromkeys(values)) + 1, 2)
                range_name = f"CB_OPTIONS_{list_index}"
                workbook.defined_names.add(
                    DefinedName(
                        range_name,
                        attr_text=f"'선택값'!${option_column}$2:${option_column}${last_row}",
                    )
                )
                validation = DataValidation(
                    type="list",
                    formula1=range_name,
                    allow_blank=True,
                )
                validation.error = "선택값 시트에 있는 값 중 하나를 선택하세요."
                validation.errorTitle = "지원하지 않는 값"
                validation.prompt = "목록에서 값을 선택할 수 있습니다."
                validation.promptTitle = header
                upload_sheet.add_data_validation(validation)
                validation.add(
                    f"{get_column_letter(header_index)}2:{get_column_letter(header_index)}5000"
                )
            _style_header(option_sheet)
            option_sheet.sheet_state = "hidden"

        try:
            _atomic_workbook_save(workbook, target)
            sheet_names = tuple(workbook.sheetnames)
        finally:
            workbook.close()
        return WorkbookExportResult(
            output_path=str(target),
            sheet_names=sheet_names,
            row_count=0,
            column_count=len(headers),
        )

    def export_failed_upload_report(
        self,
        failed_df: pd.DataFrame | None,
        unresolved_df: pd.DataFrame | None,
        output_path: str | Path,
    ) -> WorkbookExportResult:
        """원문 요청/응답과 비밀값을 제외한 사용자용 실패 보고서를 만든다."""
        target = _output_path(output_path)
        failed = (
            failed_df.copy()
            if isinstance(failed_df, pd.DataFrame)
            else pd.DataFrame()
        )
        unresolved = (
            unresolved_df.copy()
            if isinstance(unresolved_df, pd.DataFrame)
            else pd.DataFrame()
        )

        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "요약"
        summary_sheet.append(["업로드 결과", "건수"])
        summary_sheet.append(["실패", len(failed.index)])
        summary_sheet.append(["상위 항목 미해결", len(unresolved.index)])
        summary_sheet.append(["확인 필요 합계", int(len(failed.index) + len(unresolved.index))])
        summary_sheet.append([])
        summary_sheet.append([
            "보안 안내",
            "요청 payload, 서버 응답 원문, 인증값과 원본 입력값은 이 보고서에 포함하지 않습니다.",
        ])
        _style_header(summary_sheet)
        summary_sheet.freeze_panes = "A2"
        fit_column_widths(summary_sheet)

        maximum_column_count = 0
        for sheet_name, frame in (("실패", failed), ("미해결", unresolved)):
            sheet = workbook.create_sheet(sheet_name)
            visible_columns = [
                column
                for column in _FAILED_REPORT_COLUMN_LABELS
                if column in frame.columns
            ]
            if not visible_columns:
                visible_columns = list(_FAILED_REPORT_COLUMN_LABELS)
            maximum_column_count = max(maximum_column_count, len(visible_columns))
            sheet.append(
                [_FAILED_REPORT_COLUMN_LABELS[column] for column in visible_columns]
            )
            for row_number, (_, row) in enumerate(frame.iterrows(), start=2):
                for column_number, column in enumerate(visible_columns, start=1):
                    cell_name = f"{sheet_name}!{get_column_letter(column_number)}{row_number}"
                    sheet.cell(
                        row=row_number,
                        column=column_number,
                        value=_safe_output_value(
                            _failure_display_value(column, row.get(column)),
                            cell=cell_name,
                        ),
                    )
            _style_header(sheet)
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.sheet_view.showGridLines = False
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            fit_column_widths(sheet)

        try:
            _atomic_workbook_save(workbook, target)
        finally:
            workbook.close()
        return WorkbookExportResult(
            output_path=str(target),
            sheet_names=("요약", "실패", "미해결"),
            row_count=int(len(failed.index) + len(unresolved.index)),
            column_count=maximum_column_count,
        )


__all__ = ["UploadWorkbookService", "WorkbookExportResult"]

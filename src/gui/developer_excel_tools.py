from __future__ import annotations

import contextlib
import csv
import os
import re
import tempfile
from dataclasses import dataclass
from dataclasses import field
from datetime import date
from datetime import datetime
from datetime import time
from datetime import timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from .excel_export_style import EXCEL_MAX_CELL_TEXT
from .excel_export_style import fit_column_widths


# 공개된 이름은 유지하고 값은 공용 상수를 단일 출처로 사용한다.
EXCEL_CELL_MAX_CHARACTERS = EXCEL_MAX_CELL_TEXT
_SUPPORTED_INPUT_SUFFIXES = {".xlsx", ".xlsm", ".xls", ".csv"}
_FORMULA_PREFIXES = ("=", "+", "-", "@")
_FORMULA_ERRORS = {
    "#DIV/0!",
    "#N/A",
    "#NAME?",
    "#NULL!",
    "#NUM!",
    "#REF!",
    "#VALUE!",
}
_ILLEGAL_XML_CONTROL_RE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


class DeveloperExcelToolError(RuntimeError):
    pass


@dataclass(frozen=True)
class ExcelInspectionIssue:
    severity: str
    code: str
    message: str
    sheet_name: str = ""
    cell: str = ""


@dataclass(frozen=True)
class ExcelInspectionReport:
    file_path: str
    file_type: str
    sheet_names: tuple[str, ...]
    selected_sheet: str
    row_count: int
    column_count: int
    headers: tuple[str, ...]
    formula_count: int
    merged_range_count: int
    hidden_sheet_count: int
    hidden_row_count: int
    hidden_column_count: int
    issues: tuple[ExcelInspectionIssue, ...] = field(default_factory=tuple)


@dataclass(frozen=True)
class ExcelConversionResult:
    input_path: str
    output_path: str
    sheet_name: str
    row_count: int
    column_count: int
    warnings: tuple[str, ...] = field(default_factory=tuple)


def _validate_input_path(file_path: str | Path) -> Path:
    path = Path(file_path).expanduser().resolve()
    if not path.is_file():
        raise DeveloperExcelToolError(f"파일을 찾을 수 없습니다: {path}")
    if path.suffix.lower() not in _SUPPORTED_INPUT_SUFFIXES:
        raise DeveloperExcelToolError(
            "지원하지 않는 형식입니다. .xlsx, .xlsm, .xls, .csv만 사용할 수 있습니다."
        )
    return path


def _validate_output_path(input_path: Path, output_path: str | Path, suffix: str) -> Path:
    target = Path(output_path).expanduser().resolve()
    if target.suffix.lower() != suffix:
        raise DeveloperExcelToolError(f"출력 파일 확장자는 {suffix}이어야 합니다.")
    if target == input_path:
        raise DeveloperExcelToolError("원본 파일과 같은 경로에는 저장할 수 없습니다.")
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


def _atomic_workbook_save(workbook: Workbook, target: Path) -> None:
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{target.stem}-",
            suffix=target.suffix,
            dir=target.parent,
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        verification = load_workbook(temporary_path, read_only=True, data_only=False)
        verification.close()
        os.replace(temporary_path, target)
        temporary_path = None
    finally:
        if temporary_path is not None:
            with contextlib.suppress(OSError):
                temporary_path.unlink(missing_ok=True)


def _atomic_text_save(text: str, target: Path, *, encoding: str) -> None:
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{target.stem}-",
            suffix=target.suffix,
            dir=target.parent,
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
        temporary_path.write_text(text, encoding=encoding, newline="")
        os.replace(temporary_path, target)
        temporary_path = None
    finally:
        if temporary_path is not None:
            with contextlib.suppress(OSError):
                temporary_path.unlink(missing_ok=True)


def _normalize_header(value: Any, index: int) -> str:
    text = str(value or "").strip()
    return text or f"Unnamed_{index}"


def _safe_output_value(value: Any, *, cell: str) -> Any:
    if value is None or isinstance(
        value,
        (bool, int, float, Decimal, datetime, date, time, timedelta),
    ):
        return value
    text = str(value)
    if len(text) > EXCEL_CELL_MAX_CHARACTERS:
        raise DeveloperExcelToolError(
            f"{cell} 값이 Excel 셀 길이 제한({EXCEL_CELL_MAX_CHARACTERS}자)을 초과합니다."
        )
    if _ILLEGAL_XML_CONTROL_RE.search(text):
        raise DeveloperExcelToolError(
            f"{cell} 값에 Excel에 저장할 수 없는 제어문자가 있습니다."
        )
    if text.lstrip(" \t\r\n").startswith(_FORMULA_PREFIXES):
        return f"'{text}"
    return text


def _read_xls_values(
    path: Path,
    sheet_name: str | None,
) -> tuple[list[list[Any]], str, tuple[str, ...]]:
    """Read the original .xls cell matrix without upload-specific normalization."""
    try:
        import xlwings as xw
    except ImportError as exc:  # pragma: no cover - depends on local installation
        raise DeveloperExcelToolError(
            ".xls 처리는 Microsoft Excel과 xlwings 설치가 필요합니다."
        ) from exc

    try:
        app = xw.App(visible=False, add_book=False)
    except Exception as exc:  # pragma: no cover - depends on local Excel
        raise DeveloperExcelToolError(
            ".xls 처리를 위해 Microsoft Excel을 시작할 수 없습니다. "
            "Excel 설치 상태를 확인하세요."
        ) from exc

    workbook = None
    try:
        app.display_alerts = False
        app.screen_updating = False
        try:
            workbook = app.books.open(str(path), read_only=True, update_links=False)
        except Exception as exc:
            raise DeveloperExcelToolError(
                ".xls 파일을 열 수 없습니다. "
                "파일 형식과 Microsoft Excel 상태를 확인하세요."
            ) from exc

        sheet_names = tuple(sheet.name for sheet in workbook.sheets)
        if not sheet_names:
            raise DeveloperExcelToolError("시트가 없는 Excel 파일입니다.")
        selected = sheet_name or sheet_names[0]
        if selected not in sheet_names:
            raise DeveloperExcelToolError(f"시트를 찾을 수 없습니다: {selected}")

        sheet = workbook.sheets[selected]
        used_range = sheet.used_range
        last_row = max(int(used_range.last_cell.row or 0), 0)
        last_column = max(int(used_range.last_cell.column or 0), 0)
        if last_row <= 0 or last_column <= 0:
            return [], selected, sheet_names

        # Start at A1 so leading blank rows/columns and the absolute header row
        # remain at the same positions in the converted workbook.
        raw_values = sheet.range((1, 1), (last_row, last_column)).options(ndim=2).value
        rows = [list(row) for row in (raw_values or [])]
        if len(rows) < last_row:
            rows.extend([[None] * last_column for _ in range(last_row - len(rows))])
        normalized_rows = []
        for row in rows[:last_row]:
            normalized = list(row[:last_column])
            if len(normalized) < last_column:
                normalized.extend([None] * (last_column - len(normalized)))
            normalized_rows.append(normalized)
        return normalized_rows, selected, sheet_names
    finally:
        if workbook is not None:
            with contextlib.suppress(Exception):
                workbook.close()
        with contextlib.suppress(Exception):
            app.quit()


def _csv_rows(path: Path, *, encoding: str, delimiter: str) -> list[list[Any]]:
    if len(delimiter) != 1:
        raise DeveloperExcelToolError("CSV 구분자는 한 글자여야 합니다.")
    try:
        with path.open("r", encoding=encoding, newline="") as handle:
            return [list(row) for row in csv.reader(handle, delimiter=delimiter)]
    except UnicodeDecodeError as exc:
        raise DeveloperExcelToolError(
            f"CSV를 {encoding} 인코딩으로 읽을 수 없습니다."
        ) from exc


def _select_sheet(workbook, sheet_name: str | None):
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise DeveloperExcelToolError(f"시트를 찾을 수 없습니다: {sheet_name}")
        return workbook[sheet_name]
    if not workbook.sheetnames:
        raise DeveloperExcelToolError("시트가 없는 Excel 파일입니다.")
    return workbook[workbook.sheetnames[0]]


class DeveloperExcelToolService:
    """개발자 도구에서 원본을 바꾸지 않고 Excel을 검사하고 변환한다."""

    def inspect(
        self,
        file_path: str | Path,
        *,
        sheet_name: str | None = None,
        header_row: int = 1,
        summary_column: str = "Summary",
        csv_encoding: str = "utf-8-sig",
        csv_delimiter: str = ",",
    ) -> ExcelInspectionReport:
        path = _validate_input_path(file_path)
        if header_row < 1:
            raise DeveloperExcelToolError("헤더 행은 1 이상이어야 합니다.")
        suffix = path.suffix.lower()
        if suffix == ".csv":
            return self._inspect_csv(
                path,
                header_row=header_row,
                summary_column=summary_column,
                encoding=csv_encoding,
                delimiter=csv_delimiter,
            )
        if suffix == ".xls":
            return self._inspect_xls(
                path,
                sheet_name=sheet_name,
                header_row=header_row,
                summary_column=summary_column,
            )
        return self._inspect_openxml(
            path,
            sheet_name=sheet_name,
            header_row=header_row,
            summary_column=summary_column,
        )

    def _inspect_csv(
        self,
        path: Path,
        *,
        header_row: int,
        summary_column: str,
        encoding: str,
        delimiter: str,
    ) -> ExcelInspectionReport:
        rows = _csv_rows(path, encoding=encoding, delimiter=delimiter)
        values = rows[header_row - 1] if len(rows) >= header_row else []
        headers = tuple(_normalize_header(value, index) for index, value in enumerate(values))
        issues = self._header_issues(headers, summary_column, "CSV")
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                cell = f"{get_column_letter(column_index)}{row_index}"
                issues.extend(self._value_issues(value, "CSV", cell))
        return ExcelInspectionReport(
            file_path=str(path),
            file_type="csv",
            sheet_names=("CSV",),
            selected_sheet="CSV",
            row_count=max(len(rows) - header_row, 0),
            column_count=max((len(row) for row in rows), default=0),
            headers=headers,
            formula_count=0,
            merged_range_count=0,
            hidden_sheet_count=0,
            hidden_row_count=0,
            hidden_column_count=0,
            issues=tuple(issues),
        )

    def _inspect_xls(
        self,
        path: Path,
        *,
        sheet_name: str | None,
        header_row: int,
        summary_column: str,
    ) -> ExcelInspectionReport:
        rows, selected, sheet_names = _read_xls_values(path, sheet_name)
        header_values = rows[header_row - 1] if len(rows) >= header_row else []
        headers = tuple(
            _normalize_header(value, index)
            for index, value in enumerate(header_values)
        )
        issues = self._header_issues(headers, summary_column, selected)
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                cell = f"{get_column_letter(column_index)}{row_index}"
                issues.extend(self._value_issues(value, selected, cell))
        return ExcelInspectionReport(
            file_path=str(path),
            file_type="xls",
            sheet_names=sheet_names,
            selected_sheet=selected,
            row_count=max(len(rows) - header_row, 0),
            column_count=max((len(row) for row in rows), default=0),
            headers=headers,
            formula_count=0,
            merged_range_count=0,
            hidden_sheet_count=0,
            hidden_row_count=0,
            hidden_column_count=0,
            issues=tuple(issues),
        )

    def _inspect_openxml(
        self,
        path: Path,
        *,
        sheet_name: str | None,
        header_row: int,
        summary_column: str,
    ) -> ExcelInspectionReport:
        workbook = load_workbook(path, read_only=False, data_only=False, keep_vba=False)
        try:
            worksheet = _select_sheet(workbook, sheet_name)
            header_values = [cell.value for cell in worksheet[header_row]]
            headers = tuple(
                _normalize_header(value, index) for index, value in enumerate(header_values)
            )
            issues = self._header_issues(headers, summary_column, worksheet.title)
            formula_count = 0
            for row in worksheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if cell.data_type == "f":
                        formula_count += 1
                    issues.extend(self._value_issues(value, worksheet.title, cell.coordinate))
            if formula_count:
                issues.append(
                    ExcelInspectionIssue(
                        "안내",
                        "FORMULAS_PRESENT",
                        f"수식 셀 {formula_count}개가 있습니다. 값 전용 변환 전에 계산 결과를 확인하세요.",
                        worksheet.title,
                    )
                )
            hidden_sheets = sum(sheet.sheet_state != "visible" for sheet in workbook.worksheets)
            hidden_rows = sum(bool(dimension.hidden) for dimension in worksheet.row_dimensions.values())
            hidden_columns = sum(
                bool(dimension.hidden) for dimension in worksheet.column_dimensions.values()
            )
            merged_count = len(worksheet.merged_cells.ranges)
            if merged_count:
                issues.append(
                    ExcelInspectionIssue(
                        "안내",
                        "MERGED_CELLS",
                        f"병합 범위 {merged_count}개가 있습니다.",
                        worksheet.title,
                    )
                )
            if hidden_sheets or hidden_rows or hidden_columns:
                issues.append(
                    ExcelInspectionIssue(
                        "안내",
                        "HIDDEN_CONTENT",
                        f"숨김 시트 {hidden_sheets}개, 행 {hidden_rows}개, 열 {hidden_columns}개가 있습니다.",
                        worksheet.title,
                    )
                )
            return ExcelInspectionReport(
                file_path=str(path),
                file_type=path.suffix.lower().lstrip("."),
                sheet_names=tuple(workbook.sheetnames),
                selected_sheet=worksheet.title,
                row_count=max(worksheet.max_row - header_row, 0),
                column_count=worksheet.max_column,
                headers=headers,
                formula_count=formula_count,
                merged_range_count=merged_count,
                hidden_sheet_count=hidden_sheets,
                hidden_row_count=hidden_rows,
                hidden_column_count=hidden_columns,
                issues=tuple(issues),
            )
        finally:
            workbook.close()

    @staticmethod
    def _header_issues(
        headers: tuple[str, ...],
        summary_column: str,
        sheet_name: str,
    ) -> list[ExcelInspectionIssue]:
        issues: list[ExcelInspectionIssue] = []
        unnamed = [header for header in headers if header.startswith("Unnamed_")]
        if unnamed:
            issues.append(
                ExcelInspectionIssue(
                    "오류",
                    "EMPTY_HEADERS",
                    f"빈 헤더 {len(unnamed)}개가 있습니다.",
                    sheet_name,
                )
            )
        folded: dict[str, list[str]] = {}
        for header in headers:
            folded.setdefault(header.casefold(), []).append(header)
        duplicates = [values[0] for values in folded.values() if len(values) > 1]
        if duplicates:
            issues.append(
                ExcelInspectionIssue(
                    "오류",
                    "DUPLICATE_HEADERS",
                    f"중복 헤더가 있습니다: {', '.join(duplicates)}",
                    sheet_name,
                )
            )
        if summary_column not in headers:
            issues.append(
                ExcelInspectionIssue(
                    "오류",
                    "SUMMARY_MISSING",
                    f"'{summary_column}' 컬럼을 찾을 수 없습니다.",
                    sheet_name,
                )
            )
        table_headers = [header for header in headers if "." in header]
        if table_headers:
            issues.append(
                ExcelInspectionIssue(
                    "안내",
                    "TABLE_HEADERS",
                    f"TableField 형태 헤더 {len(table_headers)}개를 찾았습니다.",
                    sheet_name,
                )
            )
        return issues

    @staticmethod
    def _value_issues(value: Any, sheet_name: str, cell: str) -> list[ExcelInspectionIssue]:
        if value is None:
            return []
        text = str(value)
        issues: list[ExcelInspectionIssue] = []
        if len(text) > EXCEL_CELL_MAX_CHARACTERS:
            issues.append(
                ExcelInspectionIssue(
                    "오류",
                    "CELL_TOO_LONG",
                    f"셀 값이 {EXCEL_CELL_MAX_CHARACTERS}자를 초과합니다.",
                    sheet_name,
                    cell,
                )
            )
        if _ILLEGAL_XML_CONTROL_RE.search(text):
            issues.append(
                ExcelInspectionIssue(
                    "오류",
                    "ILLEGAL_XML_CONTROL",
                    "Excel에 저장할 수 없는 제어문자가 있습니다.",
                    sheet_name,
                    cell,
                )
            )
        if text in _FORMULA_ERRORS:
            issues.append(
                ExcelInspectionIssue(
                    "오류",
                    "FORMULA_ERROR",
                    f"수식 오류 값 {text}이 있습니다.",
                    sheet_name,
                    cell,
                )
            )
        return issues

    def convert_to_value_xlsx(
        self,
        file_path: str | Path,
        output_path: str | Path,
        *,
        sheet_name: str | None = None,
        header_row: int = 1,
        summary_column: str = "Summary",
        renamed_summary: str | None = None,
        normalize_headers: bool = False,
        csv_encoding: str = "utf-8-sig",
        csv_delimiter: str = ",",
    ) -> ExcelConversionResult:
        path = _validate_input_path(file_path)
        target = _validate_output_path(path, output_path, ".xlsx")
        rows, selected, warnings = self._value_rows(
            path,
            sheet_name=sheet_name,
            header_row=header_row,
            csv_encoding=csv_encoding,
            csv_delimiter=csv_delimiter,
        )
        if not rows:
            raise DeveloperExcelToolError("변환할 데이터가 없습니다.")
        if header_row < 1 or header_row > len(rows):
            raise DeveloperExcelToolError("헤더 행이 데이터 범위를 벗어났습니다.")
        header_values = list(rows[header_row - 1])
        if normalize_headers:
            header_values = [str(value or "").strip() for value in header_values]
        if renamed_summary:
            header_values = [
                renamed_summary if str(value or "").strip() == summary_column else value
                for value in header_values
            ]
        rows[header_row - 1] = header_values

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = selected[:31] or "Data"
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                coordinate = f"{get_column_letter(column_index)}{row_index}"
                worksheet.cell(
                    row=row_index,
                    column=column_index,
                    value=_safe_output_value(value, cell=coordinate),
                )
        header_range = worksheet[header_row]
        fill = PatternFill("solid", fgColor="1F4E78")
        for cell in header_range:
            cell.fill = fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        worksheet.freeze_panes = f"A{header_row + 1}"
        worksheet.auto_filter.ref = worksheet.dimensions
        fit_column_widths(worksheet)
        try:
            _atomic_workbook_save(workbook, target)
        finally:
            workbook.close()
        return ExcelConversionResult(
            input_path=str(path),
            output_path=str(target),
            sheet_name=selected,
            row_count=max(len(rows) - header_row, 0),
            column_count=max((len(row) for row in rows), default=0),
            warnings=tuple(warnings),
        )

    def convert_to_csv(
        self,
        file_path: str | Path,
        output_path: str | Path,
        *,
        sheet_name: str | None = None,
        output_encoding: str = "utf-8-sig",
        delimiter: str = ",",
        input_encoding: str = "utf-8-sig",
        input_delimiter: str = ",",
    ) -> ExcelConversionResult:
        path = _validate_input_path(file_path)
        target = _validate_output_path(path, output_path, ".csv")
        if len(delimiter) != 1:
            raise DeveloperExcelToolError("CSV 구분자는 한 글자여야 합니다.")
        rows, selected, warnings = self._value_rows(
            path,
            sheet_name=sheet_name,
            header_row=1,
            csv_encoding=input_encoding,
            csv_delimiter=input_delimiter,
        )
        output = []
        for row_index, row in enumerate(rows, start=1):
            safe_row = []
            for column_index, value in enumerate(row, start=1):
                coordinate = f"{get_column_letter(column_index)}{row_index}"
                safe_row.append(_safe_output_value(value, cell=coordinate))
            output.append(safe_row)
        from io import StringIO

        buffer = StringIO(newline="")
        writer = csv.writer(buffer, delimiter=delimiter, lineterminator="\n")
        writer.writerows(output)
        _atomic_text_save(buffer.getvalue(), target, encoding=output_encoding)
        return ExcelConversionResult(
            input_path=str(path),
            output_path=str(target),
            sheet_name=selected,
            row_count=max(len(rows) - 1, 0),
            column_count=max((len(row) for row in rows), default=0),
            warnings=tuple(warnings),
        )

    def _value_rows(
        self,
        path: Path,
        *,
        sheet_name: str | None,
        header_row: int,
        csv_encoding: str = "utf-8-sig",
        csv_delimiter: str = ",",
    ) -> tuple[list[list[Any]], str, list[str]]:
        suffix = path.suffix.lower()
        if suffix == ".csv":
            return _csv_rows(path, encoding=csv_encoding, delimiter=csv_delimiter), "CSV", []
        if suffix == ".xls":
            rows, selected, _sheet_names = _read_xls_values(path, sheet_name)
            if header_row > len(rows):
                raise DeveloperExcelToolError("헤더 행이 데이터 범위를 벗어났습니다.")
            return rows, selected, []

        formula_workbook = load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_vba=False,
        )
        value_workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_vba=False,
        )
        warnings: list[str] = []
        try:
            formula_sheet = _select_sheet(formula_workbook, sheet_name)
            value_sheet = value_workbook[formula_sheet.title]
            rows: list[list[Any]] = []
            for formula_row, value_row in zip(
                formula_sheet.iter_rows(),
                value_sheet.iter_rows(), strict=False,
            ):
                output_row: list[Any] = []
                for formula_cell, value_cell in zip(formula_row, value_row, strict=False):
                    if formula_cell.data_type == "f" and value_cell.value is None:
                        raise DeveloperExcelToolError(
                            f"{formula_sheet.title}!{formula_cell.coordinate} 수식의 계산된 값이 없어 값 전용 변환을 중단했습니다. Excel에서 계산 후 저장하세요."
                        )
                    output_row.append(value_cell.value)
                rows.append(output_row)
            if suffix == ".xlsm":
                warnings.append("매크로는 출력 파일에 포함되지 않습니다.")
            return rows, formula_sheet.title, warnings
        finally:
            formula_workbook.close()
            value_workbook.close()


__all__ = [
    "EXCEL_CELL_MAX_CHARACTERS",
    "DeveloperExcelToolError",
    "DeveloperExcelToolService",
    "ExcelConversionResult",
    "ExcelInspectionIssue",
    "ExcelInspectionReport",
]

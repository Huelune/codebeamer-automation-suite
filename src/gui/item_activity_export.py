"""아이템 활동 이력 Excel 내보내기.

화면에서 거른 구간과 유형의 항목을 그대로 받아 요약, 필드별 요약, 이력, 긴 값 전체보기 시트를 만든다.
서버에 다시 묻지 않는다.
"""

from __future__ import annotations

import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.utils import get_column_letter

from .excel_export_style import BODY_FONT
from .excel_export_style import CENTER_WRAP
from .excel_export_style import EXCEL_MAX_CELL_LINE_FEEDS
from .excel_export_style import EXCEL_MAX_CELL_TEXT
from .excel_export_style import GROUP_FILL
from .excel_export_style import HEADER_FILL
from .excel_export_style import HEADER_FONT
from .excel_export_style import LINK_FONT
from .excel_export_style import THIN_BORDER
from .excel_export_style import TOP_WRAP
from .excel_export_style import WHITE_FONT
from .excel_export_style import normalize_excel_text
from .excel_long_value import LONG_VALUE_SHEET_TITLE
from .excel_long_value import cell_preview
from .excel_long_value import long_value_preview
from .excel_long_value import requires_long_value_sheet
from .excel_long_value import split_long_value
from .item_activity import ACTIVITY_KIND_LABELS
from .item_activity import BaselineMarker
from .item_activity import ItemActivityEntry
from .item_activity import ItemActivityKind
from .item_activity import ItemActivitySummary
from .item_activity import is_after_baseline
from .tracker_baseline_compare import display_tracker_value


EXCEL_MAX_ROWS = 1048576
HISTORY_SHEET_TITLE = "이력"
_HISTORY_HEADERS = ("순번", "유형", "버전", "시각", "사용자", "대상", "이전 값", "새 값", "구간")
_HISTORY_WIDTHS = (8, 11, 8, 22, 18, 26, 40, 40, 14)
_VALUE_COLUMNS = (7, 8)
_FIELD_HEADERS = ("필드", "구간 처음 값", "구간 마지막 값", "변경 횟수", "마지막 수정 시각", "마지막 수정자")
_FIELD_WIDTHS = (26, 40, 40, 10, 22, 18)
_LONG_HEADERS = ("순번", "유형", "대상", "칸", "부분", "전체 내용", "이동")
_LONG_WIDTHS = (8, 11, 26, 10, 8, 80, 14)


class ItemActivityExportError(RuntimeError):
    pass


@dataclass(frozen=True)
class ItemActivityExportRequest:
    """화면에서 거른 결과 그대로. 내보내기는 이 값만으로 통합 문서를 만든다."""

    item_id: int
    item_name: str
    tracker_name: str
    scope_label: str
    kind_labels: tuple[str, ...]
    entries: tuple[ItemActivityEntry, ...]
    summary: ItemActivitySummary
    # Baseline 상세 창에서만 둔다. `구간` 칸에 `Baseline 이후`를 적는다.
    baseline: BaselineMarker | None = None


@dataclass(frozen=True)
class ItemActivityExportSummary:
    entry_count: int
    row_count: int
    long_value_count: int
    output_path: str = ""


@dataclass(frozen=True)
class _LongValue:
    number: int
    kind: str
    target: str
    column: str
    parts: tuple[str, ...]
    start_row: int
    source: str


def create_item_activity_workbook(
    request: ItemActivityExportRequest,
    *,
    generated_at: datetime | None = None,
) -> tuple[Workbook, ItemActivityExportSummary]:
    if not request.entries:
        raise ItemActivityExportError("내보낼 이력이 없습니다.")
    rows = _history_rows(request)
    if len(rows) + 1 > EXCEL_MAX_ROWS:
        raise ItemActivityExportError(
            f"아이템 #{request.item_id}의 이력이 Excel 행 한도({EXCEL_MAX_ROWS:,}행)를 넘습니다."
        )
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "요약"
    field_sheet = workbook.create_sheet("필드별 요약")
    history_sheet = workbook.create_sheet(HISTORY_SHEET_TITLE)
    long_values = _populate_history(history_sheet, rows, item_id=request.item_id)
    if long_values:
        _populate_long_values(workbook.create_sheet(LONG_VALUE_SHEET_TITLE), long_values)
    _populate_fields(field_sheet, request.summary)
    _populate_summary(
        summary_sheet,
        request,
        generated_at=generated_at or datetime.now().astimezone(),
        long_value_count=len(long_values),
    )
    workbook.properties.title = "Codebeamer 아이템 변경 이력"
    workbook.properties.subject = f"#{request.item_id}"
    workbook.properties.creator = "Codebeamer Automation Suite"
    return workbook, ItemActivityExportSummary(len(request.entries), len(rows), len(long_values))


def export_item_activity_xlsx(
    request: ItemActivityExportRequest,
    output_path: str | Path,
) -> ItemActivityExportSummary:
    path = Path(output_path).expanduser()
    if path.suffix.casefold() != ".xlsx":
        path = path.with_suffix(".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook, summary = create_item_activity_workbook(request)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{path.stem}-", suffix=".xlsx", dir=path.parent, delete=False
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        temporary_path.replace(path)
    except Exception as exc:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
        raise ItemActivityExportError("Excel 파일을 저장하지 못했습니다.") from exc
    return ItemActivityExportSummary(
        summary.entry_count, summary.row_count, summary.long_value_count, output_path=str(path)
    )


def _history_rows(request: ItemActivityExportRequest) -> list[tuple[Any, ...]]:
    """항목마다 순번을 붙이고, 필드 변경은 바뀐 필드마다 한 행으로 편다."""
    rows: list[tuple[Any, ...]] = []
    for number, entry in enumerate(request.entries, start=1):
        kind = ACTIVITY_KIND_LABELS[entry.kind]
        after = request.baseline is not None and is_after_baseline(entry, request.baseline)
        scope = "Baseline 이후" if after else ""
        version = entry.version if entry.version is not None else ""
        if entry.kind is ItemActivityKind.FIELD_CHANGE and entry.changes:
            for change in entry.changes:
                rows.append(
                    (
                        number, kind, version, entry.at, entry.by, change.field_name,
                        display_tracker_value(change.old_value), display_tracker_value(change.new_value), scope,
                    )
                )
            continue
        if entry.kind is ItemActivityKind.FIELD_CHANGE:
            text = entry.change_summary or "바뀐 필드 정보 없음"
            rows.append((number, kind, version, entry.at, entry.by, "", "", text, scope))
            continue
        text = entry.body
        if entry.attachment_names:
            text += "\n\n첨부: " + ", ".join(entry.attachment_names)
        rows.append((number, kind, "", entry.at, entry.by, entry.title, "", text, scope))
    return rows


def _populate_history(sheet, rows: list[tuple[Any, ...]], *, item_id: int) -> list[_LongValue]:
    _write_headers(sheet, _HISTORY_HEADERS, _HISTORY_WIDTHS)
    long_values: list[_LongValue] = []
    next_long_row = 2
    for offset, row in enumerate(rows):
        row_index = offset + 2
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row_index, column)
            if column in _VALUE_COLUMNS and requires_long_value_sheet(normalize_excel_text(value)):
                parts = split_long_value(value)
                if next_long_row + len(parts) - 1 > EXCEL_MAX_ROWS:
                    raise ItemActivityExportError(
                        f"아이템 #{item_id}의 긴 값을 나누는 중 Excel 행 한도({EXCEL_MAX_ROWS:,}행)를 넘습니다."
                    )
                long_values.append(
                    _LongValue(
                        row[0], row[1], row[5], _HISTORY_HEADERS[column - 1], parts, next_long_row, cell.coordinate
                    )
                )
                _write(cell, long_value_preview(value))
                _link(cell, LONG_VALUE_SHEET_TITLE, f"F{next_long_row}")
                next_long_row += len(parts)
                continue
            _write(cell, value)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(_HISTORY_HEADERS))}{len(rows) + 1}"
    return long_values


def _populate_long_values(sheet, long_values: list[_LongValue]) -> None:
    _write_headers(sheet, _LONG_HEADERS, _LONG_WIDTHS)
    for record in long_values:
        for index, part in enumerate(record.parts):
            row_index = record.start_row + index
            values = (
                record.number, record.kind, record.target, record.column,
                f"{index + 1}/{len(record.parts)}", part, "이력으로 이동" if index == 0 else "",
            )
            for column, value in enumerate(values, start=1):
                _write(sheet.cell(row_index, column), value)
            if index == 0:
                _link(sheet.cell(row_index, 7), HISTORY_SHEET_TITLE, record.source)
    sheet.freeze_panes = "A2"


def _populate_fields(sheet, summary: ItemActivitySummary) -> None:
    _write_headers(sheet, _FIELD_HEADERS, _FIELD_WIDTHS)
    for row_index, field in enumerate(summary.fields, start=2):
        values = (
            field.field_name,
            display_tracker_value(field.first_value),
            display_tracker_value(field.last_value),
            field.change_count,
            field.last_at,
            field.last_by,
        )
        for column, value in enumerate(values, start=1):
            _write(sheet.cell(row_index, column), value)
    sheet.freeze_panes = "A2"
    if summary.fields:
        sheet.auto_filter.ref = f"A1:F{len(summary.fields) + 1}"


def _populate_summary(
    sheet,
    request: ItemActivityExportRequest,
    *,
    generated_at: datetime,
    long_value_count: int,
) -> None:
    summary = request.summary
    rows = (
        ("아이템", f"#{request.item_id} {request.item_name}".strip()),
        ("트래커", request.tracker_name or "-"),
        ("구간", request.scope_label),
        ("유형", ", ".join(request.kind_labels)),
        ("생성 시각", generated_at.isoformat(timespec="seconds")),
        ("이력 항목", len(request.entries)),
        ("버전", summary.version_count),
        ("필드 변경", summary.field_change_count),
        ("바뀐 필드", summary.changed_field_count),
        ("필드 변경 정보 없는 버전", summary.versions_without_changes),
        ("댓글", summary.comment_count),
        ("첨부", summary.attachment_count),
        ("긴 값 분할", f"{long_value_count:,}개 값" if long_value_count else "없음"),
    )
    for row_index, (label, value) in enumerate(rows, start=1):
        label_cell = sheet.cell(row_index, 1)
        _write(label_cell, label)
        label_cell.font = HEADER_FONT
        label_cell.fill = HEADER_FILL
        _write(sheet.cell(row_index, 2), value)
        if label == "긴 값 분할" and long_value_count:
            _link(sheet.cell(row_index, 2), LONG_VALUE_SHEET_TITLE, "A1")
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 60
    sheet.sheet_view.showGridLines = False


def _write_headers(sheet, headers: tuple[str, ...], widths: tuple[int, ...]) -> None:
    for column, (label, width) in enumerate(zip(headers, widths, strict=True), start=1):
        cell = sheet.cell(1, column, label)
        cell.fill = GROUP_FILL
        cell.font = WHITE_FONT
        cell.alignment = CENTER_WRAP
        cell.border = THIN_BORDER
        sheet.column_dimensions[get_column_letter(column)].width = width


def _write(cell: Cell, value: Any) -> None:
    """수식으로 읽힐 수 있는 머리글자는 작은따옴표로 막고, 셀 한도를 넘는 글은 줄인다."""
    cell.font = BODY_FONT
    cell.alignment = TOP_WRAP
    cell.border = THIN_BORDER
    if not isinstance(value, str):
        cell.value = value
        return
    text = cell_preview(
        value,
        max_text_units=EXCEL_MAX_CELL_TEXT - 1,
        max_line_feeds=EXCEL_MAX_CELL_LINE_FEEDS,
        suffix="\n…",
    )
    if text.startswith(("=", "+", "@")) or (text.startswith("-") and text != "-"):
        text = f"'{text}"
    cell.value = text or None


def _link(cell: Cell, sheet_title: str, coordinate: str) -> None:
    escaped = sheet_title.replace("'", "''")
    cell.hyperlink = f"#'{escaped}'!{coordinate}"
    cell.font = LINK_FONT

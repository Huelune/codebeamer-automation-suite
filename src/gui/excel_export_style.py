"""Excel 내보내기에서 공유하는 셀 한도, 서식 객체와 텍스트 정규화 helper.

`tracker_hierarchy_export`, `tracker_baseline_export`, `upload_workbook_tools`,
`developer_excel_tools`가 같은 서식과 같은 열 너비 규칙을 쓰도록 한 곳에 모은다.
긴 값을 별도 시트로 분할하는 로직은 계층 내보내기와 Baseline 비교가 서로 다른
행·열 구조를 다루므로 각 모듈에 남겨 둔다.
"""

from __future__ import annotations

from typing import Any

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment
from openpyxl.styles import Border
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Side
from openpyxl.utils import get_column_letter


# Excel 한 셀이 담을 수 있는 최대 문자 수와 줄바꿈 수.
EXCEL_MAX_CELL_TEXT = 32767
EXCEL_MAX_CELL_LINE_FEEDS = 253

EXPORT_FONT_NAME = "Arial Unicode MS"

GROUP_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")

WHITE_FONT = Font(name=EXPORT_FONT_NAME, color="FFFFFF", bold=True)
HEADER_FONT = Font(name=EXPORT_FONT_NAME, color="1F2937", bold=True)
BODY_FONT = Font(name=EXPORT_FONT_NAME, color="1F2937")
LINK_FONT = Font(name=EXPORT_FONT_NAME, color="0563C1", underline="single")

_THIN_SIDE = Side(style="thin", color="CBD5E1")
THIN_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)

TOP_WRAP = Alignment(vertical="top", wrap_text=True)
CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)


def normalize_excel_text(value: Any) -> str:
    """openpyxl이 거부하는 제어 문자를 지운 문자열을 만든다."""
    return ILLEGAL_CHARACTERS_RE.sub("", str(value or ""))


def excel_text_units(value: str) -> int:
    """셀 길이 한도와 비교할 때 쓰는 UTF-16 code unit 수를 센다.

    Excel의 셀 길이 한도는 code unit 기준이므로 surrogate pair 문자는 2로 센다.
    """
    return len(value.encode("utf-16-le")) // 2


def fit_column_widths(
    sheet,
    *,
    sample_rows: int = 200,
    minimum: int = 10,
    maximum: int = 60,
) -> None:
    """앞쪽 몇 행의 값 길이를 기준으로 열 너비를 맞춘다.

    전체 행을 훑으면 큰 시트에서 느려지므로 `sample_rows`까지만 본다.
    """
    last_sampled_row = min(sheet.max_row, sample_rows)
    for column_index in range(1, sheet.max_column + 1):
        lengths = [
            len(str(sheet.cell(row=row_index, column=column_index).value or ""))
            for row_index in range(1, last_sampled_row + 1)
        ]
        width = max(lengths, default=8) + 2
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(width, minimum), maximum
        )

from __future__ import annotations

import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from .excel_export_style import BODY_FONT
from .excel_export_style import CENTER_WRAP
from .excel_export_style import EXCEL_MAX_CELL_LINE_FEEDS
from .excel_export_style import EXCEL_MAX_CELL_TEXT
from .excel_export_style import EXPORT_FONT_NAME
from .excel_export_style import GROUP_FILL
from .excel_export_style import HEADER_FILL
from .excel_export_style import HEADER_FONT
from .excel_export_style import LINK_FONT
from .excel_export_style import THIN_BORDER
from .excel_export_style import TOP_WRAP
from .excel_export_style import WHITE_FONT
from .excel_export_style import excel_text_units
from .excel_export_style import normalize_excel_text
from .excel_long_value import LONG_VALUE_CELL_LINE_FEEDS
from .excel_long_value import LONG_VALUE_CELL_TEXT
from .excel_long_value import LONG_VALUE_SHEET_TITLE
from .excel_long_value import cell_preview
from .excel_long_value import long_value_preview
from .excel_long_value import requires_long_value_sheet
from .excel_long_value import split_long_value
from .tracker_baseline_compare import BaselineComparisonKind
from .tracker_baseline_compare import BaselineComparisonResult
from .tracker_baseline_compare import TrackerFieldDifference
from .tracker_baseline_compare import TrackerItemComparison
from .tracker_baseline_compare import TrackerTableColumn
from .tracker_baseline_compare import comparison_value_key
from .tracker_baseline_compare import table_field_rows


EXCEL_MAX_COLUMNS = 16384
EXCEL_MAX_ROWS = 1048576


class BaselineExportError(RuntimeError):
    pass


@dataclass(frozen=True)
class BaselineExportField:
    field_key: str
    label: str
    is_table: bool = False
    table_columns: tuple[TrackerTableColumn, ...] = ()

    @property
    def column_count(self) -> int:
        return max(len(self.table_columns), 1)


@dataclass(frozen=True)
class BaselineExportSummary:
    item_count: int
    data_row_count: int
    selected_field_count: int
    long_value_count: int = 0
    long_value_part_count: int = 0
    output_path: str = ""


@dataclass(frozen=True)
class _LongValueRecord:
    kind: BaselineComparisonKind
    item_id: int
    item_name: str
    field_label: str
    table_row: int | None
    table_column: str
    reference_parts: tuple[str, ...]
    comparison_parts: tuple[str, ...]
    reference_is_long: bool
    comparison_is_long: bool
    reference_source: str
    comparison_source: str
    is_changed: bool
    reference_empty: bool
    comparison_empty: bool
    start_row: int

    @property
    def row_count(self) -> int:
        return max(len(self.reference_parts), len(self.comparison_parts), 1)


@dataclass(frozen=True)
class _LongValueTarget:
    start_row: int
    reference_is_long: bool
    comparison_is_long: bool


class _LongValueCollector:
    def __init__(self) -> None:
        self.records: list[_LongValueRecord] = []
        self.data_row_count = 0
        self.value_count = 0

    def add_if_needed(
        self,
        *,
        kind: BaselineComparisonKind,
        item_id: int,
        item_name: str,
        field_label: str,
        reference_text: str,
        comparison_text: str,
        reference_source: str,
        comparison_source: str,
        is_changed: bool,
        reference_empty: bool,
        comparison_empty: bool,
        table_row: int | None = None,
        table_column: str = "",
    ) -> _LongValueTarget | None:
        normalized_reference = normalize_excel_text(reference_text)
        normalized_comparison = normalize_excel_text(comparison_text)
        reference_is_long = requires_long_value_sheet(normalized_reference)
        comparison_is_long = requires_long_value_sheet(normalized_comparison)
        if not reference_is_long and not comparison_is_long:
            return None

        reference_parts = split_long_value(normalized_reference)
        comparison_parts = split_long_value(normalized_comparison)
        start_row = 3 + self.data_row_count
        row_count = max(len(reference_parts), len(comparison_parts), 1)
        if start_row + row_count - 1 > EXCEL_MAX_ROWS:
            raise BaselineExportError(
                f"아이템 #{item_id}의 긴 값을 분할하는 중 Excel 행 한도"
                f"({EXCEL_MAX_ROWS:,}행)를 초과합니다."
            )
        self.records.append(
            _LongValueRecord(
                kind=kind,
                item_id=item_id,
                item_name=item_name,
                field_label=field_label,
                table_row=table_row,
                table_column=table_column,
                reference_parts=reference_parts,
                comparison_parts=comparison_parts,
                reference_is_long=reference_is_long,
                comparison_is_long=comparison_is_long,
                reference_source=reference_source,
                comparison_source=comparison_source,
                is_changed=is_changed,
                reference_empty=reference_empty,
                comparison_empty=comparison_empty,
                start_row=start_row,
            )
        )
        self.data_row_count += row_count
        self.value_count += int(reference_is_long) + int(comparison_is_long)
        return _LongValueTarget(start_row, reference_is_long, comparison_is_long)


_KIND_LABELS = {
    BaselineComparisonKind.ADDED: "신규",
    BaselineComparisonKind.REMOVED: "삭제",
    BaselineComparisonKind.CHANGED: "변경",
    BaselineComparisonKind.UNCHANGED: "변경 없음",
}

_ADDED_FILL = PatternFill("solid", fgColor="DCFCE7")
_REMOVED_FILL = PatternFill("solid", fgColor="FEE2E2")
_CHANGED_FILL = PatternFill("solid", fgColor="FEF3C7")
_UNCHANGED_FILL = PatternFill("solid", fgColor="ECFDF5")


def baseline_export_fields(
    result: BaselineComparisonResult,
) -> tuple[BaselineExportField, ...]:
    """전체 비교 결과의 필드를 처음 등장한 순서로 합친다."""
    fields: list[BaselineExportField] = []
    indexes: dict[str, int] = {}
    for item in result.items:
        for difference in item.fields:
            existing_index = indexes.get(difference.field_key)
            if existing_index is None:
                indexes[difference.field_key] = len(fields)
                fields.append(
                    BaselineExportField(
                        difference.field_key,
                        difference.label,
                        difference.is_table,
                        difference.table_columns,
                    )
                )
                continue
            existing = fields[existing_index]
            columns = _merge_columns(existing.table_columns, difference.table_columns)
            fields[existing_index] = BaselineExportField(
                existing.field_key,
                existing.label or difference.label,
                existing.is_table or difference.is_table,
                columns,
            )
    return tuple(fields)


def create_baseline_comparison_workbook(
    result: BaselineComparisonResult,
    *,
    tracker_name: str,
    reference_label: str,
    comparison_label: str,
    selected_field_keys: tuple[str, ...] | list[str],
    generated_at: datetime | None = None,
) -> tuple[Workbook, BaselineExportSummary]:
    fields_by_key = {field.field_key: field for field in baseline_export_fields(result)}
    selected_fields = tuple(
        fields_by_key[key]
        for key in dict.fromkeys(str(key) for key in selected_field_keys)
        if key in fields_by_key
    )
    if not selected_fields:
        raise BaselineExportError("내보낼 필드를 하나 이상 선택하세요.")

    flattened_count = sum(field.column_count for field in selected_fields)
    total_columns = 3 + flattened_count * 2
    if total_columns > EXCEL_MAX_COLUMNS:
        raise BaselineExportError(
            f"선택한 필드가 Excel 열 한도({EXCEL_MAX_COLUMNS:,}열)를 초과합니다."
        )

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "요약"
    comparison_sheet = workbook.create_sheet("비교 결과")
    generated = generated_at or datetime.now().astimezone()

    filtered_items = tuple(result.items)
    counts = _selected_counts(filtered_items, selected_fields)
    long_values = _LongValueCollector()
    data_row_count = _populate_comparison_sheet(
        comparison_sheet,
        items=filtered_items,
        selected_fields=selected_fields,
        reference_label=reference_label,
        comparison_label=comparison_label,
        long_values=long_values,
    )
    if long_values.records:
        long_value_sheet = workbook.create_sheet(LONG_VALUE_SHEET_TITLE)
        _populate_long_value_sheet(
            long_value_sheet,
            records=tuple(long_values.records),
            reference_label=reference_label,
            comparison_label=comparison_label,
        )
    _populate_summary_sheet(
        summary_sheet,
        tracker_name=tracker_name,
        reference_label=reference_label,
        comparison_label=comparison_label,
        generated_at=generated,
        selected_fields=selected_fields,
        item_count=len(filtered_items),
        counts=counts,
        long_value_count=long_values.value_count,
        long_value_part_count=long_values.data_row_count,
    )
    workbook.properties.title = "Codebeamer Baseline 비교 결과"
    workbook.properties.subject = str(tracker_name or "Tracker")
    workbook.properties.creator = "Codebeamer Automation Suite"
    return workbook, BaselineExportSummary(
        item_count=len(filtered_items),
        data_row_count=data_row_count,
        selected_field_count=len(selected_fields),
        long_value_count=long_values.value_count,
        long_value_part_count=long_values.data_row_count,
    )


def export_baseline_comparison_xlsx(
    result: BaselineComparisonResult,
    output_path: str | Path,
    *,
    tracker_name: str,
    reference_label: str,
    comparison_label: str,
    selected_field_keys: tuple[str, ...] | list[str],
) -> BaselineExportSummary:
    path = Path(output_path).expanduser()
    if path.suffix.casefold() != ".xlsx":
        path = path.with_suffix(".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook, summary = create_baseline_comparison_workbook(
        result,
        tracker_name=tracker_name,
        reference_label=reference_label,
        comparison_label=comparison_label,
        selected_field_keys=selected_field_keys,
    )
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{path.stem}-",
            suffix=".xlsx",
            dir=path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        temporary_path.replace(path)
    except Exception as exc:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
        raise BaselineExportError("Excel 파일을 저장하지 못했습니다.") from exc
    return BaselineExportSummary(
        item_count=summary.item_count,
        data_row_count=summary.data_row_count,
        selected_field_count=summary.selected_field_count,
        long_value_count=summary.long_value_count,
        long_value_part_count=summary.long_value_part_count,
        output_path=str(path),
    )


def _merge_columns(
    *groups: tuple[TrackerTableColumn, ...],
) -> tuple[TrackerTableColumn, ...]:
    merged: list[TrackerTableColumn] = []
    seen: set[str] = set()
    for group in groups:
        for column in group:
            if column.column_key in seen:
                continue
            seen.add(column.column_key)
            merged.append(column)
    return tuple(merged)


def _selected_item_kind(
    item: TrackerItemComparison,
    selected_keys: set[str],
) -> BaselineComparisonKind:
    if item.kind in {BaselineComparisonKind.ADDED, BaselineComparisonKind.REMOVED}:
        return item.kind
    return (
        BaselineComparisonKind.CHANGED
        if any(
            field.field_key in selected_keys and field.is_changed
            for field in item.fields
        )
        else BaselineComparisonKind.UNCHANGED
    )


def _selected_counts(
    items: tuple[TrackerItemComparison, ...],
    fields: tuple[BaselineExportField, ...],
) -> dict[BaselineComparisonKind, int]:
    selected_keys = {field.field_key for field in fields}
    return {
        kind: sum(_selected_item_kind(item, selected_keys) == kind for item in items)
        for kind in BaselineComparisonKind
    }


def _populate_summary_sheet(
    sheet,
    *,
    tracker_name: str,
    reference_label: str,
    comparison_label: str,
    generated_at: datetime,
    selected_fields: tuple[BaselineExportField, ...],
    item_count: int,
    counts: dict[BaselineComparisonKind, int],
    long_value_count: int,
    long_value_part_count: int,
) -> None:
    sheet["A1"] = "Baseline 비교 내보내기"
    sheet["A1"].font = Font(
        name=EXPORT_FONT_NAME,
        size=16,
        bold=True,
        color="FFFFFF",
    )
    sheet["A1"].fill = GROUP_FILL
    sheet.merge_cells("A1:B1")
    rows = (
        ("트래커", tracker_name),
        ("기준", reference_label),
        ("비교", comparison_label),
        ("생성 시각", generated_at.isoformat(timespec="seconds")),
        (
            "선택 필드",
            cell_preview(
                "\n".join(field.label for field in selected_fields),
                max_text_units=LONG_VALUE_CELL_TEXT,
                max_line_feeds=LONG_VALUE_CELL_LINE_FEEDS,
                suffix=f"\n… 총 {len(selected_fields):,}개 필드",
            ),
        ),
        ("전체 아이템", item_count),
        ("신규", counts[BaselineComparisonKind.ADDED]),
        ("삭제", counts[BaselineComparisonKind.REMOVED]),
        ("변경", counts[BaselineComparisonKind.CHANGED]),
        ("변경 없음", counts[BaselineComparisonKind.UNCHANGED]),
        (
            "긴 값 분할",
            (
                f"{long_value_count:,}개 값 · {long_value_part_count:,}개 분할 행"
                if long_value_count
                else "없음"
            ),
        ),
    )
    for row_index, (label, value) in enumerate(rows, start=3):
        sheet.cell(row_index, 1, label)
        _set_safe_value(sheet.cell(row_index, 2), value, context=label)
        if label == "긴 값 분할" and long_value_count:
            _set_internal_link(
                sheet.cell(row_index, 2),
                LONG_VALUE_SHEET_TITLE,
                "A1",
            )
        sheet.cell(row_index, 1).font = HEADER_FONT
        sheet.cell(row_index, 1).fill = HEADER_FILL
        for column in (1, 2):
            sheet.cell(row_index, column).border = THIN_BORDER
            sheet.cell(row_index, column).alignment = TOP_WRAP
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 48
    sheet.freeze_panes = "A3"
    sheet.sheet_view.showGridLines = False


def _flattened_headers(
    fields: tuple[BaselineExportField, ...],
) -> tuple[tuple[BaselineExportField, TrackerTableColumn | None, str], ...]:
    headers: list[tuple[BaselineExportField, TrackerTableColumn | None, str]] = []
    for field in fields:
        if field.is_table and field.table_columns:
            for column in field.table_columns:
                headers.append((field, column, f"{field.label}.{column.label}"))
        else:
            headers.append((field, None, field.label))
    return tuple(headers)


def _populate_comparison_sheet(
    sheet,
    *,
    items: tuple[TrackerItemComparison, ...],
    selected_fields: tuple[BaselineExportField, ...],
    reference_label: str,
    comparison_label: str,
    long_values: _LongValueCollector,
) -> int:
    headers = _flattened_headers(selected_fields)
    reference_start = 4
    reference_end = reference_start + len(headers) - 1
    comparison_start = reference_end + 1
    comparison_end = comparison_start + len(headers) - 1

    for column, label in enumerate(("결과", "아이템 ID", "아이템명"), start=1):
        sheet.merge_cells(start_row=1, start_column=column, end_row=2, end_column=column)
        cell = sheet.cell(1, column, label)
        _style_group_header(cell)
    sheet.merge_cells(
        start_row=1,
        start_column=reference_start,
        end_row=1,
        end_column=reference_end,
    )
    sheet.merge_cells(
        start_row=1,
        start_column=comparison_start,
        end_row=1,
        end_column=comparison_end,
    )
    reference_group = sheet.cell(1, reference_start)
    comparison_group = sheet.cell(1, comparison_start)
    _set_safe_value(reference_group, f"기준 · {reference_label}", context="기준 제목")
    _set_safe_value(comparison_group, f"비교 · {comparison_label}", context="비교 제목")
    _style_group_header(reference_group)
    _style_group_header(comparison_group)

    for offset, (_field, _column, label) in enumerate(headers):
        for start in (reference_start, comparison_start):
            cell = sheet.cell(2, start + offset, label)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = CENTER_WRAP
            cell.border = THIN_BORDER

    selected_keys = {field.field_key for field in selected_fields}
    row_index = 3
    for item in items:
        difference_by_key = {field.field_key: field for field in item.fields}
        row_count = _item_row_count(difference_by_key, selected_fields)
        if row_index + row_count - 1 > EXCEL_MAX_ROWS:
            raise BaselineExportError(
                f"아이템 #{item.item_id}에서 Excel 행 한도({EXCEL_MAX_ROWS:,}행)를 초과합니다."
            )
        start_row = row_index
        end_row = row_index + row_count - 1
        kind = _selected_item_kind(item, selected_keys)
        _merge_item_value(sheet, start_row, end_row, 1, _KIND_LABELS[kind])
        _merge_item_value(sheet, start_row, end_row, 2, item.item_id)
        reference_name = item.reference.name if item.reference is not None else ""
        comparison_name = item.comparison.name if item.comparison is not None else ""
        name_target = long_values.add_if_needed(
            kind=kind,
            item_id=item.item_id,
            item_name=item.name,
            field_label="아이템명",
            reference_text=reference_name,
            comparison_text=comparison_name,
            reference_source=sheet.cell(start_row, 3).coordinate,
            comparison_source=sheet.cell(start_row, 3).coordinate,
            is_changed=reference_name != comparison_name,
            reference_empty=item.reference is None,
            comparison_empty=item.comparison is None,
        )
        item_name_value = item.name
        item_name_link_column = 8
        if name_target is not None:
            sheet.row_dimensions[start_row].height = max(
                sheet.row_dimensions[start_row].height or 0,
                72,
            )
            if item.reference is not None and name_target.reference_is_long:
                item_name_value = long_value_preview(reference_name)
                item_name_link_column = 8
            elif name_target.comparison_is_long:
                item_name_value = long_value_preview(comparison_name)
                item_name_link_column = 9
        _merge_item_value(sheet, start_row, end_row, 3, item_name_value)
        if name_target is not None:
            _set_internal_link(
                sheet.cell(start_row, 3),
                LONG_VALUE_SHEET_TITLE,
                name_target.start_row,
                item_name_link_column,
            )
        sheet.cell(start_row, 1).fill = _kind_fill(kind)
        sheet.cell(start_row, 1).font = Font(name=EXPORT_FONT_NAME, bold=True)

        for offset, (export_field, table_column, _label) in enumerate(headers):
            difference = difference_by_key.get(export_field.field_key)
            reference_column = reference_start + offset
            comparison_column = comparison_start + offset
            if difference is None:
                _merge_item_value(sheet, start_row, end_row, reference_column, "-")
                _merge_item_value(sheet, start_row, end_row, comparison_column, "-")
                continue
            if not export_field.is_table or table_column is None:
                reference_text = difference.reference_text()
                comparison_text = difference.comparison_text()
                long_target = long_values.add_if_needed(
                    kind=kind,
                    item_id=item.item_id,
                    item_name=item.name,
                    field_label=difference.label,
                    reference_text=reference_text,
                    comparison_text=comparison_text,
                    reference_source=sheet.cell(start_row, reference_column).coordinate,
                    comparison_source=sheet.cell(start_row, comparison_column).coordinate,
                    is_changed=difference.is_changed,
                    reference_empty=_is_empty(difference.reference),
                    comparison_empty=_is_empty(difference.comparison),
                )
                if long_target is not None:
                    sheet.row_dimensions[start_row].height = max(
                        sheet.row_dimensions[start_row].height or 0,
                        72,
                    )
                _merge_item_value(
                    sheet,
                    start_row,
                    end_row,
                    reference_column,
                    (
                        long_value_preview(reference_text)
                        if long_target is not None and long_target.reference_is_long
                        else reference_text
                    ),
                    context=f"#{item.item_id} {difference.label} 기준",
                )
                _merge_item_value(
                    sheet,
                    start_row,
                    end_row,
                    comparison_column,
                    (
                        long_value_preview(comparison_text)
                        if long_target is not None and long_target.comparison_is_long
                        else comparison_text
                    ),
                    context=f"#{item.item_id} {difference.label} 비교",
                )
                if long_target is not None and long_target.reference_is_long:
                    _set_internal_link(
                        sheet.cell(start_row, reference_column),
                        LONG_VALUE_SHEET_TITLE,
                        long_target.start_row,
                        8,
                    )
                if long_target is not None and long_target.comparison_is_long:
                    _set_internal_link(
                        sheet.cell(start_row, comparison_column),
                        LONG_VALUE_SHEET_TITLE,
                        long_target.start_row,
                        9,
                    )
                _style_scalar_difference(
                    sheet.cell(start_row, reference_column),
                    sheet.cell(start_row, comparison_column),
                    difference,
                )
                continue
            _write_table_column(
                sheet,
                start_row=start_row,
                row_count=row_count,
                reference_column=reference_column,
                comparison_column=comparison_column,
                item_id=item.item_id,
                item_name=item.name,
                kind=kind,
                difference=difference,
                table_column=table_column,
                long_values=long_values,
            )

        for row in range(start_row, end_row + 1):
            for column in range(1, comparison_end + 1):
                cell = sheet.cell(row, column)
                if not isinstance(cell, Cell):
                    continue
                cell.border = THIN_BORDER
                cell.alignment = TOP_WRAP
        row_index = end_row + 1

    sheet.freeze_panes = "D3"
    widths = {1: 13, 2: 13, 3: 32}
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    for column in range(reference_start, comparison_end + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 24
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 36
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:2"
    return max(row_index - 3, 0)


def _style_group_header(cell: Cell) -> None:
    cell.fill = GROUP_FILL
    cell.font = WHITE_FONT
    cell.alignment = CENTER_WRAP
    cell.border = THIN_BORDER


def _item_row_count(
    difference_by_key: dict[str, TrackerFieldDifference],
    selected_fields: tuple[BaselineExportField, ...],
) -> int:
    row_count = 1
    for field in selected_fields:
        if not field.is_table:
            continue
        difference = difference_by_key.get(field.field_key)
        if difference is None:
            continue
        row_count = max(
            row_count,
            len(table_field_rows(difference.reference)),
            len(table_field_rows(difference.comparison)),
        )
    return row_count


def _merge_item_value(
    sheet,
    start_row: int,
    end_row: int,
    column: int,
    value: Any,
    *,
    context: str = "셀",
) -> None:
    if end_row > start_row:
        sheet.merge_cells(
            start_row=start_row,
            start_column=column,
            end_row=end_row,
            end_column=column,
        )
    cell = sheet.cell(start_row, column)
    _set_safe_value(cell, value, context=context)
    cell.alignment = TOP_WRAP
    cell.border = THIN_BORDER


def _write_table_column(
    sheet,
    *,
    start_row: int,
    row_count: int,
    reference_column: int,
    comparison_column: int,
    item_id: int,
    item_name: str,
    kind: BaselineComparisonKind,
    difference: TrackerFieldDifference,
    table_column: TrackerTableColumn,
    long_values: _LongValueCollector,
) -> None:
    reference_rows = table_field_rows(difference.reference)
    comparison_rows = table_field_rows(difference.comparison)
    missing = object()
    for offset in range(row_count):
        reference_value = (
            reference_rows[offset].get(table_column.column_key, missing)
            if offset < len(reference_rows)
            else missing
        )
        comparison_value = (
            comparison_rows[offset].get(table_column.column_key, missing)
            if offset < len(comparison_rows)
            else missing
        )
        reference_cell = sheet.cell(start_row + offset, reference_column)
        comparison_cell = sheet.cell(start_row + offset, comparison_column)
        reference_text = (
            "" if reference_value is missing else _display_cell_value(reference_value)
        )
        comparison_text = (
            "" if comparison_value is missing else _display_cell_value(comparison_value)
        )
        long_target = long_values.add_if_needed(
            kind=kind,
            item_id=item_id,
            item_name=item_name,
            field_label=difference.label,
            table_row=offset + 1,
            table_column=table_column.label,
            reference_text=reference_text,
            comparison_text=comparison_text,
            reference_source=reference_cell.coordinate,
            comparison_source=comparison_cell.coordinate,
            is_changed=(
                reference_value is missing
                or comparison_value is missing
                or comparison_value_key(reference_value)
                != comparison_value_key(comparison_value)
            ),
            reference_empty=(reference_value is missing or _is_empty(reference_value)),
            comparison_empty=(comparison_value is missing or _is_empty(comparison_value)),
        )
        if long_target is not None:
            target_row = start_row + offset
            sheet.row_dimensions[target_row].height = max(
                sheet.row_dimensions[target_row].height or 0,
                72,
            )
        _set_safe_value(
            reference_cell,
            (
                long_value_preview(reference_text)
                if long_target is not None and long_target.reference_is_long
                else reference_text
            ),
            context=f"#{item_id} {difference.label}.{table_column.label} 기준",
        )
        _set_safe_value(
            comparison_cell,
            (
                long_value_preview(comparison_text)
                if long_target is not None and long_target.comparison_is_long
                else comparison_text
            ),
            context=f"#{item_id} {difference.label}.{table_column.label} 비교",
        )
        if long_target is not None and long_target.reference_is_long:
            _set_internal_link(
                reference_cell,
                LONG_VALUE_SHEET_TITLE,
                long_target.start_row,
                8,
            )
        if long_target is not None and long_target.comparison_is_long:
            _set_internal_link(
                comparison_cell,
                LONG_VALUE_SHEET_TITLE,
                long_target.start_row,
                9,
            )
        reference_cell.alignment = TOP_WRAP
        comparison_cell.alignment = TOP_WRAP
        if reference_value is missing and comparison_value is missing:
            continue
        if comparison_value is missing:
            reference_cell.fill = _ADDED_FILL
        elif reference_value is missing:
            comparison_cell.fill = _REMOVED_FILL
        elif comparison_value_key(reference_value) != comparison_value_key(comparison_value):
            reference_cell.fill = _CHANGED_FILL
            comparison_cell.fill = _CHANGED_FILL


def _display_cell_value(value: Any) -> str:
    if value is None or value == "" or value == [] or value == {}:
        return ""
    if isinstance(value, bool):
        return "예" if value else "아니요"
    if isinstance(value, dict):
        name = value.get("name") or value.get("summary") or value.get("label")
        reference_id = value.get("id")
        if name not in (None, "") and reference_id not in (None, ""):
            return f"{name} (ID {reference_id})"
        if name not in (None, ""):
            return str(name)
        if reference_id not in (None, ""):
            return f"ID {reference_id}"
        return "\n".join(
            f"{key}: {_display_cell_value(nested)}" for key, nested in value.items()
        )
    if isinstance(value, (list, tuple)):
        return "\n".join(f"• {_display_cell_value(item)}" for item in value)
    return str(value)


def _populate_long_value_sheet(
    sheet,
    *,
    records: tuple[_LongValueRecord, ...],
    reference_label: str,
    comparison_label: str,
) -> None:
    sheet.merge_cells("A1:G1")
    title_cell = sheet["A1"]
    title_cell.value = LONG_VALUE_SHEET_TITLE
    _style_group_header(title_cell)
    reference_header = sheet["H1"]
    comparison_header = sheet["I1"]
    _set_safe_value(
        reference_header,
        cell_preview(
            f"기준 · {reference_label}",
            max_text_units=1000,
            max_line_feeds=4,
            suffix="\n…",
        ),
        context="긴 값 기준 제목",
    )
    _set_safe_value(
        comparison_header,
        cell_preview(
            f"비교 · {comparison_label}",
            max_text_units=1000,
            max_line_feeds=4,
            suffix="\n…",
        ),
        context="긴 값 비교 제목",
    )
    _style_group_header(reference_header)
    _style_group_header(comparison_header)
    sheet.merge_cells("J1:J2")
    navigation_header = sheet["J1"]
    navigation_header.value = "이동"
    _style_group_header(navigation_header)

    headers = (
        "결과",
        "아이템 ID",
        "아이템명",
        "필드",
        "Table 행",
        "Table 열",
        "부분",
        "전체 내용",
        "전체 내용",
    )
    for column, label in enumerate(headers, start=1):
        cell = sheet.cell(2, column, label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER_WRAP
        cell.border = THIN_BORDER

    for record in records:
        for offset in range(record.row_count):
            row = record.start_row + offset
            reference_part = (
                record.reference_parts[offset]
                if offset < len(record.reference_parts)
                else ""
            )
            comparison_part = (
                record.comparison_parts[offset]
                if offset < len(record.comparison_parts)
                else ""
            )
            values = (
                _KIND_LABELS[record.kind],
                record.item_id,
                cell_preview(
                    record.item_name,
                    max_text_units=1000,
                    max_line_feeds=12,
                    suffix="\n…",
                ),
                cell_preview(
                    record.field_label,
                    max_text_units=1000,
                    max_line_feeds=12,
                    suffix="\n…",
                ),
                record.table_row or "",
                cell_preview(
                    record.table_column,
                    max_text_units=1000,
                    max_line_feeds=12,
                    suffix="\n…",
                ),
                f"{offset + 1}/{record.row_count}",
                reference_part,
                comparison_part,
                "비교 결과로 이동",
            )
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row, column)
                _set_safe_value(
                    cell,
                    value,
                    context=f"#{record.item_id} 긴 값 {offset + 1}/{record.row_count}",
                )
                cell.alignment = TOP_WRAP
                cell.border = THIN_BORDER
            sheet.cell(row, 1).fill = _kind_fill(record.kind)
            sheet.cell(row, 1).font = Font(name=EXPORT_FONT_NAME, bold=True)
            _style_long_value_pair(
                sheet.cell(row, 8),
                sheet.cell(row, 9),
                record,
            )
            source = (
                record.reference_source
                if record.reference_is_long
                else record.comparison_source
            )
            _set_internal_link(sheet.cell(row, 10), "비교 결과", source)
            sheet.row_dimensions[row].height = 90

    last_row = max((record.start_row + record.row_count - 1 for record in records), default=2)
    sheet.auto_filter.ref = f"A2:J{last_row}"
    sheet.freeze_panes = "A3"
    widths = {
        1: 13,
        2: 13,
        3: 30,
        4: 26,
        5: 10,
        6: 22,
        7: 10,
        8: 60,
        9: 60,
        10: 18,
    }
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 36
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:2"


def _set_internal_link(
    cell: Cell,
    sheet_title: str,
    row_or_coordinate: int | str,
    column: int | None = None,
) -> None:
    if isinstance(row_or_coordinate, int):
        if column is None:
            raise ValueError("내부 링크 열 번호가 필요합니다.")
        coordinate = f"{get_column_letter(column)}{row_or_coordinate}"
    else:
        coordinate = row_or_coordinate
    escaped_title = sheet_title.replace("'", "''")
    cell.hyperlink = f"#'{escaped_title}'!{coordinate}"
    cell.font = LINK_FONT


def _style_long_value_pair(
    reference_cell: Cell,
    comparison_cell: Cell,
    record: _LongValueRecord,
) -> None:
    if not record.is_changed:
        return
    if record.comparison_empty and not record.reference_empty:
        reference_cell.fill = _ADDED_FILL
    elif record.reference_empty and not record.comparison_empty:
        comparison_cell.fill = _REMOVED_FILL
    else:
        reference_cell.fill = _CHANGED_FILL
        comparison_cell.fill = _CHANGED_FILL


def _style_scalar_difference(
    reference_cell: Cell,
    comparison_cell: Cell,
    difference: TrackerFieldDifference,
) -> None:
    if not difference.is_changed:
        return
    if _is_empty(difference.comparison) and not _is_empty(difference.reference):
        reference_cell.fill = _ADDED_FILL
    elif _is_empty(difference.reference) and not _is_empty(difference.comparison):
        comparison_cell.fill = _REMOVED_FILL
    else:
        reference_cell.fill = _CHANGED_FILL
        comparison_cell.fill = _CHANGED_FILL


def _is_empty(value: Any) -> bool:
    return value is None or value == "" or value == [] or value == {}


def _kind_fill(kind: BaselineComparisonKind) -> PatternFill:
    return {
        BaselineComparisonKind.ADDED: _ADDED_FILL,
        BaselineComparisonKind.REMOVED: _REMOVED_FILL,
        BaselineComparisonKind.CHANGED: _CHANGED_FILL,
        BaselineComparisonKind.UNCHANGED: _UNCHANGED_FILL,
    }[kind]


def _set_safe_value(cell: Cell, value: Any, *, context: str) -> None:
    cell.font = BODY_FONT
    if isinstance(value, str):
        normalized = normalize_excel_text(value)
        text_units = excel_text_units(normalized)
        if text_units > EXCEL_MAX_CELL_TEXT:
            raise BaselineExportError(
                f"{context} 내용이 Excel 셀 길이 한도({EXCEL_MAX_CELL_TEXT:,}자)를 "
                f"초과합니다(현재 {text_units:,}자)."
            )
        line_feeds = normalized.count("\n")
        if line_feeds > EXCEL_MAX_CELL_LINE_FEEDS:
            raise BaselineExportError(
                f"{context} 내용이 Excel 셀 줄바꿈 한도"
                f"({EXCEL_MAX_CELL_LINE_FEEDS:,}개)를 초과합니다"
                f"(현재 {line_feeds:,}개)."
            )
        if normalized.startswith(("=", "+", "@")) or (
            normalized.startswith("-") and normalized != "-"
        ):
            normalized = f"'{normalized}"
        cell.value = normalized
        return
    cell.value = value


__all__ = [
    "EXCEL_MAX_CELL_LINE_FEEDS",
    "EXCEL_MAX_CELL_TEXT",
    "LONG_VALUE_CELL_LINE_FEEDS",
    "LONG_VALUE_CELL_TEXT",
    "LONG_VALUE_SHEET_TITLE",
    "BaselineExportError",
    "BaselineExportField",
    "BaselineExportSummary",
    "baseline_export_fields",
    "create_baseline_comparison_workbook",
    "export_baseline_comparison_xlsx",
]

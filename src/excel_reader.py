from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook


OPENPYXL_PREVIEW_SCAN_ROW_LIMIT = 10_000


class ExcelReader:
    """Excel 파일에서 raw DataFrame만 읽어오는 입력 전용 reader다."""

    def __init__(self, header_row: int = 1, summary_col: str = "Summary", logger=None):
        self.header_row = header_row
        self.summary_col = summary_col
        self.logger = logger

    @staticmethod
    def is_blank(value: Any) -> bool:
        if value is None:
            return True
        if isinstance(value, float) and pd.isna(value):
            return True
        if isinstance(value, str) and value.strip() == "":
            return True
        return False

    @staticmethod
    def _normalize_headers(headers: list[Any]) -> list[str]:
        return [
            str(header).strip() if header is not None else f"Unnamed_{index}"
            for index, header in enumerate(headers)
        ]

    @classmethod
    def _trim_trailing_blank_headers(cls, headers: list[Any]) -> list[Any]:
        """서식만 남은 오른쪽 열을 실제 헤더 범위에서 제외한다."""
        last_value_index = next(
            (
                index
                for index in range(len(headers) - 1, -1, -1)
                if not cls.is_blank(headers[index])
            ),
            None,
        )
        if last_value_index is None:
            return headers
        return headers[: last_value_index + 1]

    @staticmethod
    def _supports_openpyxl(file_path: str) -> bool:
        return Path(file_path).suffix.lower() in {".xlsx", ".xlsm", ".xltx", ".xltm"}

    @staticmethod
    def _resolve_openpyxl_sheet(workbook, sheet_name: str | int):
        if isinstance(sheet_name, str):
            normalized = sheet_name.strip()
            if normalized in workbook.sheetnames:
                return workbook[normalized]
            if normalized.isdigit():
                sheet_index = int(normalized)
                if 0 <= sheet_index < len(workbook.worksheets):
                    return workbook.worksheets[sheet_index]
        elif isinstance(sheet_name, int) and 0 <= sheet_name < len(workbook.worksheets):
            return workbook.worksheets[sheet_name]

        raise ValueError(f"시트를 찾을 수 없습니다: {sheet_name}")

    @staticmethod
    def _normalize_row(values: list[Any], width: int) -> list[Any]:
        normalized = list(values)
        if len(normalized) < width:
            normalized += [None] * (width - len(normalized))
        elif len(normalized) > width:
            normalized = normalized[:width]
        return normalized

    @classmethod
    def _normalize_cell_value(cls, value: Any) -> Any:
        if value is None:
            return None
        if isinstance(value, float):
            if pd.isna(value):
                return None
        return value

    @classmethod
    def _normalize_dataframe_values(cls, dataframe: pd.DataFrame) -> pd.DataFrame:
        if dataframe.empty:
            return dataframe

        work = dataframe.copy().astype(object)
        for column in work.columns:
            work[column] = pd.Series(
                [cls._normalize_cell_value(value) for value in work[column].tolist()],
                dtype=object,
            )
        return work

    def _openpyxl_sheet_names(self, file_path: str) -> list[str]:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()

    def _openpyxl_headers(self, file_path: str, sheet_name: str | int) -> list[str]:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            worksheet = self._resolve_openpyxl_sheet(workbook, sheet_name)
            row = next(
                worksheet.iter_rows(
                    min_row=self.header_row,
                    max_row=self.header_row,
                ),
                (),
            )
            raw_headers = self._trim_trailing_blank_headers(
                [cell.value for cell in row]
            )
            return self._normalize_headers(raw_headers)
        finally:
            workbook.close()

    @staticmethod
    def _create_xlwings_app(visible: bool = False):
        try:
            import xlwings as xw
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("Excel 구형 포맷 처리에는 xlwings 패키지가 필요합니다.") from exc

        app = xw.App(visible=visible, add_book=False)
        app.display_alerts = False
        app.screen_updating = False
        return app

    @staticmethod
    def _resolve_xlwings_sheet(workbook, sheet_name: str | int):
        if isinstance(sheet_name, str):
            normalized = sheet_name.strip()
            if normalized.isdigit():
                sheet_index = int(normalized)
                return workbook.sheets[sheet_index]
            return workbook.sheets[normalized]
        return workbook.sheets[sheet_name]

    def _xlwings_sheet_names(self, file_path: str) -> list[str]:
        app = self._create_xlwings_app()
        workbook = None

        try:
            workbook = app.books.open(file_path)
            return [sheet.name for sheet in workbook.sheets]
        finally:
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass
            app.quit()

    def _xlwings_headers(self, file_path: str, sheet_name: str | int) -> list[str]:
        app = self._create_xlwings_app()
        workbook = None

        try:
            workbook = app.books.open(file_path)
            sheet = self._resolve_xlwings_sheet(workbook, sheet_name)
            used_range = sheet.used_range
            column_count = max(int(used_range.columns.count or 0), 0)
            if column_count <= 0:
                return []
            first_column = int(used_range.column)
            last_column = first_column + column_count - 1
            values = sheet.range(
                (self.header_row, first_column),
                (self.header_row, last_column),
            ).value
            rows = self._normalize_xlwings_matrix(
                values,
                row_count=1,
                column_count=column_count,
            )
            return self._normalize_headers(rows[0] if rows else [])
        finally:
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass
            app.quit()

    def list_sheet_names(self, file_path: str) -> list[str]:
        if self._supports_openpyxl(file_path):
            return self._openpyxl_sheet_names(file_path)
        return self._xlwings_sheet_names(file_path)

    def read_headers(self, file_path: str, sheet_name: str | int) -> list[str]:
        if self._supports_openpyxl(file_path):
            return self._openpyxl_headers(file_path, sheet_name)
        return self._xlwings_headers(file_path, sheet_name)

    @classmethod
    def _normalize_xlwings_matrix(
        cls,
        values: Any,
        *,
        row_count: int,
        column_count: int,
    ) -> list[list[Any]]:
        """xlwings의 scalar/1D/2D 반환값을 요청한 범위 모양으로 맞춘다."""
        if row_count <= 0 or column_count <= 0:
            return []

        if isinstance(values, list):
            if values and isinstance(values[0], list):
                rows = [list(row) for row in values]
            elif row_count == 1:
                rows = [list(values)]
            elif column_count == 1:
                rows = [[value] for value in values]
            else:
                rows = [list(values)]
        else:
            rows = [[values]]

        normalized = [cls._normalize_row(row, column_count) for row in rows[:row_count]]
        while len(normalized) < row_count:
            normalized.append([None] * column_count)
        return normalized

    def read_preview_rows(
        self,
        file_path: str,
        sheet_name: str | int,
        *,
        max_rows: int = 10,
    ) -> tuple[list[str], list[list[Any]]]:
        """헤더와 제한된 데이터 행만 읽어 전체 시트 로딩을 피한다."""
        normalized_max_rows = max(int(max_rows), 0)
        if self._supports_openpyxl(file_path):
            workbook = load_workbook(file_path, read_only=True, data_only=True)
            try:
                worksheet = self._resolve_openpyxl_sheet(workbook, sheet_name)
                header_cells = next(
                    worksheet.iter_rows(
                        min_row=self.header_row,
                        max_row=self.header_row,
                    ),
                    (),
                )
                raw_headers = self._trim_trailing_blank_headers(
                    [cell.value for cell in header_cells]
                )
                headers = self._normalize_headers(raw_headers)
                rows: list[list[Any]] = []
                if normalized_max_rows <= 0:
                    return headers, rows
                scan_end_row = min(
                    int(worksheet.max_row or self.header_row),
                    self.header_row + OPENPYXL_PREVIEW_SCAN_ROW_LIMIT,
                )
                for row in worksheet.iter_rows(
                    min_row=self.header_row + 1,
                    max_row=scan_end_row,
                    max_col=len(headers),
                ):
                    normalized = self._normalize_row([cell.value for cell in row], len(headers))
                    if all(self.is_blank(value) for value in normalized):
                        continue
                    rows.append([self._normalize_cell_value(value) for value in normalized])
                    if len(rows) >= normalized_max_rows:
                        break
                if (
                    len(rows) < normalized_max_rows
                    and scan_end_row < int(worksheet.max_row or scan_end_row)
                ):
                    raise ValueError(
                        "Excel 시트의 사용 범위가 지나치게 큽니다. "
                        f"미리보기는 헤더 다음 {OPENPYXL_PREVIEW_SCAN_ROW_LIMIT:,}행까지만 확인했습니다. "
                        "실제 데이터 아래의 불필요한 행과 오른쪽의 불필요한 열을 삭제한 사본으로 다시 시도하세요."
                    )
                return headers, rows
            finally:
                workbook.close()

        app = self._create_xlwings_app()
        workbook = None
        try:
            workbook = app.books.open(file_path)
            sheet = self._resolve_xlwings_sheet(workbook, sheet_name)
            used_range = sheet.used_range
            column_count = max(int(used_range.columns.count or 0), 0)
            if column_count <= 0:
                return [], []
            first_column = int(used_range.column)
            last_column = first_column + column_count - 1
            header_values = sheet.range(
                (self.header_row, first_column),
                (self.header_row, last_column),
            ).value
            header_rows = self._normalize_xlwings_matrix(
                header_values,
                row_count=1,
                column_count=column_count,
            )
            headers = self._normalize_headers(header_rows[0] if header_rows else [])
            if normalized_max_rows <= 0:
                return headers, []

            last_used_row = int(used_range.last_cell.row)
            if last_used_row <= self.header_row:
                return headers, []

            rows: list[list[Any]] = []
            next_row = self.header_row + 1
            while next_row <= last_used_row and len(rows) < normalized_max_rows:
                remaining = normalized_max_rows - len(rows)
                chunk_size = max(remaining * 2, 16)
                chunk_end_row = min(last_used_row, next_row + chunk_size - 1)
                preview_values = sheet.range(
                    (next_row, first_column),
                    (chunk_end_row, last_column),
                ).value
                preview_rows = self._normalize_xlwings_matrix(
                    preview_values,
                    row_count=chunk_end_row - next_row + 1,
                    column_count=column_count,
                )
                for raw_row in preview_rows:
                    if all(self.is_blank(value) for value in raw_row):
                        continue
                    rows.append(
                        [self._normalize_cell_value(value) for value in raw_row]
                    )
                    if len(rows) >= normalized_max_rows:
                        break
                next_row = chunk_end_row + 1
            return headers, rows
        finally:
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass
            app.quit()

    def count_upload_rows(self, file_path: str, sheet_name: str | int) -> int:
        if self._supports_openpyxl(file_path):
            workbook = load_workbook(file_path, read_only=True, data_only=True)
            try:
                worksheet = self._resolve_openpyxl_sheet(workbook, sheet_name)
                header_row = next(
                    worksheet.iter_rows(
                        min_row=self.header_row,
                        max_row=self.header_row,
                    ),
                    (),
                )
                raw_headers = self._trim_trailing_blank_headers(
                    [cell.value for cell in header_row]
                )
                headers = self._normalize_headers(raw_headers)
                if self.summary_col not in headers:
                    raise ValueError(f"'{self.summary_col}' 컬럼을 찾을 수 없습니다.")
                summary_index = headers.index(self.summary_col)
                upload_row_count = 0
                for row in worksheet.iter_rows(
                    min_row=self.header_row + 1,
                    max_col=len(headers),
                ):
                    normalized_row = self._normalize_row([cell.value for cell in row], len(headers))

                    if all(self.is_blank(value) for value in normalized_row):
                        continue
                    if not self.is_blank(normalized_row[summary_index]):
                        upload_row_count += 1
                return upload_row_count
            finally:
                workbook.close()

        headers = self.read_headers(file_path, sheet_name)
        if self.summary_col not in headers:
            raise ValueError(f"'{self.summary_col}' 컬럼을 찾을 수 없습니다.")
        raw_df = self.read_excel(file_path=file_path, sheet_name=sheet_name)
        if raw_df.empty:
            return 0
        return int((~raw_df[self.summary_col].apply(self.is_blank)).sum())

    def _resolve_indent_level(self, cell: Any, summary_value: Any) -> int:
        try:
            alignment = getattr(cell, "alignment", None)
            indent = getattr(alignment, "indent", None) if alignment is not None else None
            if indent not in (None, ""):
                return int(indent)
        except Exception:
            pass
        try:
            return int(cell.api.IndentLevel)
        except Exception:
            if isinstance(summary_value, str):
                leading_spaces = len(summary_value) - len(summary_value.lstrip(" "))
                return leading_spaces // 4
        return 0

    def read_excel(self, file_path: str, sheet_name: str | int = 0, visible: bool = False) -> pd.DataFrame:
        """Excel 시트를 읽고 `_excel_row`, `_summary_indent` 메타정보를 포함한 raw DataFrame을 만든다."""
        if self._supports_openpyxl(file_path):
            workbook = load_workbook(file_path, read_only=True, data_only=True)
            try:
                worksheet = self._resolve_openpyxl_sheet(workbook, sheet_name)
                header_cells = next(
                    worksheet.iter_rows(
                        min_row=self.header_row,
                        max_row=self.header_row,
                    ),
                    (),
                )
                raw_headers = self._trim_trailing_blank_headers(
                    [cell.value for cell in header_cells]
                )
                headers = self._normalize_headers(raw_headers)
                if self.summary_col not in headers:
                    raise ValueError(f"'{self.summary_col}' 컬럼을 찾을 수 없습니다.")

                summary_col_index = headers.index(self.summary_col)
                records = []
                for excel_row, row in enumerate(
                    worksheet.iter_rows(
                        min_row=self.header_row + 1,
                        max_col=len(headers),
                    ),
                    start=self.header_row + 1,
                ):
                    normalized_row = self._normalize_row([cell.value for cell in row], len(headers))
                    if all(self.is_blank(value) for value in normalized_row):
                        continue

                    summary_cell = row[summary_col_index] if summary_col_index < len(row) else None
                    summary_value = normalized_row[summary_col_index]
                    indent_level = self._resolve_indent_level(summary_cell, summary_value)

                    record = dict(zip(headers, normalized_row))
                    record["_excel_row"] = excel_row
                    record["_summary_indent"] = indent_level
                    records.append(record)

                return self._normalize_dataframe_values(pd.DataFrame(records, dtype=object))
            finally:
                workbook.close()

        app = self._create_xlwings_app(visible=visible)
        workbook = None

        try:
            workbook = app.books.open(file_path)
            sheet = self._resolve_xlwings_sheet(workbook, sheet_name)
            used_range = sheet.used_range
            column_count = max(int(used_range.columns.count or 0), 0)
            if column_count <= 0:
                return pd.DataFrame()
            used_start_col = int(used_range.column)
            last_used_col = used_start_col + column_count - 1
            header_values = sheet.range(
                (self.header_row, used_start_col),
                (self.header_row, last_used_col),
            ).value
            header_rows = self._normalize_xlwings_matrix(
                header_values,
                row_count=1,
                column_count=column_count,
            )
            headers = self._normalize_headers(header_rows[0] if header_rows else [])
            if self.summary_col not in headers:
                raise ValueError(f"'{self.summary_col}' 컬럼을 찾을 수 없습니다.")

            summary_col_idx_1based = headers.index(self.summary_col) + 1
            last_used_row = int(used_range.last_cell.row)
            if last_used_row <= self.header_row:
                return pd.DataFrame(
                    columns=[*headers, "_excel_row", "_summary_indent"],
                    dtype=object,
                )
            data_values = sheet.range(
                (self.header_row + 1, used_start_col),
                (last_used_row, last_used_col),
            ).value
            data_rows = self._normalize_xlwings_matrix(
                data_values,
                row_count=last_used_row - self.header_row,
                column_count=column_count,
            )

            records = []
            for excel_row, normalized_row in enumerate(
                data_rows,
                start=self.header_row + 1,
            ):
                if all(self.is_blank(value) for value in normalized_row):
                    continue

                excel_col = used_start_col + (summary_col_idx_1based - 1)
                summary_cell = sheet.range((excel_row, excel_col))
                summary_value = normalized_row[summary_col_idx_1based - 1]
                indent_level = self._resolve_indent_level(summary_cell, summary_value)

                record = dict(zip(headers, normalized_row))
                record["_excel_row"] = excel_row
                record["_summary_indent"] = indent_level
                records.append(record)

            return self._normalize_dataframe_values(pd.DataFrame(records, dtype=object))
        finally:
            if workbook is not None:
                try:
                    workbook.close()
                except Exception:
                    pass
            app.quit()
